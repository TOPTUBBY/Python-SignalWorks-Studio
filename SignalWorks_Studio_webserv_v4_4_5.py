###################################################################################
# Module:   SignalWorks_Studio_webserv_v4_4_5.py
# Category: Engineering Data Visualization & Analysis Studio
# Version:  V4.4.5.2026
#
# Author: Patiphan Phakdeeburti
# Develop: DET DQT EVSBG Patiphan.Phak
#
# Description
# ------------------------------------------------------------------------------
# SignalWorks Studio: portable engineering CSV visualization, evaluation, analysis and reporting WebServer.
# Replaces Tkinter GUI with a browser-based interface while preserving:
#   - Smart CSV header detection from Timestamp / Date&Time / Date / Time markers
#   - Automatic timestamp parsing
#   - Eight Nissan OBC preset tabs
#   - Per-signal Primary/Left or Secondary/Right Y-axis assignment
#   - High-contrast signal palettes
#   - Solid Primary traces / dashed Secondary traces
#   - Separate compact Primary/Secondary legends
#   - Per-tab title and axis-range configuration
#   - Engineering Reference/Limit lines, Range bands and graph Comments
#   - PNG plot download
#   - Selectable custom-wide PDF pages + optional A4 multi-page export
#   - A4 landscape Word report generation
#
# Web architecture
# ------------------------------------------------------------------------------
# Framework : FastAPI + Uvicorn
# Plot      : Matplotlib using non-GUI Agg backend
# UI        : Embedded HTML/CSS/JavaScript + locally served Plotly.js (no CDN)
# Default   : localhost only (127.0.0.1)
# Port      : 8800
# LAN mode  : python SignalWorks_Studio_webserv_v4_4_5.py --lan
#
# Changelog
# ------------------------------------------------------------------------------
# V4.4.5 GitHub issue hotfix                                           2026-09-15
# Fixed / Improved:
#   - Parse dd.MM.yyyy HH:mm:ss[.fff] timestamps explicitly (Issue #2).
#   - Keep legends to a maximum of three rows by expanding columns horizontally.
#   - Preserve zoom/pan, drag mode and legend visibility when switching plot pages.
#   - Apply the saved interactive view to PNG/PDF/Word static exports.
#   - Add per-signal line-style selection (Auto/Solid/Dash/Dot/Dash-dot/Long dash).
#   - Improve contrast/readability of Dark mode surfaces and form controls.
#   - Add a dependency-free DOCX fallback when python-docx is unavailable (Issue #1).
#
# V4.4.5                                                               2026-09-03
# Added:
#   - Multi-file CSV selection from the browser file picker.
#   - Confirmation dialog listing selected CSV files before loading.
#   - Initial file ordering by browser-visible Last Modified time.
#   - Drag-and-drop file reordering before data binding.
#   - Bind Data workflow: keep the first file header and append only data rows
#     from subsequent files after removing their metadata/header sections.
#   - Header/column compatibility validation before binding.
#   - Bind progress indicator and source-file/row summary.
#   - Download Bound CSV before or after loading the merged dataset.
#   - Bound CSV can be loaded into the existing plot/preset/analysis workflow.
# Note:
#   - Web browsers do not expose Windows file Creation Time. Initial ordering
#     therefore uses File.lastModified; users can override the order by drag/drop.
#
# V4.4.4                                                               2026-09-02
# Added / Improved:
#   - Add Cycle Analysis profile mode: Auto / Temp cycling / Thermal shock
#   - Auto mode distinguishes slow-ramp temperature cycling from fast-transition thermal shock
#   - Thermal shock analysis uses dominant plateau temperatures instead of absolute overshoot extrema
#   - Report observed extrema separately from analysis Tmin/Tmax plateau centers
#   - Show median transition time and estimated transition rate for cycle-profile diagnostics
#   - Preserve legacy extrema-based behavior for temperature-cycling profiles
#   - Extend Cycle Analysis CSV export with profile and plateau diagnostics
#
# V4.4.3.1                                                             2026-09-01
# Fixed:
#   - Launcher startup timeout hotfix; plotting/export logic remains based on V4.4.3.
#   - Extended startup grace period for slow first-run Matplotlib/font-cache initialization.
#   - Failed launcher startup cleans up the spawned child process to avoid stale port 8000 ownership.
#
# V4.4.3                                                               2026-09-01
# Added:
#   - Select report tabs/pages before PDF or Word export.
#   - Browser Save As workflow with user-defined filename for PDF, Word, PNG, cycle CSV and .preset exports.
#   - User-defined preset display name when exporting a .preset file.
#   - File System Access API support for choosing save folder/path where the browser permits it, with download fallback.
# Changed:
#   - Report endpoints accept an explicit tab selection instead of always exporting all 8 plot tabs.
#
# V4.4.2.1                                                             2026-09-01
# Fixed:
#  - Replaced continuous long-page PDF with separate custom-wide report pages.
#  - Default wide report page size is 1200 x 650 pt (larger than A4 landscape 842 x 595 pt).
#  - Kept legacy A4 multi-page PDF as a separate export option.
# V4.4.2                                                               2026-09-01
# Added:
#   - Continuous long-page PDF report export with all plot tabs stacked on one PDF page
#   - Legacy A4 PDF endpoint retained as an alternate export option
#   - Rich graph comment styling: arrow toggle, border color/width, box width/height
#   - Multiline comment rendering preserves user-entered line breaks in Plotly and reports
# Changed:
#   - Comment manager provides post-placement box size and border controls
#
# V4.4.1                                                               2026-09-01
# Fixed / Improved:
#   - Fix transparent Graph Comment modal by defining panel/soft theme surfaces and strengthen modal backdrop
#   - Increase spacing between Plot Mode, Plot Setup, Data Series and Annotation cards
#   - Make the sticky sidebar reliably scroll to the final Annotation card even when all sections are expanded
#   - Preserve sidebar card open/closed state and scroll position across form submissions / page refreshes
#   - Preserve Reference/Range creator panel and annotation item expansion state per tab
#   - Add direct Plotly drag editing for Reference lines, Range bands and graph Comment boxes
#   - Synchronize dragged annotation positions back to the server and Annotation Manager
#   - Add editable X/Y anchor coordinates and text-box offset fields for graph Comments
#
# V4.4.0.1                                                             2026-09-01
# Fixed / Improved:
#   - Make Data Series card collapsible like Plot setup
#   - Convert Reference Line and Range Band creation forms into compact collapsed sub-panels
#   - Prevent Annotation & Limits controls from extending awkwardly beyond the visible sidebar
#   - Increase spacing between sidebar, plot and analysis cards for a cleaner visual hierarchy
#
# V4.4.0                                                               2026-09-01
# Added / Modified:
#   - Add Graph Annotation & Limits workspace for engineering review markup
#   - Add configurable horizontal Reference / Limit lines on Primary or Secondary Y axes
#   - Add configurable tolerance/range bands with fill color, opacity, border style and label
#   - Add click-to-place graph comments with arrow, text/background colors and font size
#   - Add annotation manager with show/hide, edit and delete controls
#   - Include Reference lines, Range bands and Comments in Interactive Plotly view and static PNG/PDF/Word exports
#   - Store Reference lines and Range bands in portable .preset files; keep Comments session-specific by default
#   - Preserve Cursor A/B overlay while combining it with persistent engineering annotations
#
# V4.3.1.2                                                             2026-09-01
# Fixed / Improved:
#   - Fix Preset dropdown form layout so the Import preset button remains inside the menu panel
#   - Make header dropdown menus mutually exclusive: opening one automatically closes the others
#   - Close open header dropdowns when clicking outside the menu or pressing Escape
#
# V4.3.1                                                               2026-09-01
# Added / Modified:
#   - Add System / Light / Dark appearance modes with browser-local preference memory
#   - Default appearance follows the operating-system/browser color scheme automatically
#   - Apply the selected appearance to the Plotly interactive graph as well as the surrounding Web UI
#   - Remove the visible Header Auto Detect control; Smart Header Detection continues silently in the background
#   - Reorganize the crowded top toolbar into compact Preset, Export and Theme dropdown menus
#   - Keep the CSV file selector and Load CSV action directly accessible as the primary workflow
#
# V4.3.0                                                               2026-09-01
# Added / Modified:
#   - Add Smart Header Detection: scan the first CSV column for Timestamp / Date&Time / Date / Time
#   - Automatically use the detected row as the CSV header; no manual Skip Rows counting required
#   - Add first-row fallback for ordinary CSV files whose first header is not one of the marker names
#   - Show detected header row and timestamp source in the dataset information strip
#   - Add web-based Cycle Analysis derived from graphplot_common.py V3.1 cycle-analysis logic
#   - Detect Tmin/Tmax tolerance-band plateau events and complete A -> B -> A cycles
#   - Calculate cycle count, per-cycle time, soak time at Tmin/Tmax, and average durations
#   - Support real CSV timestamps or user-defined fallback seconds/sample when timestamp parsing fails
#   - Add Cycle Analysis CSV export from the browser
#
# V4.2.4.1                                                             2026-08-31
# Added / Modified:
#   - Add stale GraphPlot server detection for port 8000 during Windows launch
#   - Identify the PID/process command line that owns port 8000 before any forced stop
#   - Offer to stop older V4.x GraphPlot servers automatically and then continue startup
#   - Use clean /api/shutdown first for V4.2.4+ servers, with PID taskkill fallback only for recognized GraphPlot processes
#   - Never auto-kill unknown/non-GraphPlot applications using port 8000
#   - Add launcher inspection mode and clearer Stable / Hidden / Debug lifecycle behavior
#   - Add VERSION_HISTORY.md covering the WebServer development history from V4.0 onward
#
# V4.2.4                                                               2026-08-31
# Added / Modified:
#   - Add browser-client heartbeat and automatic WebServer shutdown after all tabs close
#   - Add page-close notification plus stale-client fallback for reliable process cleanup
#   - Add /api/health and localhost-only /api/shutdown endpoints for launcher management
#   - Replace uvicorn.run shortcut with a managed Uvicorn Server instance
#   - Add robust launcher helper that reuses an already-running V4.2.4 server instead of colliding on port 8000
#   - Reorganize Windows launchers into Stable / Hidden / Debug folders
#
# V4.2.3                                                               2026-08-31
# Added / Modified:
#   - Shorten signal axis selector labels to Primary / Secondary
#   - Add semantic color styling to top toolbar actions for faster visual scanning
#   - Add hidden Windows launchers using pythonw/VBScript so no Command Prompt remains open
#   - Keep debug console launchers available separately for troubleshooting
#
# V4.2.2                                                               2026-08-31
# Added / Modified:
#   - Align all top toolbar controls on one compact horizontal control plane
#   - Normalize header input/button heights and file-picker styling
#   - Widen the settings sidebar for improved signal readability
#   - Rework Data Series rows so signal names receive the main available width
#   - Move PRESET marker into the signal-name cell instead of reserving a grid column
#   - Preserve full signal names in hover tooltips while keeping rows compact
#
# V4.2.1                                                               2026-08-31
# Added / Modified:
#   - Fix persistent Loading graph placeholder before Plotly rendering
#   - Make interactive graph height responsive to browser viewport
#   - Add per-signal manual color selection with one-click Auto color restore
#   - Add portable JSON-based .preset import/export for signal selections and settings
#   - Add Normal Plot Mode when no preset is loaded (no automatic signal selection)
#   - Allow preset import before or after CSV upload and match signals by column name
#   - Export title, axes, selected signals, axis assignment and manual colors
#   - Preserve imported preset across subsequent CSV uploads in the same browser session
#
# V4.2.0                                                               2026-08-31
# Added / Modified:
#   - Professional minimal UI redesign with a compact top bar and two-column workspace
#   - Move plot settings and signal selection into a clean control sidebar
#   - Add Cursor A / Cursor B point inspection directly on the interactive graph
#   - Add A-B shaded analysis region and Delta Time / Delta X calculation
#   - Add per-signal Cursor A value, Cursor B value and Delta Value table
#   - Add Min / Max / Average / sample-count statistics for current zoom or A-B range
#   - Statistics automatically follow zoom/pan and visible Plotly legend traces
#   - Add compact analysis cards and a scrollable professional statistics table
#   - Preserve V4.1 offline Plotly, dual-axis behavior and Matplotlib report exports
#
# V4.1.0                                                               2026-08-31
# Added / Modified:
#   - Replace static browser Graph Preview image with interactive Plotly graph
#   - Zoom / Pan / mouse-wheel zoom using Plotly browser controls
#   - Hover inspection with timestamp and signal values
#   - Click Legend to show/hide traces; double-click to isolate a trace
#   - Group Legend into PRIMARY / LEFT and SECONDARY / RIGHT sections
#   - Add Reset View, Autoscale and Show All Signals controls
#   - Use WebGL traces for improved responsiveness with large CSV datasets
#   - Serve Plotly.js locally from FastAPI (offline/LAN capable; no CDN)
#   - Keep Matplotlib export engine unchanged for PNG/PDF/Word reports
#
# V4.0.0                                                               2026-08-31
# Added / Modified:
#   - Replace Tkinter desktop GUI with local WebServer UI
#   - Browser-based CSV upload and signal configuration
#   - Browser navigation for all 8 preset tabs
#   - Primary/Secondary axis assignment visible per signal
#   - Bulk assign checked signals to Primary or Secondary axis
#   - Server-side Matplotlib plotting; no internet/CDN needed
#   - Preserve V3.5.2 high-contrast colors and line-style separation
#   - Preserve separate compact Primary/Secondary legends
#   - PNG, PDF and Word export through browser downloads
#   - Per-browser session data to reduce conflicts between clients
#   - Localhost-only default with optional --lan access
###################################################################################

from __future__ import annotations

import argparse
import csv
import html
import io
import ipaddress
import json
import math
import os
import re
import socket
import struct
import sys
import threading
import zipfile
import textwrap
import time
import uuid
import webbrowser
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple
from urllib.parse import quote

import matplotlib

# IMPORTANT: Web server has no Tkinter display. Use a headless Matplotlib backend.
matplotlib.use("Agg")

import matplotlib.dates as mdates
import matplotlib.image as mpimg
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.ticker import MultipleLocator

try:
    import plotly.graph_objects as go
    import plotly.io as pio
    from plotly.offline import get_plotlyjs
except ImportError as exc:
    raise SystemExit(
        "Missing Plotly library. Install it with: pip install plotly\n"
        f"Original error: {exc}"
    )

try:
    from fastapi import FastAPI, Request, UploadFile, File
    from fastapi.responses import HTMLResponse, RedirectResponse, Response, JSONResponse
    from starlette.datastructures import FormData
except ImportError as exc:
    raise SystemExit(
        "Missing WebServer libraries. Install them with:\n\n"
        "  pip install fastapi uvicorn python-multipart\n\n"
        f"Original error: {exc}"
    )


APP_VERSION = "V4.4.5.2026"

# Custom wide PDF report page size (PDF points, 72 pt = 1 inch).
# Standard A4 landscape is approximately 842 x 595 pt.
WIDE_REPORT_WIDTH_PT = 1200
WIDE_REPORT_HEIGHT_PT = 650
WIDE_REPORT_FIGSIZE_IN = (WIDE_REPORT_WIDTH_PT / 72.0, WIDE_REPORT_HEIGHT_PT / 72.0)
APP_TITLE = "SignalWorks Studio"
APP_TAGLINE = "Engineering Data Visualization · Evaluation · Analysis · Reporting"
DEFAULT_PORT = 8800
SESSION_COOKIE = "signalworks_studio_sid"
SESSION_TTL_SECONDS = 12 * 60 * 60

# WebServer lifecycle management (V4.2.4)
# A browser tab sends a heartbeat every few seconds. Closing/reloading a page also
# sends a best-effort release message. When all browser clients are gone, the
# managed Uvicorn server exits cleanly so the Python process and TCP port are freed.
CLIENT_HEARTBEAT_STALE_SECONDS = 60.0
DEFAULT_IDLE_SHUTDOWN_SECONDS = 10.0
CLIENT_HEARTBEAT_INTERVAL_MS = 2500

_lifecycle_lock = threading.Lock()
_client_last_seen: Dict[str, float] = {}
_ever_had_browser_client = False
_no_client_since: Optional[float] = None
_managed_uvicorn_server = None
_auto_shutdown_enabled = True
_idle_shutdown_seconds = DEFAULT_IDLE_SHUTDOWN_SECONDS

# Interactive browser graph settings (V4.2)
INTERACTIVE_PLOT_HEIGHT = 620
PLOTLY_JS = get_plotlyjs()


@dataclass
class TabConfig:
    name: str
    title: str
    left_min: float
    left_max: float
    left_step: float
    left_label: str
    right_min: Optional[float] = None
    right_max: Optional[float] = None
    right_step: Optional[float] = None
    right_label: Optional[str] = None
    signals: List[str] = field(default_factory=list)


@dataclass
class TabState:
    title: str
    x_label: str
    left_label: str
    right_label: str
    left_min: str
    left_max: str
    left_step: str
    right_min: str
    right_max: str
    right_step: str
    auto_title: bool = False
    selected_signals: List[str] = field(default_factory=list)
    axis_assignments: Dict[str, str] = field(default_factory=dict)
    # Manual color overrides only. Missing signal => automatic palette color.
    color_assignments: Dict[str, str] = field(default_factory=dict)
    # Optional per-signal line style. Missing/"auto" keeps the axis default style.
    line_style_assignments: Dict[str, str] = field(default_factory=dict)
    # V4.4 engineering markup. Reference/range items are portable preset data.
    # Comments are intentionally session-specific unless a future format opts in.
    reference_lines: List[dict] = field(default_factory=list)
    range_bands: List[dict] = field(default_factory=list)
    comments: List[dict] = field(default_factory=list)
    revision: int = 0
    # Reset/preset replacement creates a fresh view; styling alone keeps zoom.
    view_token: str = field(default_factory=lambda: uuid.uuid4().hex[:16])


@dataclass
class BrowserSession:
    sid: str
    created_at: float = field(default_factory=time.time)
    last_access: float = field(default_factory=time.time)
    df: Optional[pd.DataFrame] = None
    filename: str = ""
    numeric_cols: List[str] = field(default_factory=list)
    timestamp_col_name: str = ""
    parsed_info: str = ""
    row_count: int = 0
    # V4.3 smart-header diagnostics. header_row_index is zero-based.
    header_row_index: int = 0
    header_detection: str = "Auto header pending"
    header_marker: str = ""
    delimiter: str = ","
    status: str = "Ready"
    status_type: str = "info"
    tabs: List[TabState] = field(default_factory=list)
    plot_cache: Dict[int, Tuple[int, bytes]] = field(default_factory=dict)
    # Portable signal preset currently loaded in this browser session.
    loaded_preset: Optional[dict] = None
    preset_name: str = ""
    preset_filename: str = ""
    # V4.4.5 multi-file CSV binding workspace.
    bound_csv_bytes: bytes = b""
    bound_filename: str = ""
    bound_source_files: List[str] = field(default_factory=list)
    bound_source_rows: List[int] = field(default_factory=list)
    bound_total_rows: int = 0
    loaded_source_files: List[str] = field(default_factory=list)
    # Changes whenever a CSV/bound dataset is loaded; scopes browser view persistence.
    dataset_token: str = ""


# ------------------------------------------------------------------------------
# Presets from V3.5.2
# ------------------------------------------------------------------------------
TAB_PRESETS: List[TabConfig] = [
    TabConfig(
        name="Tab1",
        title="Overall Parameter",
        left_min=0,
        left_max=1000,
        left_step=50,
        left_label="Voltage (V)",
        right_min=-110,
        right_max=110,
        right_step=10,
        right_label="Current (A)/ Temperature (°C)/ Humidity (%Rh)",
        signals=[
            "Vac_In[In]", "Vac_L1[Ex]", "Vdc_Out[In]", "Vdc_Out[Ex]",
            "Idc_Out[In]", "Idc_Out[Ex]", "Cham_Act_Temp", "Cham_Set_Humi",
            "Cham_Act_Humi", "Cool_Act_Temp",
        ],
    ),
    TabConfig(
        name="Tab2",
        title="Input Parameter",
        left_min=0,
        left_max=500,
        left_step=20,
        left_label="Voltage (V)",
        right_min=0,
        right_max=50,
        right_step=2,
        right_label="Current (A)",
        signals=["Vdc_Out[In]", "Vdc_Out[Ex]", "Idc_Out[In]", "Idc_Out[Ex]", "HVCurrentRequest"],
    ),
    TabConfig(
        name="Tab3",
        title="Ubatt (KL30) Parameter",
        left_min=0,
        left_max=16,
        left_step=1,
        left_label="Voltage (V)",
        right_min=0,
        right_max=2,
        right_step=0.2,
        right_label="Current (A)",
        signals=["Ubatt[In]", "Ubatt[Ex]", "Ibatt[Ex]"],
    ),
    TabConfig(
        name="Tab4",
        title="Output Parameter",
        left_min=0,
        left_max=300,
        left_step=10,
        left_label="Voltage (V)",
        right_min=0,
        right_max=60,
        right_step=2,
        right_label="Current (A)",
        signals=[
            "Vac_In[In]", "Vac_L1[Ex]", "Iac_In[In]", "Iac_L1[Ex]", "Freq[Ex]",
            "V2G_VoltSetPoint", "V2L_VoltSetPoint_v2", "V2L_OutletCurlim",
        ],
    ),
    TabConfig(
        name="Tab5",
        title="Power Parameter",
        left_min=0,
        left_max=12000,
        left_step=500,
        left_label="Power (W)",
        right_min=0,
        right_max=12,
        right_step=1,
        right_label="Active Power (kW)",
        signals=["P_In[In]", "P_In[Ex]", "Pout[In]", "P_Out[In]", "P_Out[Ex]", "V2G_ActivePower"],
    ),
    TabConfig(
        name="Tab6",
        title="Internal status",
        left_min=0,
        left_max=7,
        left_step=1,
        left_label="StateRequest/LoadState (1: Charge, 2: PowernOn, 5: V2G, 6: V2L)",
        right_min=0,
        right_max=4,
        right_step=1,
        right_label=(
            "Failure (0: No Alert, 1: Alert)/ FaultType (0: No failure, 1: Power derating)/ "
            "Interlock (1: Open, 2: Close)"
        ),
        signals=["CHGStateRequest", "CHGLoadState", "CHG_FaultType", "CHG_ChargeFailureDisplay", "CHGMainsACInterlockState"],
    ),
    TabConfig(
        name="Tab7",
        title="Internal/unit Temperature",
        left_min=-20,
        left_max=140,
        left_step=10,
        left_label="Temperature (°C)",
        right_min=None,
        right_max=None,
        right_step=None,
        right_label=None,
        signals=["Temp_DUT[In]", "Temp_LLC_Mag[In]", "Temp_PFC_Mag[In]", "Temp_WaterDUT[In]", "Temp_Inlet[In]"],
    ),
    TabConfig(
        name="Tab8",
        title="Chamber/Chiller",
        left_min=-20,
        left_max=120,
        left_step=5,
        left_label="Temperature (°C)/Humidity (%Rh)",
        right_min=None,
        right_max=None,
        right_step=None,
        right_label="",
        signals=["Cham_Act_Temp", "Cham_Set_Humi", "Cham_Act_Humi", "Cool_Act_Temp"],
    ),
]


# ------------------------------------------------------------------------------
# V3.5.2 visual theme
# ------------------------------------------------------------------------------
PRIMARY_AXIS_COLOR = "#1F4E79"
SECONDARY_AXIS_COLOR = "#A61C00"

PRIMARY_LINE_COLORS = [
    "#0072B2", "#E69F00", "#009E73", "#CC79A7",
    "#D55E00", "#56B4E9", "#6A3D9A", "#B15928",
    "#17BECF", "#BCBD22", "#E377C2", "#1F77B4",
    "#2CA02C", "#FF7F0E", "#7F7F7F", "#111111",
]

SECONDARY_LINE_COLORS = [
    "#C62828", "#1565C0", "#6A1B9A", "#00897B",
    "#EF6C00", "#AD1457", "#558B2F", "#5D4037",
    "#00838F", "#F9A825", "#4527A0", "#2E7D32",
    "#D81B60", "#8E24AA", "#616161", "#000000",
]

SECONDARY_DASH_PATTERNS = [
    (6, 2),
    (3, 1.5),
    (8, 2, 2, 2),
    (2, 1.5),
]

DEFAULT_SUBPLOT_PARAMS = {
    "left": 0.09,
    "bottom": 0.25,
    "right": 0.915,
    "top": 0.94,
    "wspace": 0.2,
    "hspace": 0.2,
}


# ------------------------------------------------------------------------------
# Session management
# ------------------------------------------------------------------------------
SESSIONS: Dict[str, BrowserSession] = {}
SESSIONS_LOCK = threading.RLock()


def _fmt_default(value: Optional[float]) -> str:
    return "" if value is None else str(value)


def create_default_tab_states() -> List[TabState]:
    """Create eight clean tabs for Normal Plot Mode.

    V4.2.2 deliberately does not auto-select Nissan signals unless a .preset file
    has been imported. This makes the WebServer reusable with other projects.
    """
    states: List[TabState] = []
    for idx in range(len(TAB_PRESETS)):
        states.append(
            TabState(
                title=f"Plot {idx + 1}",
                x_label="Time",
                left_label="Primary Y",
                right_label="Secondary Y",
                left_min="",
                left_max="",
                left_step="",
                right_min="",
                right_max="",
                right_step="",
            )
        )
    return states


def _cleanup_sessions() -> None:
    now = time.time()
    with SESSIONS_LOCK:
        expired = [sid for sid, state in SESSIONS.items() if now - state.last_access > SESSION_TTL_SECONDS]
        for sid in expired:
            SESSIONS.pop(sid, None)


def get_or_create_session(request: Request) -> Tuple[BrowserSession, bool]:
    _cleanup_sessions()
    sid = request.cookies.get(SESSION_COOKIE, "")
    with SESSIONS_LOCK:
        state = SESSIONS.get(sid)
        if state is not None:
            state.last_access = time.time()
            return state, False

        sid = uuid.uuid4().hex
        state = BrowserSession(sid=sid, tabs=create_default_tab_states())
        SESSIONS[sid] = state
        return state, True


def response_with_session_cookie(response: Response, state: BrowserSession, is_new: bool) -> Response:
    if is_new:
        response.set_cookie(
            key=SESSION_COOKIE,
            value=state.sid,
            max_age=SESSION_TTL_SECONDS,
            httponly=True,
            samesite="lax",
        )
    return response


# ------------------------------------------------------------------------------
# CSV helpers
# ------------------------------------------------------------------------------
def is_current_col(name: str) -> bool:
    n = name.lower()
    return (
        ("_i_" in n)
        or n.startswith("i")
        or n.endswith("_i")
        or ("current" in n)
        or ("interlock" in n)
        or ("failure" in n)
        or ("fault" in n)
        or ("cur" in n)
    )


def default_axis_for_signal(tab_id: int, header_name: str) -> str:
    if tab_id == 0:
        special_right_signals = {
            "Idc_Out[In]",
            "Idc_Out[Ex]",
            "Cham_Act_Temp",
            "Cham_Set_Humi",
            "Cham_Act_Humi",
            "Cool_Act_Temp",
        }
        if header_name in special_right_signals:
            return "right"

    if tab_id == 4 and header_name == "V2G_ActivePower":
        return "right"

    return "right" if is_current_col(header_name) else "left"


def _valid_hex_color(value: str) -> bool:
    value = str(value or "").strip()
    if len(value) != 7 or not value.startswith("#"):
        return False
    try:
        int(value[1:], 16)
        return True
    except ValueError:
        return False


def _auto_color_for_position(axis_target: str, primary_index: int, secondary_index: int) -> str:
    if axis_target == "right":
        return SECONDARY_LINE_COLORS[secondary_index % len(SECONDARY_LINE_COLORS)]
    return PRIMARY_LINE_COLORS[primary_index % len(PRIMARY_LINE_COLORS)]


SIGNAL_LINE_STYLES = ("solid", "dash", "dot", "dashdot", "longdash")


def _effective_signal_color(ts: TabState, header: str, axis_target: str, primary_index: int, secondary_index: int) -> str:
    manual = ts.color_assignments.get(header, "")
    if _valid_hex_color(manual):
        return manual
    return _auto_color_for_position(axis_target, primary_index, secondary_index)


def _effective_signal_dash(ts: TabState, header: str, axis_target: str, secondary_index: int) -> str:
    # Missing/auto keeps the original Primary solid / Secondary patterned behavior.
    manual = str(ts.line_style_assignments.get(header, "auto") or "auto").lower()
    if manual in SIGNAL_LINE_STYLES:
        return manual
    if axis_target == "right":
        defaults = ("dash", "dot", "dashdot", "longdash")
        return defaults[secondary_index % len(defaults)]
    return "solid"


def _generic_tab_state(tab_id: int, revision: int = 0) -> TabState:
    return TabState(
        title=f"Plot {tab_id + 1}",
        x_label="Time",
        left_label="Primary Y",
        right_label="Secondary Y",
        left_min="",
        left_max="",
        left_step="",
        right_min="",
        right_max="",
        right_step="",
        auto_title=False,
        revision=revision,
    )


def reset_tab_normal(state: BrowserSession, tab_id: int) -> None:
    old_revision = state.tabs[tab_id].revision if 0 <= tab_id < len(state.tabs) else 0
    state.tabs[tab_id] = _generic_tab_state(tab_id, old_revision + 1)
    state.plot_cache.pop(tab_id, None)


def initialize_tabs_normal(state: BrowserSession) -> None:
    """Normal Plot Mode: eight clean tabs and no automatic signal selection."""
    old_revisions = [t.revision for t in state.tabs] if state.tabs else [0] * len(TAB_PRESETS)
    state.tabs = [
        _generic_tab_state(i, (old_revisions[i] if i < len(old_revisions) else 0) + 1)
        for i in range(len(TAB_PRESETS))
    ]
    state.plot_cache.clear()


def _clean_annotation_axis(value: object) -> str:
    return "right" if str(value or "left").lower() in ("right", "secondary", "y2") else "left"


def _clean_dash(value: object) -> str:
    value = str(value or "dash").lower()
    return value if value in ("solid", "dash", "dot", "dashdot", "longdash") else "dash"


def _clean_width(value: object, default: float = 1.5) -> float:
    try:
        return max(0.5, min(8.0, float(value)))
    except Exception:
        return default


def _clean_opacity(value: object, default: float = 0.10) -> float:
    try:
        v = float(value)
        if v > 1.0:
            v /= 100.0
        return max(0.0, min(0.8, v))
    except Exception:
        return default


def _new_annotation_id(prefix: str) -> str:
    return f"{prefix}_{uuid.uuid4().hex[:10]}"


def _normalize_reference_lines(items: object) -> List[dict]:
    out: List[dict] = []
    if not isinstance(items, list):
        return out
    for item in items[:100]:
        if not isinstance(item, dict):
            continue
        try:
            value = float(item.get("value"))
        except Exception:
            continue
        color = str(item.get("color") or "#D92D20").upper()
        if not _valid_hex_color(color):
            color = "#D92D20"
        out.append({
            "id": str(item.get("id") or _new_annotation_id("line"))[:80],
            "label": str(item.get("label") or "Reference")[:120],
            "value": value,
            "axis": _clean_annotation_axis(item.get("axis")),
            "color": color,
            "width": _clean_width(item.get("width"), 1.5),
            "dash": _clean_dash(item.get("dash")),
            "show_label": bool(item.get("show_label", True)),
            "visible": bool(item.get("visible", True)),
        })
    return out


def _normalize_range_bands(items: object) -> List[dict]:
    out: List[dict] = []
    if not isinstance(items, list):
        return out
    for item in items[:100]:
        if not isinstance(item, dict):
            continue
        try:
            lo = float(item.get("min")); hi = float(item.get("max"))
        except Exception:
            continue
        if lo > hi:
            lo, hi = hi, lo
        line_color = str(item.get("line_color") or "#12B76A").upper()
        fill_color = str(item.get("fill_color") or "#12B76A").upper()
        if not _valid_hex_color(line_color): line_color = "#12B76A"
        if not _valid_hex_color(fill_color): fill_color = "#12B76A"
        out.append({
            "id": str(item.get("id") or _new_annotation_id("band"))[:80],
            "label": str(item.get("label") or "Accept range")[:120],
            "min": lo, "max": hi,
            "axis": _clean_annotation_axis(item.get("axis")),
            "line_color": line_color, "fill_color": fill_color,
            "opacity": _clean_opacity(item.get("opacity"), 0.10),
            "width": _clean_width(item.get("width"), 1.2),
            "dash": _clean_dash(item.get("dash")),
            "show_label": bool(item.get("show_label", True)),
            "visible": bool(item.get("visible", True)),
        })
    return out


def _normalize_comments(items: object) -> List[dict]:
    out: List[dict] = []
    if not isinstance(items, list):
        return out
    for item in items[:200]:
        if not isinstance(item, dict) or not str(item.get("text", "")).strip():
            continue
        try:
            y = float(item.get("y"))
        except Exception:
            continue
        tc = str(item.get("text_color") or "#101828").upper()
        bg = str(item.get("bg_color") or "#FFFFFF").upper()
        if not _valid_hex_color(tc): tc = "#101828"
        if not _valid_hex_color(bg): bg = "#FFFFFF"
        try:
            font_size = max(8, min(28, int(float(item.get("font_size", 12)))))
        except Exception:
            font_size = 12
        try:
            ax_offset = float(item.get("ax", 32))
        except Exception:
            ax_offset = 32.0
        try:
            ay_offset = float(item.get("ay", -38))
        except Exception:
            ay_offset = -38.0
        border_color = str(item.get("border_color") or tc).upper()
        if not _valid_hex_color(border_color):
            border_color = tc
        try:
            border_width = max(0.0, min(8.0, float(item.get("border_width", 1.0))))
        except Exception:
            border_width = 1.0
        try:
            box_width = max(0, min(900, int(float(item.get("box_width", 0)))))
        except Exception:
            box_width = 0
        try:
            box_height = max(0, min(600, int(float(item.get("box_height", 0)))))
        except Exception:
            box_height = 0
        arrow_color = str(item.get("arrow_color") or border_color).upper()
        if not _valid_hex_color(arrow_color):
            arrow_color = border_color
        try:
            arrow_width = max(0.5, min(8.0, float(item.get("arrow_width", 1.2))))
        except Exception:
            arrow_width = 1.2
        out.append({
            "id": str(item.get("id") or _new_annotation_id("comment"))[:80],
            "text": str(item.get("text"))[:1200],
            "x": item.get("x"), "y": y,
            "axis": _clean_annotation_axis(item.get("axis")),
            "text_color": tc, "bg_color": bg,
            "border_color": border_color, "border_width": border_width,
            "box_width": box_width, "box_height": box_height,
            "font_size": font_size,
            "show_arrow": bool(item.get("show_arrow", True)),
            "arrow_color": arrow_color, "arrow_width": arrow_width,
            "ax": max(-800.0, min(800.0, ax_offset)),
            "ay": max(-800.0, min(800.0, ay_offset)),
            "visible": bool(item.get("visible", True)),
        })
    return out


def _preset_axis_dict(tab: dict, key: str) -> dict:
    axes = tab.get("axes", {}) if isinstance(tab.get("axes", {}), dict) else {}
    value = axes.get(key, {})
    return value if isinstance(value, dict) else {}


def normalize_preset_payload(payload: object) -> dict:
    if not isinstance(payload, dict):
        raise ValueError("Preset root must be a JSON object.")
    fmt = str(payload.get("format", ""))
    if fmt not in ("SignalWorksStudioPreset", "GraphPlotSignalPreset", "CSVDataPlotterPreset"):
        raise ValueError("Unsupported preset format.")
    version = int(payload.get("format_version", 1))
    if version != 1:
        raise ValueError(f"Unsupported preset format version: {version}")
    tabs = payload.get("tabs")
    if not isinstance(tabs, list) or not tabs:
        raise ValueError("Preset does not contain any tabs.")
    if len(tabs) > 32:
        raise ValueError("Preset contains too many tabs.")
    return payload


def apply_loaded_preset(state: BrowserSession) -> Tuple[int, int]:
    """Apply the stored portable preset. Returns (matched_signals, requested_signals)."""
    if not state.loaded_preset:
        initialize_tabs_normal(state)
        return 0, 0

    payload = normalize_preset_payload(state.loaded_preset)
    preset_tabs = payload.get("tabs", [])
    old_revisions = [t.revision for t in state.tabs] if state.tabs else [0] * len(TAB_PRESETS)
    new_tabs: List[TabState] = []
    matched_total = 0
    requested_total = 0

    for tab_id in range(len(TAB_PRESETS)):
        revision = (old_revisions[tab_id] if tab_id < len(old_revisions) else 0) + 1
        if tab_id >= len(preset_tabs) or not isinstance(preset_tabs[tab_id], dict):
            new_tabs.append(_generic_tab_state(tab_id, revision))
            continue

        ptab = preset_tabs[tab_id]
        primary = _preset_axis_dict(ptab, "primary")
        secondary = _preset_axis_dict(ptab, "secondary")
        ts = TabState(
            title=str(ptab.get("title") or ptab.get("name") or f"Plot {tab_id + 1}"),
            x_label=str(ptab.get("x_label", "Time")),
            left_label=str(primary.get("label", "Primary Y")),
            right_label=str(secondary.get("label", "Secondary Y")),
            left_min=str(primary.get("min", "") if primary.get("min", "") is not None else ""),
            left_max=str(primary.get("max", "") if primary.get("max", "") is not None else ""),
            left_step=str(primary.get("step", "") if primary.get("step", "") is not None else ""),
            right_min=str(secondary.get("min", "") if secondary.get("min", "") is not None else ""),
            right_max=str(secondary.get("max", "") if secondary.get("max", "") is not None else ""),
            right_step=str(secondary.get("step", "") if secondary.get("step", "") is not None else ""),
            auto_title=bool(ptab.get("auto_title", False)),
            revision=revision,
        )

        ts.reference_lines = _normalize_reference_lines(ptab.get("reference_lines", []))
        ts.range_bands = _normalize_range_bands(ptab.get("range_bands", []))

        signals = ptab.get("signals", [])
        if not isinstance(signals, list):
            signals = []
        for entry in signals:
            if isinstance(entry, str):
                name, axis, color, line_style = entry, "left", None, "auto"
            elif isinstance(entry, dict):
                name = str(entry.get("name", "")).strip()
                axis = str(entry.get("axis", "left")).lower()
                color = entry.get("color")
                line_style = str(entry.get("line_style", entry.get("dash", "auto")) or "auto").lower()
            else:
                continue
            if not name:
                continue
            requested_total += 1
            if state.df is not None and name not in state.numeric_cols:
                continue
            ts.selected_signals.append(name)
            ts.axis_assignments[name] = "right" if axis in ("right", "secondary", "y2") else "left"
            if color and _valid_hex_color(str(color)):
                ts.color_assignments[name] = str(color).upper()
            if line_style in SIGNAL_LINE_STYLES:
                ts.line_style_assignments[name] = line_style
            matched_total += 1

        new_tabs.append(ts)

    state.tabs = new_tabs
    state.plot_cache.clear()
    return matched_total, requested_total


def reset_tab_to_loaded_mode(state: BrowserSession, tab_id: int) -> None:
    if state.loaded_preset:
        # Rebuild all tabs from the imported preset for deterministic matching,
        # then preserve the other tabs exactly as they currently are.
        current = state.tabs[:]
        saved_comments = list(current[tab_id].comments)
        apply_loaded_preset(state)
        rebuilt = state.tabs[tab_id]
        rebuilt.comments = saved_comments
        state.tabs = current
        rebuilt.revision = current[tab_id].revision + 1
        state.tabs[tab_id] = rebuilt
        state.plot_cache.pop(tab_id, None)
    else:
        reset_tab_normal(state, tab_id)


def build_preset_payload_from_state(state: BrowserSession, preset_name: Optional[str] = None) -> dict:
    tabs = []
    for tab_id, ts in enumerate(state.tabs):
        signal_entries = []
        for signal in ts.selected_signals:
            axis = ts.axis_assignments.get(signal, "left")
            signal_entries.append({
                "name": signal,
                "axis": axis,
                # null means use automatic palette when the preset is imported.
                "color": ts.color_assignments.get(signal) if _valid_hex_color(ts.color_assignments.get(signal, "")) else None,
                # "auto" preserves legacy Primary solid / Secondary patterned defaults.
                "line_style": ts.line_style_assignments.get(signal, "auto"),
            })
        tabs.append({
            "name": f"Tab{tab_id + 1}",
            "title": ts.title,
            "x_label": ts.x_label,
            "auto_title": ts.auto_title,
            "axes": {
                "primary": {"label": ts.left_label, "min": ts.left_min, "max": ts.left_max, "step": ts.left_step},
                "secondary": {"label": ts.right_label, "min": ts.right_min, "max": ts.right_max, "step": ts.right_step},
            },
            "signals": signal_entries,
            "reference_lines": [dict(item) for item in ts.reference_lines],
            "range_bands": [dict(item) for item in ts.range_bands],
        })

    return {
        "format": "SignalWorksStudioPreset",
        "format_version": 1,
        "name": preset_name or state.preset_name or "Signal Preset",
        "created_by": APP_TITLE,
        "app_version": APP_VERSION,
        "description": "Portable signal, axis, color, line-style, reference-line and range-band preset. No CSV data or test-specific comments are stored.",
        "tabs": tabs,
    }


def build_builtin_nissan_preset() -> dict:
    """Convert the legacy V3/V4 Nissan defaults to the portable V4.2.x format."""
    tabs = []
    for tab_id, cfg in enumerate(TAB_PRESETS):
        signals = [
            {"name": sig, "axis": default_axis_for_signal(tab_id, sig), "color": None}
            for sig in cfg.signals
        ]
        tabs.append({
            "name": cfg.name,
            "title": cfg.title,
            "x_label": "Time",
            "auto_title": False,
            "axes": {
                "primary": {"label": cfg.left_label, "min": _fmt_default(cfg.left_min), "max": _fmt_default(cfg.left_max), "step": _fmt_default(cfg.left_step)},
                "secondary": {"label": cfg.right_label or "", "min": _fmt_default(cfg.right_min), "max": _fmt_default(cfg.right_max), "step": _fmt_default(cfg.right_step)},
            },
            "signals": signals,
        })
    return {
        "format": "SignalWorksStudioPreset",
        "format_version": 1,
        "name": "Nissan OBC Default",
        "created_by": APP_TITLE,
        "app_version": APP_VERSION,
        "description": "Legacy Nissan OBC V3.5.2/V4.x default preset converted to portable format.",
        "tabs": tabs,
    }


HEADER_MARKER_NORMALIZED = {"timestamp", "datetime", "date", "time"}


def _normalize_header_marker(value: object) -> str:
    """Normalize header text so Date&Time, Date/Time, Date Time, etc. become datetime."""
    text = str(value or "").lstrip("\ufeff").strip().lower()
    return re.sub(r"[^a-z0-9]+", "", text)


def _decode_csv_text(file_bytes: bytes) -> Tuple[str, str]:
    last_error: Optional[Exception] = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return file_bytes.decode(encoding), encoding
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError(f"Cannot decode CSV file: {last_error}")


def detect_csv_header_row(file_bytes: bytes, max_scan_rows: int = 1000) -> Tuple[int, str, str, str, str]:
    """
    Detect a CSV header row by looking only at the FIRST column for one of:
    Timestamp, Date&Time (normalized as datetime), Date, or Time.

    Returns: (zero_based_row, encoding, original_marker, delimiter, detection_text)
    If no marker is found, row 0 is used as a safe ordinary-CSV fallback.
    """
    if not file_bytes:
        raise ValueError("Uploaded CSV file is empty.")

    text, encoding = _decode_csv_text(file_bytes)
    lines = text.splitlines()
    delimiters = (",", ";", "\t", "|")

    for row_idx, raw_line in enumerate(lines[:max_scan_rows]):
        if not raw_line.strip():
            continue
        for delimiter in delimiters:
            first_cell = raw_line.split(delimiter, 1)[0].strip().strip('"').strip("'")
            normalized = _normalize_header_marker(first_cell)
            if normalized in HEADER_MARKER_NORMALIZED:
                display = first_cell.strip() or normalized
                return row_idx, encoding, display, delimiter, f"auto-detected '{display}'"

    # Ordinary CSV fallback: first row is treated as the header. This keeps generic
    # projects usable while eliminating the need to count metadata rows manually
    # whenever a recognized timestamp/date/time marker is present.
    return 0, encoding, "", ",", "marker not found; first row used as header"


def _best_datetime_conversion(df: pd.DataFrame, timestamp_col_name: str) -> Tuple[Optional[pd.Series], int, str]:
    """Try common logger timestamp layouts, including separate Date + Time columns."""
    if timestamp_col_name not in df.columns:
        return None, 0, ""

    first_norm = _normalize_header_marker(timestamp_col_name)
    source = df[timestamp_col_name]
    source_description = timestamp_col_name

    # If the first header is Date and the second column is Time, combine them.
    if first_norm == "date" and len(df.columns) > 1:
        second_name = str(df.columns[1])
        if _normalize_header_marker(second_name) == "time":
            source = df[timestamp_col_name].astype(str).str.strip() + " " + df[second_name].astype(str).str.strip()
            source_description = f"{timestamp_col_name} + {second_name}"

    # Prefer an explicit day-first parser for dotted European logger dates.
    # This avoids pandas/locale inference treating e.g. 01.09.2026 as Jan 9.
    source_text = source.astype(str).str.strip()
    dotted_mask = source_text.str.match(
        r"^\d{1,2}\.\d{1,2}\.\d{4}[ _]\d{1,2}:\d{2}:\d{2}(?:\.\d+)?$",
        na=False,
    )
    # Explicit formats work on pandas 1.5 too (format="mixed" requires 2.x).
    # Combine whole/fractional seconds, and never let fallback inference replace
    # a dotted date, even when an invalid row is present in the same file.
    dotted_source = source_text.where(dotted_mask).str.replace("_", " ", regex=False)
    dotted_conv = pd.to_datetime(dotted_source, format="%d.%m.%Y %H:%M:%S", errors="coerce")
    dotted_conv = dotted_conv.fillna(pd.to_datetime(
        dotted_source, format="%d.%m.%Y %H:%M:%S.%f", errors="coerce"))
    best_conv = dotted_conv if bool(dotted_mask.any()) else None
    best_count = int(dotted_conv.notna().sum())

    formats = (
        "%Y-%m-%d_%H-%M-%S",
        "%Y-%m-%d_%H:%M:%S",
        "%d/%m/%Y_%H:%M:%S",
        "%d/%m/%Y %H:%M:%S",
        "%d.%m.%Y %H:%M:%S",
        "%d.%m.%Y_%H:%M:%S",
        "%d.%m.%Y %H:%M:%S.%f",
        "%d.%m.%Y_%H:%M:%S.%f",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M:%S.%f",
        "%Y/%m/%d %H:%M:%S",
        "%m/%d/%Y %H:%M:%S",
        None,
    )
    target_count = int(source.notna().sum())
    for fmt in formats:
        try:
            # If the explicit dotted day-first parser already covered every non-null
            # sample, do not fall through to pandas locale inference.
            if fmt is None and best_count >= target_count:
                continue
            if fmt is None:
                conv = pd.to_datetime(source.where(~dotted_mask), errors="coerce")
            else:
                conv = pd.to_datetime(source, format=fmt, errors="coerce")
            if bool(dotted_mask.any()):
                conv = conv.copy()
                conv.loc[dotted_mask] = dotted_conv.loc[dotted_mask]
        except Exception:
            continue
        cnt = int(conv.notna().sum())
        if cnt > best_count:
            best_count = cnt
            best_conv = conv
    return best_conv, best_count, source_description


def parse_csv_bytes(file_bytes: bytes) -> Tuple[pd.DataFrame, str, str, List[str], int, int, str, str, str]:
    """Load CSV using V4.3 Smart Header Detection."""
    if not file_bytes:
        raise ValueError("Uploaded CSV file is empty.")

    header_row, encoding, header_marker, delimiter, detection_text = detect_csv_header_row(file_bytes)

    try:
        df = pd.read_csv(
            io.BytesIO(file_bytes),
            skiprows=header_row,
            header=0,
            encoding=encoding,
            sep=delimiter,
        )
    except Exception as exc:
        # If fallback row 0 used the default comma but the file is another delimiter,
        # let pandas' Python engine infer it once before failing.
        if header_row == 0 and not header_marker:
            try:
                df = pd.read_csv(io.BytesIO(file_bytes), header=0, encoding=encoding, sep=None, engine="python")
                delimiter = "auto"
            except Exception:
                raise ValueError(f"Cannot read CSV file: {exc}") from exc
        else:
            raise ValueError(f"Cannot read CSV file: {exc}") from exc

    # A generic non-comma CSV can parse successfully as a single giant column.
    # When no smart marker was found, retry delimiter inference before continuing.
    if header_row == 0 and not header_marker and len(df.columns) == 1:
        try:
            inferred = pd.read_csv(io.BytesIO(file_bytes), header=0, encoding=encoding, sep=None, engine="python")
            if len(inferred.columns) > 1:
                df = inferred
                delimiter = "auto"
        except Exception:
            pass

    if df.empty:
        raise ValueError("CSV file contains no data rows after the detected header.")
    if len(df.columns) == 0:
        raise ValueError("CSV has no columns.")

    # Strip BOM/whitespace from column names without changing project-specific text.
    df.columns = [str(c).lstrip("\ufeff").strip() for c in df.columns]
    timestamp_col_name = str(df.columns[0])
    total_rows = len(df)

    if header_marker:
        best_conv, best_count, source_description = _best_datetime_conversion(df, timestamp_col_name)
    else:
        best_conv, best_count, source_description = None, 0, ""

    if best_conv is not None and best_count > 0.5 * total_rows:
        df.index = pd.DatetimeIndex(best_conv)
        parsed_info = f"timestamp parsed ({source_description})"
    else:
        if header_marker:
            parsed_info = "index mode (timestamp parsing failed)"
        else:
            parsed_info = "index mode (no timestamp/date/time header marker)"
        timestamp_col_name = "Index (Timestamp Parsing Failed)"

    # Convert numeric-looking columns before collecting plot signals. This handles
    # loggers that quote numeric values or mix a small number of blank strings.
    for col in df.columns:
        if col == str(df.columns[0]):
            continue
        if not pd.api.types.is_numeric_dtype(df[col]):
            converted = pd.to_numeric(df[col], errors="coerce")
            if int(converted.notna().sum()) > 0:
                df[col] = converted

    numeric_cols = df.select_dtypes(include=["number"]).columns.tolist()
    if not numeric_cols:
        raise ValueError("No numeric columns were found in the CSV file.")

    return (
        df,
        timestamp_col_name,
        parsed_info,
        numeric_cols,
        total_rows,
        header_row,
        detection_text,
        header_marker,
        delimiter,
    )


# ------------------------------------------------------------------------------
# V4.4.5 Multi-file CSV binding helpers
# ------------------------------------------------------------------------------
def _bind_csv_delimiter(text: str, header_row: int, detected_delimiter: str, header_marker: str) -> str:
    """Return a concrete delimiter for binding, including ordinary CSV fallback."""
    if detected_delimiter in (",", ";", "\t", "|") and header_marker:
        return detected_delimiter

    lines = text.splitlines()
    sample = "\n".join(lines[header_row: min(len(lines), header_row + 12)])
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        if dialect.delimiter in (",", ";", "\t", "|"):
            return dialect.delimiter
    except Exception:
        pass

    if detected_delimiter in (",", ";", "\t", "|"):
        return detected_delimiter
    return ","


def _bind_header_key(values: List[str]) -> List[str]:
    return [str(v).lstrip("\ufeff").strip() for v in values]


def _bound_output_filename(source_names: List[str]) -> str:
    if not source_names:
        return "Bound_Data.csv"
    first = os.path.splitext(os.path.basename(source_names[0]))[0]
    safe = re.sub(r"[^A-Za-z0-9_-]+", "_", first).strip("_") or "Data"
    if len(source_names) == 1:
        return f"{safe}.csv"
    return f"Bound_{safe}_{len(source_names)}files.csv"


def bind_csv_payloads(file_payloads: List[Tuple[str, bytes]]) -> Tuple[bytes, dict]:
    """
    Bind ordered CSV files without repeating metadata/header sections.

    The first file contributes:
      metadata rows (if present) + header + data

    Every subsequent file contributes:
      data rows only

    Column names must match the first file in the same order. Different delimiters
    are normalized to the first file's delimiter in the bound output.
    """
    if not file_payloads:
        raise ValueError("No CSV files were supplied for binding.")
    if len(file_payloads) < 2:
        raise ValueError("Bind Data requires at least two CSV files.")

    decoded = []
    total_input_bytes = 0
    for name, payload in file_payloads:
        if not payload:
            raise ValueError(f"'{name}' is empty.")
        total_input_bytes += len(payload)
        header_row, encoding, marker, detected_delim, detection_text = detect_csv_header_row(payload)
        csv_text, _ = _decode_csv_text(payload)
        delimiter = _bind_csv_delimiter(csv_text, header_row, detected_delim, marker)
        decoded.append({
            "name": os.path.basename(name or "data.csv"),
            "text": csv_text,
            "header_row": header_row,
            "marker": marker,
            "delimiter": delimiter,
            "detection": detection_text,
        })

    output = io.StringIO(newline="")
    first_info = decoded[0]
    first_lines = first_info["text"].splitlines()

    # Keep metadata/preamble from the first file so Smart Header Detection behaves
    # exactly like a normal original logger CSV after binding.
    if first_info["header_row"] > 0:
        for line in first_lines[:first_info["header_row"]]:
            output.write(line)
            output.write("\n")

    writer = csv.writer(
        output,
        delimiter=first_info["delimiter"],
        quotechar='"',
        quoting=csv.QUOTE_MINIMAL,
        lineterminator="\n",
    )

    canonical_header: Optional[List[str]] = None
    source_rows: List[int] = []
    source_details: List[dict] = []

    for file_index, info in enumerate(decoded):
        lines = info["text"].splitlines()
        csv_segment = "\n".join(lines[info["header_row"]:])
        reader = csv.reader(io.StringIO(csv_segment), delimiter=info["delimiter"])
        try:
            raw_header = next(reader)
        except StopIteration:
            raise ValueError(f"'{info['name']}' does not contain a CSV header.")

        normalized_header = _bind_header_key(raw_header)
        if not normalized_header or not any(normalized_header):
            raise ValueError(f"'{info['name']}' contains an empty CSV header.")

        if canonical_header is None:
            canonical_header = normalized_header
            writer.writerow(raw_header)
        elif normalized_header != canonical_header:
            missing = [c for c in canonical_header if c not in normalized_header][:6]
            extra = [c for c in normalized_header if c not in canonical_header][:6]
            detail_parts = []
            if missing:
                detail_parts.append("missing: " + ", ".join(missing))
            if extra:
                detail_parts.append("extra: " + ", ".join(extra))
            detail = "; ".join(detail_parts) or "column order/name mismatch"
            raise ValueError(
                f"Cannot bind '{info['name']}': CSV header does not match the first file ({detail})."
            )

        row_count = 0
        for row in reader:
            # Ignore completely empty separator rows, but preserve rows containing
            # at least one value.
            if not row or not any(str(cell).strip() for cell in row):
                continue
            if len(row) != len(canonical_header):
                raise ValueError(
                    f"Cannot bind '{info['name']}': data row {row_count + 1} has "
                    f"{len(row)} columns; expected {len(canonical_header)}."
                )
            writer.writerow(row)
            row_count += 1

        source_rows.append(row_count)
        source_details.append({
            "name": info["name"],
            "rows": row_count,
            "header_row": int(info["header_row"]) + 1,
            "delimiter": info["delimiter"],
            "header_detection": info["detection"],
        })

    bound_text = output.getvalue()
    # UTF-8 BOM improves compatibility when users open the downloaded bound file
    # directly in Excel, while parse_csv_bytes already supports utf-8-sig.
    bound_bytes = bound_text.encode("utf-8-sig")
    names = [item["name"] for item in decoded]
    return bound_bytes, {
        "filename": _bound_output_filename(names),
        "source_files": names,
        "source_rows": source_rows,
        "source_details": source_details,
        "total_rows": int(sum(source_rows)),
        "input_bytes": int(total_input_bytes),
        "output_bytes": int(len(bound_bytes)),
        "columns": len(canonical_header or []),
    }


def load_csv_payload_into_state(
    state: BrowserSession,
    file_bytes: bytes,
    filename: str,
    source_files: Optional[List[str]] = None,
) -> None:
    """Load one ordinary or already-bound CSV into the normal SignalWorks Studio workflow."""
    (
        df, timestamp_col_name, parsed_info, numeric_cols, total_rows,
        header_row_index, header_detection, header_marker, delimiter
    ) = parse_csv_bytes(file_bytes)

    state.df = df
    state.filename = os.path.basename(filename or "uploaded.csv")
    state.numeric_cols = [str(c) for c in numeric_cols]
    state.timestamp_col_name = timestamp_col_name
    state.parsed_info = parsed_info
    state.row_count = total_rows
    state.header_row_index = header_row_index
    state.header_detection = header_detection
    state.header_marker = header_marker
    state.delimiter = delimiter
    state.loaded_source_files = list(source_files or [state.filename])
    state.dataset_token = uuid.uuid4().hex[:16]
    state.plot_cache.clear()

    if state.loaded_preset:
        matched, requested = apply_loaded_preset(state)
        source_note = (
            f" · bound from {len(state.loaded_source_files)} files"
            if len(state.loaded_source_files) > 1 else ""
        )
        state.status = (
            f"Loaded {state.filename}{source_note} · header row {header_row_index + 1} "
            f"{header_detection} · preset '{state.preset_name}' matched "
            f"{matched}/{requested} configured signals"
        )
    else:
        initialize_tabs_normal(state)
        source_note = (
            f" · bound from {len(state.loaded_source_files)} files"
            if len(state.loaded_source_files) > 1 else ""
        )
        state.status = (
            f"Loaded {state.filename}{source_note} · header row {header_row_index + 1} "
            f"{header_detection} · Normal Plot Mode"
        )
    state.status_type = "success"


# ------------------------------------------------------------------------------
# V4.3 Cycle analysis helpers (derived from graphplot_common.py V3.1.1 logic)
# ------------------------------------------------------------------------------
def _cycle_duration_between(state: BrowserSession, start_pos: int, end_pos: int, fallback_sample_time: float) -> float:
    if state.df is None or end_pos <= start_pos:
        return 0.0
    if pd.api.types.is_datetime64_any_dtype(state.df.index):
        try:
            a = pd.Timestamp(state.df.index[start_pos])
            b = pd.Timestamp(state.df.index[end_pos])
            if not pd.isna(a) and not pd.isna(b):
                seconds = float((b - a).total_seconds())
                if np.isfinite(seconds) and seconds >= 0:
                    return seconds
        except Exception:
            pass
    return float(end_pos - start_pos) * float(fallback_sample_time)


def _cycle_format_duration(value: object) -> str:
    try:
        seconds = float(value)
    except Exception:
        return "-"
    if not np.isfinite(seconds):
        return "-"
    seconds = max(0.0, seconds)
    if seconds < 60:
        return f"{seconds:.1f} s"
    whole = int(round(seconds))
    hours, rem = divmod(whole, 3600)
    minutes, secs = divmod(rem, 60)
    if hours > 0:
        return f"{hours:02d}:{minutes:02d}:{secs:02d}"
    return f"{minutes:02d}:{secs:02d}"


def _cycle_format_position(state: BrowserSession, pos: Optional[int]) -> str:
    if state.df is None or pos is None:
        return "-"
    try:
        if pd.api.types.is_datetime64_any_dtype(state.df.index):
            ts = pd.Timestamp(state.df.index[pos])
            if not pd.isna(ts):
                return ts.strftime("%d/%m/%Y %H:%M:%S")
        return str(pos)
    except Exception:
        return str(pos)


def _extract_cycle_band_runs(state: BrowserSession, values: np.ndarray, low: float, high: float, label: str, fallback_sample_time: float) -> List[dict]:
    finite = np.isfinite(values)
    inside = finite & (values >= low) & (values <= high)
    runs: List[dict] = []
    n = len(values)
    i = 0
    while i < n:
        if not inside[i]:
            i += 1
            continue
        start = i
        while i + 1 < n and inside[i + 1]:
            i += 1
        last_inside = i
        exit_pos = last_inside + 1 if last_inside + 1 < n else last_inside
        closed = last_inside + 1 < n
        duration = _cycle_duration_between(state, start, exit_pos, fallback_sample_time)
        runs.append({
            "state": label,
            "start_pos": start,
            "last_inside_pos": last_inside,
            "exit_pos": exit_pos,
            "closed": closed,
            "duration": duration,
        })
        i += 1
    return runs


def _compress_cycle_plateau_events(events: List[dict]) -> List[dict]:
    if not events:
        return []
    events = sorted(events, key=lambda e: e["start_pos"])
    compressed: List[dict] = []
    group = [events[0]]
    for event in events[1:]:
        if event["state"] == group[-1]["state"]:
            group.append(event)
        else:
            compressed.append(max(group, key=lambda e: e["duration"]))
            group = [event]
    compressed.append(max(group, key=lambda e: e["duration"]))
    return compressed


def _estimate_cycle_plateau_centers(values: np.ndarray, tolerance: float) -> Tuple[float, float]:
    """Estimate dominant low/high plateau temperatures without using overshoot extrema.

    Thermal-shock samples can overshoot the real soak plateau by several degrees.
    A histogram-density estimate is intentionally used here because plateau samples
    occur repeatedly at nearly the same temperature while transition samples are
    spread over a much wider temperature range.
    """
    numeric = np.asarray(values, dtype=float)
    finite = numeric[np.isfinite(numeric)]
    if finite.size < 10:
        raise ValueError("Not enough numeric samples to estimate cycle plateaus.")

    observed_min = float(np.nanmin(finite))
    observed_max = float(np.nanmax(finite))
    span = observed_max - observed_min
    if not np.isfinite(span) or span <= 0:
        raise ValueError("Selected signal does not contain a usable temperature span.")

    # Use finer bins than the user tolerance, but cap the number of bins for
    # unusually large ranges.
    bin_width = max(0.10, min(1.00, float(tolerance) / 4.0))
    if span / bin_width > 2500:
        bin_width = span / 2500.0

    edge_start = math.floor(observed_min / bin_width) * bin_width
    edge_stop = math.ceil(observed_max / bin_width) * bin_width + bin_width
    edges = np.arange(edge_start, edge_stop + bin_width * 0.25, bin_width)
    if edges.size < 4:
        raise ValueError("Temperature span is too small for plateau estimation.")

    counts, edges = np.histogram(finite, bins=edges)
    centers = (edges[:-1] + edges[1:]) / 2.0

    q10, q90 = np.quantile(finite, [0.10, 0.90])
    split = float((q10 + q90) / 2.0)
    lower_idx = np.flatnonzero(centers < split)
    upper_idx = np.flatnonzero(centers > split)
    if lower_idx.size == 0 or upper_idx.size == 0:
        raise ValueError("Unable to separate low and high temperature plateaus.")

    lower_bin = int(lower_idx[np.argmax(counts[lower_idx])])
    upper_bin = int(upper_idx[np.argmax(counts[upper_idx])])
    lower_seed = float(centers[lower_bin])
    upper_seed = float(centers[upper_bin])

    # Refine each histogram peak using the median around that dense region. This
    # suppresses one-sample overshoot/undershoot values while preserving the
    # actual repeated soak temperature.
    refine_radius = max(bin_width, min(max(float(tolerance) * 0.50, bin_width), 2.0))
    lower_points = finite[np.abs(finite - lower_seed) <= refine_radius]
    upper_points = finite[np.abs(finite - upper_seed) <= refine_radius]
    lower = float(np.median(lower_points)) if lower_points.size else lower_seed
    upper = float(np.median(upper_points)) if upper_points.size else upper_seed

    if lower >= upper:
        raise ValueError("Unable to estimate distinct Tmin/Tmax plateaus.")
    return lower, upper


def _median_cycle_transition_seconds(
    state: BrowserSession,
    values: np.ndarray,
    tmin_center: float,
    tmax_center: float,
    tolerance: float,
    fallback_sample_time: float,
) -> Tuple[Optional[float], int]:
    """Measure typical time between leaving one plateau band and reaching the other."""
    min_runs = _extract_cycle_band_runs(
        state, values, tmin_center - tolerance, tmin_center + tolerance, "Tmin", fallback_sample_time
    )
    max_runs = _extract_cycle_band_runs(
        state, values, tmax_center - tolerance, tmax_center + tolerance, "Tmax", fallback_sample_time
    )
    events = _compress_cycle_plateau_events(min_runs + max_runs)
    transitions: List[float] = []
    for a, b in zip(events, events[1:]):
        if a.get("state") == b.get("state"):
            continue
        if int(b.get("start_pos", 0)) < int(a.get("exit_pos", 0)):
            continue
        seconds = _cycle_duration_between(
            state, int(a["exit_pos"]), int(b["start_pos"]), fallback_sample_time
        )
        if np.isfinite(seconds) and seconds >= 0:
            transitions.append(float(seconds))
    if not transitions:
        return None, 0
    return float(np.median(transitions)), len(transitions)


def _cycle_profile_display(mode: str) -> str:
    return {
        "auto": "Auto",
        "temp_cycling": "Temp cycling",
        "thermal_shock": "Thermal shock",
    }.get(str(mode), str(mode))


def analyze_cycle_signal(
    state: BrowserSession,
    header: str,
    tolerance: float = 2.0,
    fallback_sample_time: float = 10.0,
    profile_mode: str = "auto",
) -> dict:
    if state.df is None:
        raise ValueError("No CSV loaded.")
    if header not in state.df.columns:
        raise ValueError(f"Signal '{header}' was not found in the CSV.")
    if tolerance <= 0:
        raise ValueError("Tolerance must be greater than 0.")
    if fallback_sample_time <= 0:
        raise ValueError("Fallback sample time must be greater than 0.")

    profile_mode = str(profile_mode or "auto").strip().lower()
    if profile_mode not in {"auto", "temp_cycling", "thermal_shock"}:
        raise ValueError("Profile mode must be Auto, Temp cycling, or Thermal shock.")

    numeric = pd.to_numeric(state.df[header], errors="coerce").to_numpy(dtype=float)
    finite = numeric[np.isfinite(numeric)]
    if finite.size < 3:
        raise ValueError("Selected signal does not contain enough numeric samples for cycle analysis.")

    observed_tmin = float(np.nanmin(finite))
    observed_tmax = float(np.nanmax(finite))
    datetime_mode = pd.api.types.is_datetime64_any_dtype(state.df.index)
    time_source = "CSV timestamp" if datetime_mode else f"fallback {fallback_sample_time:g} sec/sample"

    plateau_tmin = observed_tmin
    plateau_tmax = observed_tmax
    plateau_detection_error = ""
    try:
        plateau_tmin, plateau_tmax = _estimate_cycle_plateau_centers(numeric, tolerance)
    except Exception as exc:
        plateau_detection_error = str(exc)

    median_transition_seconds: Optional[float] = None
    transition_count = 0
    transition_rate_c_per_min: Optional[float] = None
    if not plateau_detection_error and plateau_tmax > plateau_tmin:
        median_transition_seconds, transition_count = _median_cycle_transition_seconds(
            state,
            numeric,
            plateau_tmin,
            plateau_tmax,
            tolerance,
            fallback_sample_time,
        )
        if median_transition_seconds is not None:
            transition_minutes = max(
                median_transition_seconds / 60.0,
                max(float(fallback_sample_time), 0.001) / 60.0,
            )
            transition_rate_c_per_min = float((plateau_tmax - plateau_tmin) / transition_minutes)

    detected_profile = "temp_cycling"
    if median_transition_seconds is not None:
        transition_minutes = median_transition_seconds / 60.0
        # Uploaded validation profiles separate very clearly:
        # thermal shock ~1-3 min transition, temperature cycling ~45-75 min.
        # Keep both a time and a rate criterion so the rule scales with span.
        if transition_minutes <= 15.0 or (
            transition_rate_c_per_min is not None and transition_rate_c_per_min >= 8.0
        ):
            detected_profile = "thermal_shock"

    profile_used = detected_profile if profile_mode == "auto" else profile_mode
    if profile_used == "thermal_shock" and not plateau_detection_error:
        tmin = plateau_tmin
        tmax = plateau_tmax
    else:
        # Preserve V4.3/V4.4 legacy behavior for temperature cycling so existing
        # qualification results do not unexpectedly change.
        tmin = observed_tmin
        tmax = observed_tmax
        if profile_used == "thermal_shock" and plateau_detection_error:
            profile_used = "temp_cycling"

    if (tmax - tmin) <= 2.0 * tolerance:
        return {
            "signal": header,
            "tmin": tmin,
            "tmax": tmax,
            "observed_tmin": observed_tmin,
            "observed_tmax": observed_tmax,
            "plateau_tmin": plateau_tmin,
            "plateau_tmax": plateau_tmax,
            "tolerance": tolerance,
            "profile_mode_requested": profile_mode,
            "profile_mode_used": profile_used,
            "profile_mode_display": _cycle_profile_display(profile_used),
            "detected_profile": detected_profile,
            "detected_profile_display": _cycle_profile_display(detected_profile),
            "median_transition_seconds": median_transition_seconds,
            "median_transition": _cycle_format_duration(median_transition_seconds),
            "transition_count": transition_count,
            "transition_rate_c_per_min": transition_rate_c_per_min,
            "cycles": 0, "tmin_soak_count": 0, "tmax_soak_count": 0,
            "avg_soak_tmin_seconds": None, "avg_soak_tmax_seconds": None,
            "avg_cycle_time_seconds": None, "avg_soak_tmin": "-", "avg_soak_tmax": "-",
            "avg_cycle_time": "-", "time_source": time_source, "datetime_mode": datetime_mode,
            "note": "Tmin/Tmax tolerance bands overlap; a cycle cannot be separated reliably.",
            "details": [],
        }

    min_runs = _extract_cycle_band_runs(state, numeric, tmin - tolerance, tmin + tolerance, "Tmin", fallback_sample_time)
    max_runs = _extract_cycle_band_runs(state, numeric, tmax - tolerance, tmax + tolerance, "Tmax", fallback_sample_time)
    events = _compress_cycle_plateau_events(min_runs + max_runs)

    details: List[dict] = []
    cycle_no = 0
    i = 0
    while i + 2 < len(events):
        e0, e1, e2 = events[i], events[i + 1], events[i + 2]
        if e0["state"] == e2["state"] and e0["state"] != e1["state"]:
            cycle_no += 1
            cycle_time = _cycle_duration_between(state, e0["start_pos"], e2["start_pos"], fallback_sample_time)
            if e0["state"] == "Tmin":
                soak_tmin, soak_tmax = e0["duration"], e1["duration"]
            else:
                soak_tmax, soak_tmin = e0["duration"], e1["duration"]
            details.append({
                "cycle": cycle_no,
                "start_state": e0["state"],
                "cycle_start_pos": e0["start_pos"],
                "cycle_end_pos": e2["start_pos"],
                "cycle_start": _cycle_format_position(state, e0["start_pos"]),
                "cycle_end": _cycle_format_position(state, e2["start_pos"]),
                "soak_tmin_seconds": soak_tmin,
                "soak_tmax_seconds": soak_tmax,
                "cycle_time_seconds": cycle_time,
                "soak_tmin": _cycle_format_duration(soak_tmin),
                "soak_tmax": _cycle_format_duration(soak_tmax),
                "cycle_time": _cycle_format_duration(cycle_time),
            })
            i += 2
        else:
            i += 1

    tmin_soaks = [e["duration"] for e in events if e["state"] == "Tmin" and e["closed"] and e["duration"] >= 0]
    tmax_soaks = [e["duration"] for e in events if e["state"] == "Tmax" and e["closed"] and e["duration"] >= 0]
    cycle_times = [r["cycle_time_seconds"] for r in details]

    avg_tmin = float(np.mean(tmin_soaks)) if tmin_soaks else None
    avg_tmax = float(np.mean(tmax_soaks)) if tmax_soaks else None
    avg_cycle = float(np.mean(cycle_times)) if cycle_times else None

    analysis_note = ""
    if profile_mode == "auto":
        analysis_note = (
            f"Auto detected {_cycle_profile_display(detected_profile)}. "
            + (
                "Dominant plateau temperatures are used to reject thermal-shock overshoot."
                if profile_used == "thermal_shock"
                else "Legacy observed extrema are retained for slow-ramp temperature cycling."
            )
        )

    return {
        "signal": header,
        "tmin": tmin,
        "tmax": tmax,
        "observed_tmin": observed_tmin,
        "observed_tmax": observed_tmax,
        "plateau_tmin": plateau_tmin,
        "plateau_tmax": plateau_tmax,
        "tolerance": tolerance,
        "profile_mode_requested": profile_mode,
        "profile_mode_used": profile_used,
        "profile_mode_display": _cycle_profile_display(profile_used),
        "detected_profile": detected_profile,
        "detected_profile_display": _cycle_profile_display(detected_profile),
        "median_transition_seconds": median_transition_seconds,
        "median_transition": _cycle_format_duration(median_transition_seconds),
        "transition_count": transition_count,
        "transition_rate_c_per_min": transition_rate_c_per_min,
        "plateau_detection_error": plateau_detection_error,
        "cycles": cycle_no,
        "tmin_soak_count": len(tmin_soaks),
        "tmax_soak_count": len(tmax_soaks),
        "avg_soak_tmin_seconds": avg_tmin,
        "avg_soak_tmax_seconds": avg_tmax,
        "avg_cycle_time_seconds": avg_cycle,
        "avg_soak_tmin": _cycle_format_duration(avg_tmin),
        "avg_soak_tmax": _cycle_format_duration(avg_tmax),
        "avg_cycle_time": _cycle_format_duration(avg_cycle),
        "time_source": time_source,
        "datetime_mode": datetime_mode,
        "note": analysis_note if cycle_no else (
            "No complete A → B → A cycle pattern detected."
            + (f" Plateau detection: {plateau_detection_error}" if plateau_detection_error else "")
        ),
        "details": details,
    }


# ------------------------------------------------------------------------------
# Plot helpers
# ------------------------------------------------------------------------------
def _safe_float(text: str, field_name: str, allow_empty: bool = True) -> Optional[float]:
    value = (text or "").strip()
    if not value:
        if allow_empty:
            return None
        raise ValueError(f"{field_name} is required.")
    try:
        return float(value)
    except ValueError:
        raise ValueError(f"{field_name} must be numeric.")


def apply_axis_visual_style(ax_left, ax_right=None) -> None:
    if ax_left is not None:
        ax_left.yaxis.label.set_color(PRIMARY_AXIS_COLOR)
        ax_left.tick_params(axis="y", colors=PRIMARY_AXIS_COLOR)
        ax_left.spines["left"].set_color(PRIMARY_AXIS_COLOR)
        ax_left.spines["left"].set_linewidth(1.8)

    if ax_right is not None:
        ax_right.yaxis.label.set_color(SECONDARY_AXIS_COLOR)
        ax_right.tick_params(axis="y", colors=SECONDARY_AXIS_COLOR)
        ax_right.spines["right"].set_color(SECONDARY_AXIS_COLOR)
        ax_right.spines["right"].set_linewidth(1.8)


def compact_legend_ncol(item_count: int, max_rows: int = 3) -> int:
    # Expand columns horizontally so the legend never needs more than max_rows rows.
    if item_count <= 0:
        return 1
    return max(1, int(math.ceil(item_count / max(1, int(max_rows)))))


def create_separate_axis_legends(ax_left, ax_right=None):
    handles_left, labels_left = ax_left.get_legend_handles_labels()
    handles_right, labels_right = ([], [])
    if ax_right is not None:
        handles_right, labels_right = ax_right.get_legend_handles_labels()

    legend_left = None
    legend_right = None

    if handles_left:
        legend_left = ax_left.legend(
            handles_left,
            labels_left,
            title="PRIMARY",
            loc="upper left",
            bbox_to_anchor=(0.0, -0.14),
            borderaxespad=0.0,
            fontsize="small",
            title_fontsize="small",
            framealpha=0.95,
            ncol=compact_legend_ncol(len(labels_left)),
            handlelength=1.8,
            handletextpad=0.4,
            columnspacing=0.8,
            borderpad=0.35,
            labelspacing=0.30,
        )
        legend_left.get_title().set_fontweight("bold")
        legend_left.get_title().set_color(PRIMARY_AXIS_COLOR)
        legend_left.get_frame().set_edgecolor(PRIMARY_AXIS_COLOR)
        legend_left.get_frame().set_linewidth(1.2)

    if ax_right is not None and handles_right:
        legend_right = ax_right.legend(
            handles_right,
            labels_right,
            title="SECONDARY",
            loc="upper right",
            bbox_to_anchor=(1.0, -0.14),
            borderaxespad=0.0,
            fontsize="small",
            title_fontsize="small",
            framealpha=0.95,
            ncol=compact_legend_ncol(len(labels_right)),
            handlelength=1.8,
            handletextpad=0.4,
            columnspacing=0.8,
            borderpad=0.35,
            labelspacing=0.30,
        )
        legend_right.get_title().set_fontweight("bold")
        legend_right.get_title().set_color(SECONDARY_AXIS_COLOR)
        legend_right.get_frame().set_edgecolor(SECONDARY_AXIS_COLOR)
        legend_right.get_frame().set_linewidth(1.2)

    # Fit both legend boxes within the plot width, including long signal names.
    # The row count stays fixed; reduce text size only when columns need space.
    legends = [item for item in (legend_left, legend_right) if item is not None]
    if legends:
        fig = ax_left.figure
        renderer = fig.canvas.get_renderer()
        for _ in range(4):
            available = ax_left.get_window_extent(renderer).width * 0.96
            occupied = sum(item.get_window_extent(renderer).width for item in legends)
            if occupied <= available:
                break
            scale = min(0.95, available / occupied)
            for item in legends:
                for text in [item.get_title(), *item.get_texts()]:
                    text.set_fontsize(text.get_fontsize() * scale)
    return legend_left, legend_right


def get_plot_items(state: BrowserSession, tab_id: int):
    if state.df is None:
        return []

    ts = state.tabs[tab_id]
    items = []
    for header in ts.selected_signals:
        if header not in state.df.columns:
            continue
        series = pd.to_numeric(state.df[header], errors="coerce").values
        if series.size == 0 or np.isnan(series).all():
            continue
        axis_target = ts.axis_assignments.get(header, default_axis_for_signal(tab_id, header))
        if axis_target not in ("left", "right"):
            axis_target = "left"
        items.append((header, series, axis_target))
    return items


def _mpl_dash_style(name: str):
    return {"solid": "-", "dash": "--", "dot": ":", "dashdot": "-.", "longdash": (0, (8, 4))}.get(name, "--")


def _hex_rgba(hex_color: str, alpha: float):
    h = str(hex_color).lstrip("#")
    try:
        return tuple(int(h[i:i+2], 16) / 255.0 for i in (0, 2, 4)) + (alpha,)
    except Exception:
        return (0.07, 0.72, 0.42, alpha)


def _plotly_rgba(hex_color: str, alpha: float) -> str:
    h = str(hex_color).lstrip("#")
    try:
        r, g, b = (int(h[i:i+2], 16) for i in (0, 2, 4))
        return f"rgba({r},{g},{b},{max(0.0,min(1.0,float(alpha))):.4f})"
    except Exception:
        return f"rgba(18,183,106,{max(0.0,min(1.0,float(alpha))):.4f})"


def _annotation_x_for_mpl(state: BrowserSession, value):
    if state.df is not None and pd.api.types.is_datetime64_any_dtype(state.df.index):
        try:
            return pd.Timestamp(value).to_pydatetime()
        except Exception:
            return value
    try:
        return float(value)
    except Exception:
        return value


def _plotly_comment_text(value: object) -> str:
    """Escape user comment text and preserve explicit line breaks in Plotly annotations."""
    return html.escape(str(value or "")).replace("\r\n", "\n").replace("\r", "\n").replace("\n", "<br>")


def _mpl_comment_text(value: object, box_width_px: int, box_height_px: int, font_size: int) -> str:
    """Preserve manual line breaks and approximately wrap long report comments to the chosen box width."""
    raw = str(value or "").replace("\r\n", "\n").replace("\r", "\n")
    wrapped = []
    if box_width_px:
        # Rough character width estimate; manual line breaks are always preserved.
        chars = max(8, int(box_width_px / max(5.0, float(font_size) * 0.58)))
        for line in raw.split("\n"):
            if not line:
                wrapped.append("")
            else:
                wrapped.extend(textwrap.wrap(line, width=chars, break_long_words=False, break_on_hyphens=False) or [""])
    else:
        wrapped = raw.split("\n")
    if box_height_px:
        # Approximate a requested minimum box height in static report output.
        min_lines = max(1, int(round(box_height_px / max(8.0, float(font_size) * 1.35))))
        while len(wrapped) < min_lines:
            wrapped.append("")
    return "\n".join(wrapped)


def apply_engineering_annotations_mpl(state: BrowserSession, ts: TabState, ax_left, ax_right) -> None:
    for band in ts.range_bands:
        if not band.get("visible", True):
            continue
        ax = ax_right if band.get("axis") == "right" and ax_right is not None else ax_left
        lo, hi = float(band["min"]), float(band["max"])
        fill = _hex_rgba(band.get("fill_color", "#12B76A"), float(band.get("opacity", 0.10)))
        ax.axhspan(lo, hi, facecolor=fill, edgecolor="none", zorder=0.5)
        for y in (lo, hi):
            ax.axhline(y, color=band.get("line_color", "#12B76A"), linewidth=band.get("width", 1.2), linestyle=_mpl_dash_style(band.get("dash", "dash")), alpha=0.9, zorder=1)
        if band.get("show_label", True):
            ax.text(0.992, hi, str(band.get("label", "Range")), transform=ax.get_yaxis_transform(), ha="right", va="bottom", fontsize=8.5, color=band.get("line_color", "#12B76A"), bbox=dict(boxstyle="round,pad=0.2", fc="white", ec=band.get("line_color", "#12B76A"), alpha=0.78), zorder=5)

    for line in ts.reference_lines:
        if not line.get("visible", True):
            continue
        ax = ax_right if line.get("axis") == "right" and ax_right is not None else ax_left
        value = float(line["value"])
        ax.axhline(value, color=line.get("color", "#D92D20"), linewidth=line.get("width", 1.5), linestyle=_mpl_dash_style(line.get("dash", "dash")), alpha=0.95, zorder=4)
        if line.get("show_label", True):
            ax.text(0.992, value, str(line.get("label", "Reference")), transform=ax.get_yaxis_transform(), ha="right", va="bottom", fontsize=8.5, color=line.get("color", "#D92D20"), bbox=dict(boxstyle="round,pad=0.2", fc="white", ec=line.get("color", "#D92D20"), alpha=0.82), zorder=5)

    for comment in ts.comments:
        if not comment.get("visible", True):
            continue
        ax = ax_right if comment.get("axis") == "right" and ax_right is not None else ax_left
        x = _annotation_x_for_mpl(state, comment.get("x")); y = float(comment.get("y"))
        font_size = int(comment.get("font_size", 12))
        border_color = comment.get("border_color", comment.get("text_color", "#101828"))
        border_width = float(comment.get("border_width", 1.0))
        arrowprops = (
            dict(arrowstyle="->", color=comment.get("arrow_color", border_color), lw=float(comment.get("arrow_width", 1.2)))
            if comment.get("show_arrow", True) else None
        )
        display_text = _mpl_comment_text(comment.get("text", ""), int(comment.get("box_width", 0) or 0), int(comment.get("box_height", 0) or 0), font_size)
        ax.annotate(
            display_text, xy=(x, y),
            xytext=(float(comment.get("ax", 32)), -float(comment.get("ay", -38))),
            textcoords="offset points", fontsize=font_size,
            color=comment.get("text_color", "#101828"),
            bbox=dict(boxstyle="round,pad=0.35", fc=comment.get("bg_color", "#FFFFFF"), ec=border_color, lw=border_width, alpha=0.92),
            arrowprops=arrowprops, zorder=8
        )


def build_figure(state: BrowserSession, tab_id: int, report_header: bool = False, figsize=None, view_state: Optional[dict] = None):
    if state.df is None:
        raise ValueError("No CSV loaded.")
    if not (0 <= tab_id < len(TAB_PRESETS)):
        raise ValueError("Invalid tab.")

    ts = state.tabs[tab_id]
    plot_items = get_plot_items(state, tab_id)

    if figsize is None:
        figsize = (11.69, 8.27)
    fig, ax_left = plt.subplots(figsize=figsize, constrained_layout=False)
    fig.subplots_adjust(**DEFAULT_SUBPLOT_PARAMS)

    has_right_axis = (
        any(axis_target == "right" for _, _, axis_target in plot_items)
        or any(x.get("visible", True) and x.get("axis") == "right" for x in ts.reference_lines)
        or any(x.get("visible", True) and x.get("axis") == "right" for x in ts.range_bands)
        or any(x.get("visible", True) and x.get("axis") == "right" for x in ts.comments)
    )
    ax_right = ax_left.twinx() if has_right_axis else None
    if ax_right is not None:
        ax_right.set_zorder(ax_left.get_zorder() + 1)
        ax_right.patch.set_visible(False)

    is_datetime_index = pd.api.types.is_datetime64_any_dtype(state.df.index)
    if is_datetime_index:
        x_vals = state.df.index
        fallback_x_label = "Time"
        ax_left.xaxis.set_major_formatter(mdates.DateFormatter("%d/%m/%Y_%H:%M:%S"))

        # Keep approximately readable ticks for short and long tests.
        if len(x_vals) > 1:
            try:
                duration = x_vals[-1] - x_vals[0]
                hours = duration.total_seconds() / 3600.0
                if hours <= 12:
                    ax_left.xaxis.set_major_locator(mdates.HourLocator(interval=1))
                elif hours <= 48:
                    ax_left.xaxis.set_major_locator(mdates.HourLocator(interval=2))
                elif hours <= 120:
                    ax_left.xaxis.set_major_locator(mdates.HourLocator(interval=6))
                else:
                    ax_left.xaxis.set_major_locator(mdates.DayLocator(interval=1))
            except Exception:
                ax_left.xaxis.set_major_locator(mdates.AutoDateLocator())
        plt.setp(ax_left.get_xticklabels(), rotation=45, ha="right")
    else:
        x_vals = np.arange(len(state.df))
        fallback_x_label = state.timestamp_col_name or "Index"

    primary_index = 0
    secondary_index = 0

    visible_signals = {}
    if isinstance(view_state, dict) and isinstance(view_state.get("visible_signals"), dict):
        visible_signals = {str(k): bool(v) for k, v in view_state.get("visible_signals", {}).items()}

    for header, series, axis_target in plot_items:
        is_visible = visible_signals.get(header, True)
        if axis_target == "right" and ax_right is not None:
            color = _effective_signal_color(ts, header, axis_target, primary_index, secondary_index)
            dash_name = _effective_signal_dash(ts, header, axis_target, secondary_index)
            if is_visible:
                ax_right.plot(
                    x_vals,
                    series,
                    linestyle=_mpl_dash_style(dash_name),
                    linewidth=1.6,
                    color=color,
                    alpha=0.95,
                    label=header,
                    zorder=3,
                )
            secondary_index += 1
        else:
            color = _effective_signal_color(ts, header, axis_target, primary_index, secondary_index)
            dash_name = _effective_signal_dash(ts, header, axis_target, secondary_index)
            if is_visible:
                ax_left.plot(
                    x_vals,
                    series,
                    linestyle=_mpl_dash_style(dash_name),
                    linewidth=1.6,
                    label=header,
                    color=color,
                    alpha=0.95,
                    zorder=2,
                )
            primary_index += 1

    if ts.auto_title and ts.selected_signals:
        title_text = " / ".join(ts.selected_signals)
    else:
        title_text = ts.title.strip() or f"Plot {tab_id + 1}"

    ax_left.set_title(title_text)
    ax_left.set_xlabel(ts.x_label.strip() or fallback_x_label)
    ax_left.set_ylabel(ts.left_label.strip())
    if ax_right is not None:
        ax_right.set_ylabel(ts.right_label.strip())

    apply_axis_visual_style(ax_left, ax_right)

    # Axis limits and tick steps.
    left_min = _safe_float(ts.left_min, "Left Y Min")
    left_max = _safe_float(ts.left_max, "Left Y Max")
    left_step = _safe_float(ts.left_step, "Left Tick Step")

    if left_min is not None or left_max is not None:
        ax_left.set_ylim(left_min, left_max)
    if left_step is not None:
        if left_step <= 0:
            raise ValueError("Left Tick Step must be greater than 0.")
        ax_left.yaxis.set_major_locator(MultipleLocator(left_step))

    if ax_right is not None:
        right_min = _safe_float(ts.right_min, "Right Y Min")
        right_max = _safe_float(ts.right_max, "Right Y Max")
        right_step = _safe_float(ts.right_step, "Right Tick Step")
        if right_min is not None or right_max is not None:
            ax_right.set_ylim(right_min, right_max)
        if right_step is not None:
            if right_step <= 0:
                raise ValueError("Right Tick Step must be greater than 0.")
            ax_right.yaxis.set_major_locator(MultipleLocator(right_step))

    # Apply the current browser view to static exports when supplied.
    if isinstance(view_state, dict):
        try:
            xr = view_state.get("xaxis_range")
            if isinstance(xr, list) and len(xr) == 2:
                if is_datetime_index:
                    ax_left.set_xlim(pd.to_datetime(xr[0]), pd.to_datetime(xr[1]))
                else:
                    ax_left.set_xlim(float(xr[0]), float(xr[1]))
        except Exception:
            pass
        try:
            yr = view_state.get("yaxis_range")
            if isinstance(yr, list) and len(yr) == 2:
                ax_left.set_ylim(float(yr[0]), float(yr[1]))
        except Exception:
            pass
        if ax_right is not None:
            try:
                yr2 = view_state.get("yaxis2_range")
                if isinstance(yr2, list) and len(yr2) == 2:
                    ax_right.set_ylim(float(yr2[0]), float(yr2[1]))
            except Exception:
                pass

    apply_engineering_annotations_mpl(state, ts, ax_left, ax_right)

    if plot_items:
        create_separate_axis_legends(ax_left, ax_right)
    else:
        ax_left.text(
            0.5,
            0.5,
            "No signal selected",
            transform=ax_left.transAxes,
            ha="center",
            va="center",
            fontsize=16,
            color="#777777",
        )

    ax_left.grid(True, axis="x", linestyle="--", alpha=0.5)
    ax_left.grid(True, axis="y", linestyle="--", alpha=0.5)
    if ax_right is not None:
        ax_right.grid(False)

    if report_header:
        add_report_header(fig)

    return fig


def get_base_dir() -> str:
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    try:
        return os.path.dirname(os.path.abspath(__file__))
    except NameError:
        return os.getcwd()


def add_report_header(fig) -> None:
    logo_path = os.path.join(get_base_dir(), "Delta_Logo.png")
    img = None
    if os.path.exists(logo_path):
        try:
            img = mpimg.imread(logo_path)
        except Exception:
            img = None

    pdf_params = DEFAULT_SUBPLOT_PARAMS.copy()
    pdf_params["top"] = 0.80
    fig.subplots_adjust(**pdf_params)

    if img is not None:
        ax_logo = fig.add_axes([0.065, 0.90, 0.12, 0.07], zorder=10)
        ax_logo.imshow(img)
        ax_logo.axis("off")

    fig.text(0.08, 0.87, "Delta Electronics (Thailand) PCL", fontsize=11, color="#666666", ha="left")
    fig.text(0.92, 0.92, "TEST REPORT", fontsize=18, color="#555555", ha="right", fontweight="bold")
    fig.text(0.92, 0.87, "DELTA CONFIDENTIAL", fontsize=11, color="#555555", ha="right")
    line = plt.Line2D(
        (0.08, 0.92),
        (0.855, 0.855),
        color="#cccccc",
        linewidth=1.5,
        transform=fig.transFigure,
        clip_on=False,
    )
    fig.add_artist(line)


def render_plot_png(state: BrowserSession, tab_id: int, dpi: int = 150, use_cache: bool = True, view_state: Optional[dict] = None) -> bytes:
    ts = state.tabs[tab_id]
    cached = state.plot_cache.get(tab_id)
    if view_state is not None:
        use_cache = False
    if use_cache and cached is not None and cached[0] == ts.revision:
        return cached[1]

    fig = build_figure(state, tab_id, view_state=view_state)
    stream = io.BytesIO()
    try:
        fig.savefig(stream, format="png", dpi=dpi, bbox_inches="tight", pad_inches=0.08)
    finally:
        plt.close(fig)
    image = stream.getvalue()

    if use_cache:
        state.plot_cache[tab_id] = (ts.revision, image)
    return image


def _plotly_x_values(state: BrowserSession):
    """Return Plotly-safe x values and whether the axis is datetime based."""
    if state.df is None:
        return [], False

    is_datetime = pd.api.types.is_datetime64_any_dtype(state.df.index)
    if is_datetime:
        return state.df.index.to_pydatetime().tolist(), True
    return np.arange(len(state.df)).tolist(), False


def build_interactive_figure(state: BrowserSession, tab_id: int):
    """Build the interactive Plotly figure, including V4.4 engineering annotations."""
    if state.df is None:
        raise ValueError("No CSV loaded.")
    if not (0 <= tab_id < len(TAB_PRESETS)):
        raise ValueError("Invalid tab.")

    ts = state.tabs[tab_id]
    plot_items = get_plot_items(state, tab_id)
    x_vals, is_datetime = _plotly_x_values(state)

    if ts.auto_title and ts.selected_signals:
        title_text = " / ".join(ts.selected_signals)
    else:
        title_text = ts.title.strip() or f"Plot {tab_id + 1}"

    fig = go.Figure()
    primary_index = 0
    secondary_index = 0
    for header, series, axis_target in plot_items:
        y_vals = [None if pd.isna(v) else float(v) for v in series]

        if axis_target == "right":
            color = _effective_signal_color(ts, header, axis_target, primary_index, secondary_index)
            dash = _effective_signal_dash(ts, header, axis_target, secondary_index)
            fig.add_trace(go.Scattergl(
                x=x_vals,
                y=y_vals,
                mode="lines",
                name=header,
                yaxis="y2",
                legendgroup="secondary",
                legend="legend2",
                line={"color": color, "width": 1.8, "dash": dash},
                hovertemplate="%{x}<br>" + header + ": %{y:.4g}<extra>Secondary</extra>",
            ))
            secondary_index += 1
        else:
            color = _effective_signal_color(ts, header, axis_target, primary_index, secondary_index)
            dash = _effective_signal_dash(ts, header, axis_target, secondary_index)
            fig.add_trace(go.Scattergl(
                x=x_vals,
                y=y_vals,
                mode="lines",
                name=header,
                yaxis="y",
                legendgroup="primary",
                legend="legend",
                line={"color": color, "width": 1.8, "dash": dash},
                hovertemplate="%{x}<br>" + header + ": %{y:.4g}<extra>Primary</extra>",
            ))
            primary_index += 1

    left_min = _safe_float(ts.left_min, "Left Y Min")
    left_max = _safe_float(ts.left_max, "Left Y Max")
    left_step = _safe_float(ts.left_step, "Left Tick Step")
    if left_step is not None and left_step <= 0:
        raise ValueError("Left Tick Step must be greater than 0.")

    has_right_axis = (
        any(axis_target == "right" for _, _, axis_target in plot_items)
        or any(x.get("visible", True) and x.get("axis") == "right" for x in ts.reference_lines)
        or any(x.get("visible", True) and x.get("axis") == "right" for x in ts.range_bands)
        or any(x.get("visible", True) and x.get("axis") == "right" for x in ts.comments)
    )
    right_min = right_max = right_step = None
    if has_right_axis:
        right_min = _safe_float(ts.right_min, "Right Y Min")
        right_max = _safe_float(ts.right_max, "Right Y Max")
        right_step = _safe_float(ts.right_step, "Right Tick Step")
        if right_step is not None and right_step <= 0:
            raise ValueError("Right Tick Step must be greater than 0.")

    yaxis = {
        "title": {"text": ts.left_label.strip(), "font": {"color": PRIMARY_AXIS_COLOR}},
        "tickfont": {"color": PRIMARY_AXIS_COLOR},
        "linecolor": PRIMARY_AXIS_COLOR,
        "linewidth": 2,
        "showline": True,
        "gridcolor": "#dfe4ea",
        "zerolinecolor": "#c8ced6",
        "fixedrange": False,
    }
    if left_min is not None and left_max is not None:
        yaxis["range"] = [left_min, left_max]
    if left_step is not None:
        yaxis["dtick"] = left_step

    yaxis2 = {
        "title": {"text": ts.right_label.strip(), "font": {"color": SECONDARY_AXIS_COLOR}},
        "tickfont": {"color": SECONDARY_AXIS_COLOR},
        "linecolor": SECONDARY_AXIS_COLOR,
        "linewidth": 2,
        "showline": True,
        "overlaying": "y",
        "side": "right",
        "showgrid": False,
        "zeroline": False,
        "fixedrange": False,
        "visible": has_right_axis,
    }
    if right_min is not None and right_max is not None:
        yaxis2["range"] = [right_min, right_max]
    if right_step is not None:
        yaxis2["dtick"] = right_step

    xaxis = {
        "title": {"text": ts.x_label.strip() or ("Time" if is_datetime else state.timestamp_col_name)},
        "showgrid": True,
        "gridcolor": "#e5e9ee",
        "linecolor": "#555f6b",
        "showline": True,
        "fixedrange": False,
        "rangeslider": {"visible": False},
        "showspikes": True,
        "spikemode": "across",
        "spikesnap": "cursor",
        "spikedash": "dot",
        "spikecolor": "#8b95a1",
    }
    if is_datetime:
        xaxis.update({
            "type": "date",
            "tickformat": "%d/%m/%Y<br>%H:%M:%S",
            "hoverformat": "%d/%m/%Y %H:%M:%S",
        })

    engineering_shapes = []
    engineering_annotations = []
    # V4.4.2: browser-side maps allow Plotly drag edits to be synchronized
    # back to the correct persistent server annotation object.
    engineering_shape_map = []
    engineering_annotation_map = []
    for band in ts.range_bands:
        if not band.get("visible", True):
            continue
        yref = "y2" if band.get("axis") == "right" and has_right_axis else "y"
        engineering_shapes.append({
            "type": "rect", "xref": "paper", "yref": yref, "x0": 0, "x1": 1,
            "y0": float(band["min"]), "y1": float(band["max"]),
            "fillcolor": _plotly_rgba(band.get("fill_color", "#12B76A"), float(band.get("opacity", 0.10))),
            "line": {"color": band.get("line_color", "#12B76A"), "width": float(band.get("width", 1.2)), "dash": _clean_dash(band.get("dash"))},
            "layer": "below", "editable": True,
        })
        engineering_shape_map.append({"kind": "band", "id": str(band.get("id", "")), "axis": band.get("axis", "left")})
        if band.get("show_label", True):
            engineering_annotations.append({
                "xref": "paper", "yref": yref, "x": 0.995, "y": float(band["max"]),
                "text": str(band.get("label", "Accept range")), "showarrow": False,
                "xanchor": "right", "yanchor": "bottom",
                "font": {"size": 10, "color": band.get("line_color", "#12B76A")},
                "bgcolor": "rgba(255,255,255,0.82)", "bordercolor": band.get("line_color", "#12B76A"), "borderwidth": 1, "borderpad": 3,
            })
            engineering_annotation_map.append({"kind": "band-label", "id": str(band.get("id", ""))})

    for line in ts.reference_lines:
        if not line.get("visible", True):
            continue
        yref = "y2" if line.get("axis") == "right" and has_right_axis else "y"
        engineering_shapes.append({
            "type": "line", "xref": "paper", "yref": yref, "x0": 0, "x1": 1,
            "y0": float(line["value"]), "y1": float(line["value"]),
            "line": {"color": line.get("color", "#D92D20"), "width": float(line.get("width", 1.5)), "dash": _clean_dash(line.get("dash"))},
            "layer": "above", "editable": True,
        })
        engineering_shape_map.append({"kind": "line", "id": str(line.get("id", "")), "axis": line.get("axis", "left")})
        if line.get("show_label", True):
            engineering_annotations.append({
                "xref": "paper", "yref": yref, "x": 0.995, "y": float(line["value"]),
                "text": str(line.get("label", "Reference")), "showarrow": False,
                "xanchor": "right", "yanchor": "bottom",
                "font": {"size": 10, "color": line.get("color", "#D92D20")},
                "bgcolor": "rgba(255,255,255,0.84)", "bordercolor": line.get("color", "#D92D20"), "borderwidth": 1, "borderpad": 3,
            })
            engineering_annotation_map.append({"kind": "line-label", "id": str(line.get("id", ""))})

    for comment in ts.comments:
        if not comment.get("visible", True):
            continue
        yref = "y2" if comment.get("axis") == "right" and has_right_axis else "y"
        comment_ann = {
            "xref": "x", "yref": yref, "x": comment.get("x"), "y": float(comment["y"]),
            "text": _plotly_comment_text(comment.get("text", "")), "showarrow": bool(comment.get("show_arrow", True)),
            "arrowhead": 2, "arrowsize": 1,
            "arrowwidth": float(comment.get("arrow_width", 1.2)),
            "arrowcolor": comment.get("arrow_color", comment.get("border_color", "#101828")),
            "ax": float(comment.get("ax", 32)), "ay": float(comment.get("ay", -38)),
            "font": {"size": int(comment.get("font_size", 12)), "color": comment.get("text_color", "#101828")},
            "bgcolor": comment.get("bg_color", "#FFFFFF"),
            "bordercolor": comment.get("border_color", comment.get("text_color", "#101828")),
            "borderwidth": float(comment.get("border_width", 1.0)), "borderpad": 5,
            "align": "left", "valign": "middle",
        }
        if int(comment.get("box_width", 0) or 0) > 0:
            comment_ann["width"] = int(comment.get("box_width"))
        if int(comment.get("box_height", 0) or 0) > 0:
            comment_ann["height"] = int(comment.get("box_height"))
        engineering_annotations.append(comment_ann)
        engineering_annotation_map.append({"kind": "comment", "id": str(comment.get("id", ""))})

    # Independent legends preserve both axis headings with normal trace ordering.
    # Stacked full-width boxes avoid squeezing long names into half the plot width.
    dual_legend = primary_index > 0 and secondary_index > 0
    def axis_legend(title, count, color, y):
        return dict(
            title={"text": title, "side": "top", "font": {"color": color}},
            orientation="h", traceorder="normal", yanchor="top", yref="container",
            y=y, xanchor="left", x=0.0, groupclick="toggleitem",
            font={"size": 11}, bgcolor="rgba(255,255,255,0.88)",
            bordercolor=color, borderwidth=1,
            entrywidthmode="fraction",
            # Leave room for legend borders/padding so the last column fits.
            entrywidth=0.98 / compact_legend_ncol(count),
        )

    fig.update_layout(
        title={"text": title_text, "x": 0.5, "xanchor": "center", "font": {"size": 18}},
        xaxis=xaxis,
        yaxis=yaxis,
        yaxis2=yaxis2,
        template="plotly_white",
        autosize=True,
        margin={"l": 72, "r": 82, "t": 62, "b": 240 if dual_legend else 145},
        hovermode="x unified",
        hoverdistance=-1,
        spikedistance=-1,
        dragmode="zoom",
        showlegend=True,
        legend=axis_legend("PRIMARY", primary_index, PRIMARY_AXIS_COLOR, 0.26 if dual_legend else 0.15),
        legend2=axis_legend("SECONDARY", secondary_index, SECONDARY_AXIS_COLOR, 0.13 if dual_legend else 0.15),
        shapes=engineering_shapes,
        annotations=engineering_annotations,
        meta={
            "signalworks_shape_map": engineering_shape_map,
            "signalworks_annotation_map": engineering_annotation_map,
        },
        uirevision=f"tab-{tab_id}-rev-{ts.revision}",
    )

    if not plot_items:
        fig.add_annotation(
            x=0.5, y=0.5, xref="paper", yref="paper",
            text="No signal selected", showarrow=False,
            font={"size": 18, "color": "#777777"},
        )

    return fig


# ------------------------------------------------------------------------------
# HTML helpers
# ------------------------------------------------------------------------------
def esc(value) -> str:
    return html.escape(str(value), quote=True)


def status_class(status_type: str) -> str:
    return {
        "success": "status success",
        "error": "status error",
        "warning": "status warning",
    }.get(status_type, "status info")


def render_tabs(state: BrowserSession, active_tab: int) -> str:
    chunks = []
    for idx, ts in enumerate(state.tabs):
        css = "tab-link active" if idx == active_tab else "tab-link"
        tab_label = ts.title.strip() or f"Plot {idx + 1}"
        chunks.append(f'<a class="{css}" href="/?tab={idx}">{esc(tab_label)}</a>')
    return "".join(chunks)


def _loaded_preset_signal_names(state: BrowserSession, tab_id: int) -> set:
    names = set()
    try:
        tabs = state.loaded_preset.get("tabs", []) if state.loaded_preset else []
        if tab_id >= len(tabs) or not isinstance(tabs[tab_id], dict):
            return names
        for entry in tabs[tab_id].get("signals", []):
            if isinstance(entry, str):
                names.add(entry)
            elif isinstance(entry, dict) and entry.get("name"):
                names.add(str(entry.get("name")))
    except Exception:
        pass
    return names


def render_signal_rows(state: BrowserSession, tab_id: int) -> str:
    ts = state.tabs[tab_id]
    selected = set(ts.selected_signals)
    preset_names = _loaded_preset_signal_names(state, tab_id)
    rows = []

    # Calculate the same auto palette position used by plotting for selected traces.
    selected_positions: Dict[str, Tuple[int, int]] = {}
    p_index = 0
    s_index = 0
    for sig in ts.selected_signals:
        axis = ts.axis_assignments.get(sig, default_axis_for_signal(tab_id, sig))
        selected_positions[sig] = (p_index, s_index)
        if axis == "right":
            s_index += 1
        else:
            p_index += 1

    for idx, signal in enumerate(state.numeric_cols):
        checked = "checked" if signal in selected else ""
        axis_value = ts.axis_assignments.get(signal, default_axis_for_signal(tab_id, signal))
        left_selected = "selected" if axis_value == "left" else ""
        right_selected = "selected" if axis_value == "right" else ""
        row_class = "signal-row primary-row" if axis_value == "left" else "signal-row secondary-row"
        preset_marker = "1" if signal in preset_names else "0"

        pi, si = selected_positions.get(signal, (idx, idx))
        auto_color = _auto_color_for_position(axis_value, pi, si)
        manual_color = ts.color_assignments.get(signal, "")
        is_manual = _valid_hex_color(manual_color)
        shown_color = manual_color if is_manual else auto_color
        color_mode = "manual" if is_manual else "auto"
        auto_class = "" if is_manual else " auto"
        line_style = str(ts.line_style_assignments.get(signal, "auto") or "auto").lower()
        if line_style not in SIGNAL_LINE_STYLES:
            line_style = "auto"
        line_style_options = [("auto", "Auto"), ("solid", "Solid"), ("dash", "Dash"), ("dot", "Dot"), ("dashdot", "Dash-dot"), ("longdash", "Long dash")]
        line_style_html = "".join(
            f'<option value="{value}" {"selected" if line_style == value else ""}>{label}</option>'
            for value, label in line_style_options
        )

        rows.append(
            f"""
            <div class="{row_class}" data-signal-row data-name="{esc(signal).lower()}">
                <label class="signal-check" title="{esc(signal)}">
                    <input type="checkbox" name="selected_{idx}" value="1" {checked}>
                    <span class="signal-name">{esc(signal)}</span>
                    <span class="preset-badge {'show' if preset_marker == '1' else ''}" title="Signal is defined in the active preset">P</span>
                </label>
                <select name="axis_{idx}" class="axis-select" title="Signal axis" aria-label="Axis for {esc(signal)}" onchange="updateSignalRowClass(this); refreshAutoColorForRow(this)">
                    <option value="left" {left_selected}>Primary</option>
                    <option value="right" {right_selected}>Secondary</option>
                </select>
                <select name="line_style_{idx}" class="line-style-select" aria-label="Line style for {esc(signal)}" title="Signal line style. Auto keeps the default style for the selected axis.">{line_style_html}</select>
                <div class="color-control{auto_class}" title="Signal color. Click the color to set manually, or A to return to automatic color.">
                    <input type="hidden" class="color-mode" name="color_mode_{idx}" value="{color_mode}">
                    <input type="color" class="signal-color" name="color_{idx}" value="{esc(shown_color)}" data-auto-primary="{PRIMARY_LINE_COLORS[idx % len(PRIMARY_LINE_COLORS)]}" data-auto-secondary="{SECONDARY_LINE_COLORS[idx % len(SECONDARY_LINE_COLORS)]}" oninput="setManualColor(this)">
                    <button class="auto-color-btn" type="button" onclick="setAutoColor(this)" title="Use automatic palette">A</button>
                </div>
            </div>
            """
        )
    return "".join(rows)


def render_annotation_manager(ts: TabState, tab_id: int) -> str:
    line_items = []
    dash_values = ('solid','dash','dot','dashdot','longdash')
    for item in ts.reference_lines:
        checked = "checked" if item.get("visible", True) else ""
        dash_options = ''.join(f'<option value="{d}" {"selected" if item.get("dash")==d else ""}>{d.title()}</option>' for d in dash_values)
        line_items.append(f'''
        <details class="annotation-item" data-ui-key="ann-item-line-{esc(item.get('id',''))}">
          <summary><span class="ann-kind line-kind">LINE</span><span class="ann-name">{esc(item.get('label','Reference'))}</span><span class="ann-value">{item.get('value')}</span></summary>
          <div class="annotation-edit">
            <form method="post" action="/annotation/line/{tab_id}" data-ann-kind="line" data-ann-id="{esc(item.get('id',''))}">
              <input type="hidden" name="item_id" value="{esc(item.get('id',''))}">
              <div class="ann-grid"><label>Label<input name="label" value="{esc(item.get('label',''))}"></label><label>Value<input name="value" type="number" step="any" value="{item.get('value')}"></label></div>
              <div class="ann-grid"><label>Axis<select name="axis"><option value="left" {'selected' if item.get('axis')!='right' else ''}>Primary</option><option value="right" {'selected' if item.get('axis')=='right' else ''}>Secondary</option></select></label><label>Style<select name="dash">{dash_options}</select></label></div>
              <div class="ann-grid three"><label>Color<input name="color" type="color" value="{esc(item.get('color','#D92D20'))}"></label><label>Width<input name="width" type="number" min="0.5" max="8" step="0.5" value="{item.get('width',1.5)}"></label><label class="check-label"><input name="show_label" type="checkbox" value="1" {'checked' if item.get('show_label',True) else ''}>Label</label></div>
              <div class="ann-actions"><button class="btn compact" type="submit">Save</button><button class="btn compact danger-soft" type="button" onclick="deleteAnnotation('line','{esc(item.get('id',''))}')">Delete</button><label class="visibility-toggle"><input type="checkbox" {checked} onchange="toggleAnnotation('line','{esc(item.get('id',''))}',this.checked)"> Show</label></div>
            </form>
          </div>
        </details>''')
    band_items = []
    for item in ts.range_bands:
        checked = "checked" if item.get("visible", True) else ""
        dash_options = ''.join(f'<option value="{d}" {"selected" if item.get("dash")==d else ""}>{d.title()}</option>' for d in dash_values)
        band_items.append(f'''
        <details class="annotation-item" data-ui-key="ann-item-band-{esc(item.get('id',''))}">
          <summary><span class="ann-kind band-kind">RANGE</span><span class="ann-name">{esc(item.get('label','Range'))}</span><span class="ann-value">{item.get('min')}…{item.get('max')}</span></summary>
          <div class="annotation-edit">
            <form method="post" action="/annotation/range/{tab_id}" data-ann-kind="band" data-ann-id="{esc(item.get('id',''))}">
              <input type="hidden" name="item_id" value="{esc(item.get('id',''))}">
              <label>Label<input name="label" value="{esc(item.get('label',''))}"></label>
              <div class="ann-grid"><label>Min<input name="min" type="number" step="any" value="{item.get('min')}"></label><label>Max<input name="max" type="number" step="any" value="{item.get('max')}"></label></div>
              <div class="ann-grid"><label>Axis<select name="axis"><option value="left" {'selected' if item.get('axis')!='right' else ''}>Primary</option><option value="right" {'selected' if item.get('axis')=='right' else ''}>Secondary</option></select></label><label>Style<select name="dash">{dash_options}</select></label></div>
              <div class="ann-grid three"><label>Line<input name="line_color" type="color" value="{esc(item.get('line_color','#12B76A'))}"></label><label>Fill<input name="fill_color" type="color" value="{esc(item.get('fill_color','#12B76A'))}"></label><label>Opacity %<input name="opacity" type="number" min="0" max="80" step="1" value="{round(float(item.get('opacity',.1))*100):g}"></label></div>
              <div class="ann-grid"><label>Width<input name="width" type="number" min="0.5" max="8" step="0.5" value="{item.get('width',1.2)}"></label><label class="check-label"><input name="show_label" type="checkbox" value="1" {'checked' if item.get('show_label',True) else ''}>Show label</label></div>
              <div class="ann-actions"><button class="btn compact" type="submit">Save</button><button class="btn compact danger-soft" type="button" onclick="deleteAnnotation('band','{esc(item.get('id',''))}')">Delete</button><label class="visibility-toggle"><input type="checkbox" {checked} onchange="toggleAnnotation('band','{esc(item.get('id',''))}',this.checked)"> Show</label></div>
            </form>
          </div>
        </details>''')
    comment_items = []
    for item in ts.comments:
        checked = "checked" if item.get("visible", True) else ""
        comment_items.append(f'''
        <details class="annotation-item" data-ui-key="ann-item-comment-{esc(item.get('id',''))}">
          <summary><span class="ann-kind comment-kind">NOTE</span><span class="ann-name">{esc(item.get('text',''))}</span></summary>
          <div class="annotation-edit">
            <form method="post" action="/annotation/comment-edit/{tab_id}" data-ann-kind="comment" data-ann-id="{esc(item.get('id',''))}">
              <input type="hidden" name="item_id" value="{esc(item.get('id',''))}">
              <label>Comment<textarea name="text" rows="3">{esc(item.get('text',''))}</textarea></label>
              <div class="ann-grid"><label>X anchor<input name="x" value="{esc(item.get('x',''))}" title="Timestamp or numeric X position"></label><label>Y anchor<input name="y" type="number" step="any" value="{item.get('y')}"></label></div>
              <div class="ann-grid"><label>Text X offset<input name="ax" type="number" step="1" value="{item.get('ax',32)}"></label><label>Text Y offset<input name="ay" type="number" step="1" value="{item.get('ay',-38)}"></label></div>
              <div class="ann-grid"><label>Text color<input name="text_color" type="color" value="{esc(item.get('text_color','#101828'))}"></label><label>Background<input name="bg_color" type="color" value="{esc(item.get('bg_color','#FFFFFF'))}"></label></div>
              <div class="ann-grid"><label>Border color<input name="border_color" type="color" value="{esc(item.get('border_color', item.get('text_color','#101828')))}"></label><label>Border width<input name="border_width" type="number" min="0" max="8" step="0.5" value="{item.get('border_width',1)}"></label></div>
              <div class="ann-grid"><label>Box width (px)<input name="box_width" type="number" min="0" max="900" step="10" value="{item.get('box_width',0)}" title="0 = auto width"></label><label>Box height (px)<input name="box_height" type="number" min="0" max="600" step="10" value="{item.get('box_height',0)}" title="0 = auto height"></label></div>
              <div class="ann-grid"><label>Font size<input name="font_size" type="number" min="8" max="28" value="{item.get('font_size',12)}"></label><label class="check-label"><input name="show_arrow" type="checkbox" value="1" {'checked' if item.get('show_arrow',True) else ''}>Show arrow</label></div>
              <div class="ann-grid"><label>Arrow color<input name="arrow_color" type="color" value="{esc(item.get('arrow_color', item.get('border_color','#101828')))}"></label><label>Arrow width<input name="arrow_width" type="number" min="0.5" max="8" step="0.5" value="{item.get('arrow_width',1.2)}"></label></div>
              <div class="ann-actions"><button class="btn compact" type="submit">Save</button><button class="btn compact danger-soft" type="button" onclick="deleteAnnotation('comment','{esc(item.get('id',''))}')">Delete</button><label class="visibility-toggle"><input type="checkbox" {checked} onchange="toggleAnnotation('comment','{esc(item.get('id',''))}',this.checked)"> Show</label></div>
            </form>
          </div>
        </details>''')

    empty = '<div class="annotation-empty">No engineering annotations yet.</div>' if not (line_items or band_items or comment_items) else ''
    return f'''
    <details class="side-card collapsible annotation-card" data-ui-key="annotations" open>
      <summary><span>Annotations & limits</span><span class="summary-hint">{len(ts.reference_lines)+len(ts.range_bands)+len(ts.comments)} items</span></summary>
      <div class="details-body">
        <div class="annotation-create-stack">
          <details class="ann-create-panel" data-ui-key="ann-create-line">
            <summary><span><span class="ann-kind line-kind">LINE</span> Reference line</span><span class="ann-create-action">Add</span></summary>
            <form class="ann-create" method="post" action="/annotation/line/{tab_id}">
              <div class="ann-grid"><label>Label<input name="label" placeholder="Upper Limit"></label><label>Value<input name="value" type="number" step="any" required placeholder="410"></label></div>
              <div class="ann-grid"><label>Axis<select name="axis"><option value="left">Primary</option><option value="right">Secondary</option></select></label><label>Style<select name="dash"><option value="dash">Dash</option><option value="solid">Solid</option><option value="dot">Dot</option><option value="dashdot">Dash-dot</option><option value="longdash">Long dash</option></select></label></div>
              <div class="ann-grid three"><label>Color<input name="color" type="color" value="#D92D20"></label><label>Width<input name="width" type="number" min="0.5" max="8" step="0.5" value="1.5"></label><label class="check-label"><input name="show_label" type="checkbox" value="1" checked>Label</label></div>
              <button class="btn compact full" type="submit">+ Add reference line</button>
            </form>
          </details>
          <details class="ann-create-panel" data-ui-key="ann-create-band">
            <summary><span><span class="ann-kind band-kind">RANGE</span> Tolerance / range band</span><span class="ann-create-action">Add</span></summary>
            <form class="ann-create" method="post" action="/annotation/range/{tab_id}">
              <label>Range name<input name="label" placeholder="Accept range"></label>
              <div class="ann-grid"><label>Min<input name="min" type="number" step="any" required></label><label>Max<input name="max" type="number" step="any" required></label></div>
              <div class="ann-grid"><label>Axis<select name="axis"><option value="left">Primary</option><option value="right">Secondary</option></select></label><label>Style<select name="dash"><option value="dash">Dash</option><option value="solid">Solid</option><option value="dot">Dot</option><option value="dashdot">Dash-dot</option></select></label></div>
              <div class="ann-grid three"><label>Line<input name="line_color" type="color" value="#12B76A"></label><label>Fill<input name="fill_color" type="color" value="#12B76A"></label><label>Opacity %<input name="opacity" type="number" min="0" max="80" value="10"></label></div>
              <input name="width" type="hidden" value="1.2"><input name="show_label" type="hidden" value="1">
              <button class="btn compact full" type="submit">+ Add range band</button>
            </form>
          </details>
          <button class="btn compact full comment-add-btn" type="button" onclick="armGraphComment()">+ Add comment on graph</button>
        </div>
        <div class="annotation-drag-tip">Drag a reference line, range band, or comment box directly on the graph. Dragged values are saved automatically.</div>
        <div class="annotation-list">{''.join(line_items)}{''.join(band_items)}{''.join(comment_items)}{empty}</div>
      </div>
    </details>
    '''


def render_active_tab(state: BrowserSession, tab_id: int) -> str:
    ts = state.tabs[tab_id]
    selected_count = len(ts.selected_signals)
    left_count = sum(
        1
        for sig in ts.selected_signals
        if ts.axis_assignments.get(sig, default_axis_for_signal(tab_id, sig)) == "left"
    )
    right_count = selected_count - left_count
    dual_legend_class = " dual-legend" if left_count and right_count else ""

    signal_rows = render_signal_rows(state, tab_id)
    image_src = f"/plot/{tab_id}.png?v={ts.revision}"
    mode_title = state.preset_name if state.loaded_preset else "Normal Plot Mode"
    mode_hint = "Imported .preset active" if state.loaded_preset else "Manual signal selection"
    reset_label = "Reload preset tab" if state.loaded_preset else "Clear tab"

    cycle_candidates = [sig for sig in ts.selected_signals if sig in state.numeric_cols]
    if cycle_candidates:
        cycle_signal_options = "".join(
            f'<option value="{esc(sig)}">{esc(sig)}</option>' for sig in cycle_candidates
        )
        cycle_disabled = ""
        cycle_hint = "Select a plotted signal, then analyze its Tmin/Tmax cycle pattern."
    else:
        cycle_signal_options = '<option value="">No plotted signal</option>'
        cycle_disabled = "disabled"
        cycle_hint = "Select at least one signal and Apply & Refresh Plot first."

    return f"""
    <div class="workspace">
        <aside class="sidebar">
            <form method="post" action="/update/{tab_id}" id="tabForm">
                <input type="hidden" name="numeric_count" value="{len(state.numeric_cols)}">

                <section class="side-card overview-card">
                    <div class="side-card-head">
                        <div>
                            <div class="eyebrow">PLOT MODE · {esc(mode_hint)}</div>
                            <h2>{esc(mode_title)}</h2>
                        </div>
                    </div>
                    <div class="metric-row">
                        <div class="mini-metric"><strong>{selected_count}</strong><span>Signals</span></div>
                        <div class="mini-metric primary-metric"><strong>{left_count}</strong><span>Primary</span></div>
                        <div class="mini-metric secondary-metric"><strong>{right_count}</strong><span>Secondary</span></div>
                    </div>
                    <button class="btn primary full" type="submit">Apply & Refresh Plot</button>
                </section>

                <details class="side-card collapsible" data-ui-key="plot-setup" open>
                    <summary>
                        <span>Plot setup</span>
                        <span class="summary-hint">Title & axes</span>
                    </summary>
                    <div class="details-body">
                        <div class="field">
                            <label>Graph title</label>
                            <input type="text" name="title" value="{esc(ts.title)}">
                        </div>
                        <label class="toggle-row">
                            <input type="checkbox" name="auto_title" value="1" {'checked' if ts.auto_title else ''}>
                            <span>Auto title from selected signals</span>
                        </label>
                        <div class="field">
                            <label>X-axis label</label>
                            <input type="text" name="x_label" value="{esc(ts.x_label)}">
                        </div>

                        <div class="axis-block primary-axis-block">
                            <div class="axis-block-title"><span class="axis-dot primary-dot"></span>Primary</div>
                            <div class="field">
                                <label>Axis label</label>
                                <input type="text" name="left_label" value="{esc(ts.left_label)}">
                            </div>
                            <div class="triple-fields">
                                <div class="field"><label>Min</label><input type="text" name="left_min" value="{esc(ts.left_min)}"></div>
                                <div class="field"><label>Max</label><input type="text" name="left_max" value="{esc(ts.left_max)}"></div>
                                <div class="field"><label>Step</label><input type="text" name="left_step" value="{esc(ts.left_step)}"></div>
                            </div>
                        </div>

                        <div class="axis-block secondary-axis-block">
                            <div class="axis-block-title"><span class="axis-dot secondary-dot"></span>Secondary</div>
                            <div class="field">
                                <label>Axis label</label>
                                <input type="text" name="right_label" value="{esc(ts.right_label)}">
                            </div>
                            <div class="triple-fields">
                                <div class="field"><label>Min</label><input type="text" name="right_min" value="{esc(ts.right_min)}"></div>
                                <div class="field"><label>Max</label><input type="text" name="right_max" value="{esc(ts.right_max)}"></div>
                                <div class="field"><label>Step</label><input type="text" name="right_step" value="{esc(ts.right_step)}"></div>
                            </div>
                        </div>

                        <div class="button-grid">
                            <button class="btn" type="submit" formaction="/reset/{tab_id}">{esc(reset_label)}</button>
                            <button class="btn" type="button" onclick="openPlotExportDialog({tab_id})">Save PNG</button>
                        </div>
                    </div>
                </details>

                <details class="side-card collapsible signals-card" data-ui-key="data-series" open>
                    <summary>
                        <span class="summary-stack"><span class="summary-eyebrow">DATA SERIES</span><span>Signals</span></span>
                        <span class="summary-hint"><span class="count-badge">{len(state.numeric_cols)}</span></span>
                    </summary>
                    <div class="details-body signal-details-body">
                        <div class="signal-search-wrap">
                            <input id="signalSearch" type="search" placeholder="Search signals" oninput="filterSignals()">
                        </div>
                        <div class="signal-quick-actions">
                            <button class="text-btn" type="button" onclick="selectVisible(true)">Select visible</button>
                            <button class="text-btn" type="button" onclick="selectVisible(false)">Clear</button>
                            <button class="text-btn primary-text-btn" type="button" onclick="assignChecked('left')">→ Primary</button>
                            <button class="text-btn secondary-text-btn" type="button" onclick="assignChecked('right')">→ Secondary</button>
                        </div>
                        <div class="signal-list" id="signalList">{signal_rows}</div>
                    </div>
                </details>
            </form>
            {render_annotation_manager(ts, tab_id)}
        </aside>

        <main class="main-column">
            <section class="plot-card">
                <div class="plot-toolbar">
                    <div>
                        <div class="eyebrow">INTERACTIVE WORKSPACE</div>
                        <h2>Graph analysis</h2>
                        <div class="toolbar-help">Wheel zoom · drag zoom · legend click hide/show · double-click isolate</div>
                    </div>
                    <div class="toolbar-actions">
                        <button class="btn compact" type="button" onclick="resetInteractiveView()">Reset view</button>
                        <button class="btn compact" type="button" onclick="autoscaleInteractiveView()">Autoscale</button>
                        <button class="btn compact" type="button" onclick="showAllInteractiveSignals()">Show all</button>
                        <a class="btn compact" href="{image_src}" target="_blank" onclick="setStaticPlotHref(this)">Static PNG</a>
                    </div>
                </div>
                <div class="interactive-plot-wrap{dual_legend_class}">
                    <div id="interactivePlot" class="interactive-plot">
                        <div class="plot-loading"><span class="spinner"></span>Loading graph...</div>
                    </div>
                </div>
            </section>

            <section class="analysis-card">
                <div class="analysis-head">
                    <div>
                        <div class="eyebrow">SIGNAL ANALYSIS</div>
                        <h2>Cursor & range statistics</h2>
                        <div id="analysisRangeText" class="toolbar-help">Statistics follow the current visible X range.</div>
                    </div>
                    <div class="cursor-actions">
                        <button id="cursorABtn" class="btn cursor-a-btn" type="button" onclick="armCursor('A')">Set Cursor A</button>
                        <button id="cursorBBtn" class="btn cursor-b-btn" type="button" onclick="armCursor('B')">Set Cursor B</button>
                        <button class="btn compact" type="button" onclick="clearCursors()">Clear cursors</button>
                    </div>
                </div>

                <div id="cursorInstruction" class="cursor-instruction">Choose Cursor A or B, then click a point on the graph.</div>

                <div class="analysis-metrics">
                    <div class="analysis-metric cursor-a-metric"><span>Cursor A</span><strong id="cursorAValue">—</strong></div>
                    <div class="analysis-metric cursor-b-metric"><span>Cursor B</span><strong id="cursorBValue">—</strong></div>
                    <div class="analysis-metric"><span id="deltaXLabel">Δ Time</span><strong id="deltaXValue">—</strong></div>
                    <div class="analysis-metric"><span>Analysis range</span><strong id="analysisModeValue">Current view</strong></div>
                </div>

                <div class="stats-wrap">
                    <table class="stats-table">
                        <thead>
                            <tr>
                                <th>Signal</th>
                                <th>Axis</th>
                                <th>Cursor A</th>
                                <th>Cursor B</th>
                                <th>Δ Value</th>
                                <th>Min</th>
                                <th>Max</th>
                                <th>Average</th>
                                <th>Samples</th>
                            </tr>
                        </thead>
                        <tbody id="analysisTableBody">
                            <tr><td colspan="9" class="empty-table">Load the interactive graph to calculate statistics.</td></tr>
                        </tbody>
                    </table>
                </div>
                <div class="analysis-note">When both cursors are set, Min/Max/Average use the A↔B interval. Otherwise they use the current zoomed X range. Hidden legend traces are excluded.</div>
            </section>

            <section class="analysis-card cycle-card">
                <div class="analysis-head cycle-head">
                    <div>
                        <div class="eyebrow">CYCLE ANALYSIS</div>
                        <h2>Automatic cycle & soak analysis</h2>
                        <div id="cycleHint" class="toolbar-help">{esc(cycle_hint)}</div>
                    </div>
                    <div class="cycle-controls">
                        <label>Signal
                            <select id="cycleSignal" {cycle_disabled}>{cycle_signal_options}</select>
                        </label>
                        <label>Profile
                            <select id="cycleProfileMode" title="Auto keeps legacy extrema for slow temperature cycling and uses dominant plateau temperatures for fast thermal shock.">
                                <option value="auto" selected>Auto</option>
                                <option value="temp_cycling">Temp cycling</option>
                                <option value="thermal_shock">Thermal shock</option>
                            </select>
                        </label>
                        <label>Tolerance ±°C
                            <input id="cycleTolerance" type="number" min="0.01" step="0.1" value="2.0">
                        </label>
                        <label>Fallback sec/sample
                            <input id="cycleSampleTime" type="number" min="0.001" step="1" value="10">
                        </label>
                        <button id="cycleAnalyzeBtn" class="btn cycle-analyze-btn" type="button" onclick="runCycleAnalysis()" {cycle_disabled}>Analyze cycles</button>
                        <button id="cycleExportBtn" class="btn compact" type="button" onclick="exportCycleAnalysisCSV()" disabled>Export CSV</button>
                    </div>
                </div>
                <div id="cycleStatus" class="cycle-status">Auto profile keeps the existing extrema method for slow temperature cycling and switches to dominant plateau detection for fast thermal shock. Complete cycle = Tmin → Tmax → Tmin or Tmax → Tmin → Tmax.</div>
                <div class="cycle-metrics">
                    <div class="cycle-metric"><span>Complete cycles</span><strong id="cycleCount">—</strong></div>
                    <div class="cycle-metric"><span>Tmin</span><strong id="cycleTmin">—</strong></div>
                    <div class="cycle-metric"><span>Tmax</span><strong id="cycleTmax">—</strong></div>
                    <div class="cycle-metric"><span>Avg time / cycle</span><strong id="cycleAvgTime">—</strong></div>
                    <div class="cycle-metric"><span>Avg soak Tmin</span><strong id="cycleAvgSoakMin">—</strong></div>
                    <div class="cycle-metric"><span>Avg soak Tmax</span><strong id="cycleAvgSoakMax">—</strong></div>
                </div>
                <div class="stats-wrap cycle-table-wrap">
                    <table class="stats-table">
                        <thead><tr><th>Cycle</th><th>Start state</th><th>Cycle start</th><th>Cycle end</th><th>Soak Tmin</th><th>Soak Tmax</th><th>Time / cycle</th></tr></thead>
                        <tbody id="cycleTableBody"><tr><td colspan="7" class="empty-table">Run Cycle Analysis to calculate cycle details.</td></tr></tbody>
                    </table>
                </div>
                <div id="cycleTimeSource" class="analysis-note">Timestamp duration is used when available; otherwise the fallback sample time is used.</div>
            </section>
        </main>
    </div>

    <div id="commentModal" class="modal-backdrop" hidden onclick="if(event.target===this) closeCommentModal()">
      <div class="comment-modal" role="dialog" aria-modal="true" aria-labelledby="commentModalTitle">
        <div class="modal-head"><div><div class="eyebrow">GRAPH COMMENT</div><h3 id="commentModalTitle">Add comment</h3></div><button class="icon-close" type="button" onclick="closeCommentModal()">×</button></div>
        <div class="modal-body">
          <div id="commentPointHint" class="modal-point">Click a graph point first.</div>
          <div class="modal-help">Enable Arrow when you want the note to point to the selected graph location. After saving, drag the comment box to reposition it; resize it from the NOTE editor using Box width / Box height.</div>
          <label>Comment<textarea id="commentText" rows="4" placeholder="Describe the event or observation..."></textarea></label>
          <div class="ann-grid"><label>Axis<select id="commentAxis"><option value="left">Primary</option><option value="right">Secondary</option></select></label><label>Font size<input id="commentFontSize" type="number" min="8" max="28" value="12"></label></div>
          <div class="ann-grid"><label>Text color<input id="commentTextColor" type="color" value="#101828"></label><label>Background<input id="commentBgColor" type="color" value="#FFFFFF"></label></div>
          <div class="ann-grid"><label>Border color<input id="commentBorderColor" type="color" value="#101828"></label><label>Border width<input id="commentBorderWidth" type="number" min="0" max="8" step="0.5" value="1"></label></div>
          <div class="ann-grid"><label>Box width (px)<input id="commentBoxWidth" type="number" min="0" max="900" step="10" value="0" title="0 = automatic width"></label><label>Box height (px)<input id="commentBoxHeight" type="number" min="0" max="600" step="10" value="0" title="0 = automatic height"></label></div>
          <label class="check-label"><input id="commentArrow" type="checkbox" checked> Show arrow to selected graph point</label>
          <div class="ann-grid"><label>Arrow color<input id="commentArrowColor" type="color" value="#101828"></label><label>Arrow width<input id="commentArrowWidth" type="number" min="0.5" max="8" step="0.5" value="1.2"></label></div>
        </div>
        <div class="modal-actions"><button class="btn" type="button" onclick="closeCommentModal()">Cancel</button><button class="btn primary" type="button" onclick="saveGraphComment()">Add comment</button></div>
      </div>
    </div>
    """


V42_CSS = r"""
:root {
    --bg: #f6f7f9;
    --surface: #ffffff;
    --surface-soft: #fafbfc;
    --panel: var(--surface);
    --soft: var(--surface-soft);
    --text: #172033;
    --muted: #667085;
    --subtle: #98a2b3;
    --line: #e4e7ec;
    --line-strong: #d0d5dd;
    --accent: #1769aa;
    --accent-hover: #12598f;
    --primary-axis: #1F4E79;
    --secondary-axis: #A61C00;
    --success: #067647;
    --warning: #b54708;
    --error: #b42318;
    --radius: 12px;
    --shadow: 0 1px 2px rgba(16,24,40,.03), 0 1px 3px rgba(16,24,40,.06);
}
* { box-sizing: border-box; }
html { background: var(--bg); }
body {
    margin: 0;
    color: var(--text);
    background: var(--bg);
    font-family: Inter, "Segoe UI", Roboto, Arial, sans-serif;
    font-size: 14px;
    -webkit-font-smoothing: antialiased;
}
button, input, select { font: inherit; }
form { margin:0; }
.app-header {
    position: sticky;
    top: 0;
    z-index: 30;
    background: rgba(255,255,255,.96);
    backdrop-filter: blur(10px);
    border-bottom: 1px solid var(--line);
}
.topbar {
    width: min(1880px, calc(100% - 36px));
    min-height: 66px;
    margin: 0 auto;
    display: flex;
    align-items: center;
    justify-content: space-between;
    gap: 18px;
}
.brand { display:flex; align-items:center; gap:12px; min-width:250px; flex:0 0 auto; }
.brand-mark {
    width: 38px; height: 38px; border-radius: 10px;
    display: grid; place-items: center;
    background: linear-gradient(145deg, #eaf5fc, #f7fbfe);
    color: var(--accent);
    border: 1px solid #cfe5f4;
    box-shadow: 0 1px 2px rgba(31,111,162,.10);
}
.brand-mark svg { width: 27px; height: 27px; overflow: visible; }
.brand-mark .signal-wave { fill:none; stroke:currentColor; stroke-width:2.2; stroke-linecap:round; stroke-linejoin:round; }
.brand h1 { margin: 0; font-size: 17px; line-height: 1.2; font-weight: 720; letter-spacing: -.015em; }
.brand-studio { color: var(--accent); font-weight: 760; }
.brand .version { margin-top: 3px; font-size: 10.5px; color: var(--muted); letter-spacing:.005em; }
.header-tools {
    display:flex; align-items:center; justify-content:flex-end; gap:8px; flex-wrap:nowrap;
    min-width:0; white-space:nowrap;
}
.upload-form, .preset-form, .preset-tools, .report-actions {
    display:flex; align-items:center; gap:7px; margin:0;
}
.preset-tools { padding-left:9px; border-left:1px solid var(--line); }
.header-inline-label {
    color:var(--muted); font-size:9px; font-weight:700; text-transform:uppercase;
    letter-spacing:.045em; white-space:nowrap; margin-right:-2px;
}
.preset-file { width:166px; font-size:11px; }
.preset-mode-badge {
    display:inline-flex; align-items:center; gap:5px; height:36px; padding:0 9px;
    max-width:210px; overflow:hidden; text-overflow:ellipsis;
    border-radius:8px; background:#f8fafc; border:1px solid var(--line);
    color:#475467; font-size:10px; white-space:nowrap;
}
.preset-mode-badge.active { background:#f0fdf4; border-color:#bbf7d0; color:#166534; }
.upload-field label { display:block; font-size:10px; color:var(--muted); margin-bottom:3px; text-transform:uppercase; letter-spacing:.04em; }
.file-picker { width:228px; font-size:11px; }
.skip-input { width:60px !important; }
.header-tools input[type="file"], .header-tools input[type="number"] { height:36px; }
.header-tools input[type="file"] { padding:3px 5px; }
.header-tools input[type="file"]::file-selector-button {
    height:28px; margin-right:7px; padding:0 9px; border:0; border-right:1px solid var(--line);
    background:#f9fafb; color:#344054; font-size:10px; cursor:pointer;
}
.file-picker::file-selector-button { background:#eff8ff !important; color:#175cd3 !important; }
.preset-file::file-selector-button { background:#f4f3ff !important; color:#5925dc !important; }
input[type="text"], input[type="search"], input[type="number"], input[type="file"], select {
    border: 1px solid var(--line-strong);
    border-radius: 8px;
    padding: 8px 10px;
    color: var(--text);
    background: #fff;
    transition: border-color .15s, box-shadow .15s;
}
input[type="text"], input[type="search"], input[type="number"], select { width: 100%; }
input:focus, select:focus {
    outline: none;
    border-color: #80b7dd;
    box-shadow: 0 0 0 3px rgba(23,105,170,.10);
}
.btn {
    border: 1px solid var(--line-strong);
    background: #fff;
    color: #344054;
    border-radius: 8px;
    padding: 8px 11px;
    min-height: 36px;
    display: inline-flex;
    align-items: center;
    justify-content: center;
    gap: 6px;
    text-decoration: none;
    cursor: pointer;
    white-space: nowrap;
    font-size: 12px;
    font-weight: 550;
    transition: all .15s;
}
.btn:hover { background: #f9fafb; border-color: #b8c0cc; }
.btn.primary { background: var(--accent); color: #fff; border-color: var(--accent); }
.btn.primary:hover { background: var(--accent-hover); border-color: var(--accent-hover); }

/* Top toolbar semantic actions: restrained color coding for faster scanning */
.btn.preset-import { background:#f4f3ff; color:#5925dc; border-color:#d9d6fe; }
.btn.preset-import:hover { background:#ebe9fe; border-color:#bdb4fe; }
.btn.preset-export { background:#ecfdf3; color:#067647; border-color:#abefc6; }
.btn.preset-export:hover { background:#dcfae6; border-color:#75e0a7; }
.btn.normal-mode { background:#f2f4f7; color:#344054; border-color:#d0d5dd; }
.btn.normal-mode:hover { background:#e4e7ec; border-color:#b8c0cc; }
.btn.pdf-report { background:#fff1f0; color:#b42318; border-color:#fecdca; }
.btn.pdf-report:hover { background:#fee4e2; border-color:#fda29b; }
.btn.word-report { background:#eff8ff; color:#175cd3; border-color:#b2ddff; }
.btn.word-report:hover { background:#dff1ff; border-color:#84caff; }
.btn.full { width: 100%; }
.btn.compact { padding: 6px 9px; min-height: 30px; }
.report-actions { display:flex; align-items:center; gap:7px; }
.page-shell { width: min(1880px, calc(100% - 36px)); margin: 16px auto 28px; }
.dataset-strip {
    display: flex; align-items: center; justify-content: space-between; gap: 12px;
    margin-bottom: 10px; color: var(--muted); font-size: 12px;
}
.dataset-main { display:flex; align-items:center; gap:8px; min-width:0; }
.dataset-dot { width:8px; height:8px; border-radius:50%; background:#12b76a; box-shadow:0 0 0 3px #ecfdf3; flex:0 0 auto; }
.dataset-name { color:#344054; font-weight:600; overflow:hidden; text-overflow:ellipsis; white-space:nowrap; }
.status {
    border-radius: 9px; padding: 8px 11px; margin-bottom: 12px;
    border: 1px solid; font-size: 12px;
}
.status.info { color:#175cd3; background:#eff8ff; border-color:#b2ddff; }
.status.success { color:#067647; background:#ecfdf3; border-color:#abefc6; }
.status.warning { color:#b54708; background:#fffaeb; border-color:#fedf89; }
.status.error { color:#b42318; background:#fef3f2; border-color:#fecdca; }
.tabs {
    display:flex; gap:5px; overflow-x:auto; padding:0 0 10px; scrollbar-width:thin;
}
.tab-link {
    text-decoration:none; color:#667085; border:1px solid transparent;
    padding:7px 11px; border-radius:8px; font-size:12px; white-space:nowrap;
}
.tab-link:hover { background:#fff; border-color:var(--line); }
.tab-link.active { color:#175cd3; background:#fff; border-color:#b2ddff; box-shadow:var(--shadow); font-weight:650; }
.workspace {
    display:grid;
    grid-template-columns: 420px minmax(0, 1fr);
    gap:18px;
    align-items:start;
}
.sidebar {
    display:flex; flex-direction:column; gap:22px; position:sticky; top:82px;
    height:calc(100vh - 94px); max-height:calc(100vh - 94px); overflow-y:auto;
    padding:0 8px 34px 0; scrollbar-gutter:stable; overscroll-behavior:contain;
}
.sidebar > form { display:flex; flex-direction:column; gap:22px; flex:0 0 auto; }
.sidebar > .side-card { flex:0 0 auto; }
.side-card, .plot-card, .analysis-card {
    background:var(--surface); border:1px solid var(--line); border-radius:var(--radius); box-shadow:var(--shadow);
}
.side-card { padding:14px; }
.side-card-head { display:flex; align-items:center; justify-content:space-between; gap:10px; margin-bottom:12px; }
.side-card h2, .side-card h3, .plot-card h2, .analysis-card h2 { margin:0; letter-spacing:-.015em; }
.side-card h2 { font-size:16px; }
.side-card h3 { font-size:14px; }
.eyebrow { color:var(--subtle); font-size:9px; font-weight:700; letter-spacing:.09em; text-transform:uppercase; margin-bottom:4px; }
.metric-row { display:grid; grid-template-columns:repeat(3,1fr); gap:7px; margin:10px 0 12px; }
.mini-metric { background:var(--surface-soft); border:1px solid #eef0f3; border-radius:9px; padding:8px; }
.mini-metric strong { display:block; font-size:16px; font-weight:650; }
.mini-metric span { color:var(--muted); font-size:10px; }
.primary-metric strong { color:var(--primary-axis); }
.secondary-metric strong { color:var(--secondary-axis); }
details.collapsible { padding:0; overflow:hidden; }
details.collapsible summary {
    list-style:none; cursor:pointer; padding:12px 14px; display:flex; justify-content:space-between; align-items:center;
    font-weight:650; font-size:13px; user-select:none;
}
details.collapsible summary::-webkit-details-marker { display:none; }
details.collapsible summary::after { content:"⌄"; color:var(--subtle); font-size:14px; margin-left:8px; }
details.collapsible[open] summary::after { transform:rotate(180deg); }
.summary-hint { color:var(--subtle); font-size:10px; font-weight:500; margin-left:auto; }
.details-body { border-top:1px solid var(--line); padding:13px; }
.field { margin-bottom:10px; }
.field label { display:block; color:var(--muted); font-size:10px; font-weight:600; margin-bottom:4px; }
.toggle-row { display:flex; align-items:center; gap:8px; color:#475467; font-size:11px; margin:-2px 0 12px; cursor:pointer; }
.axis-block { border-radius:9px; padding:10px; margin-top:10px; background:#fbfcfd; border:1px solid #edf0f3; }
.primary-axis-block { border-left:3px solid var(--primary-axis); }
.secondary-axis-block { border-left:3px solid var(--secondary-axis); }
.axis-block-title { display:flex; align-items:center; gap:6px; font-size:11px; font-weight:700; margin-bottom:9px; color:#475467; }
.axis-dot { width:7px; height:7px; border-radius:50%; display:inline-block; }
.primary-dot { background:var(--primary-axis); }
.secondary-dot { background:var(--secondary-axis); }
.triple-fields { display:grid; grid-template-columns:repeat(3,1fr); gap:7px; }
.triple-fields .field { margin-bottom:0; }
.button-grid { display:grid; grid-template-columns:1fr 1fr; gap:7px; margin-top:12px; }
.signals-card { padding:0; overflow:hidden; }
.signal-details-body { padding-top:11px; }
.summary-stack { display:flex; flex-direction:column; gap:1px; min-width:0; }
.summary-eyebrow { color:var(--subtle); font-size:8px; font-weight:750; letter-spacing:.08em; text-transform:uppercase; }
.signal-head { margin-bottom:9px; }
.count-badge { min-width:27px; height:23px; display:grid; place-items:center; border-radius:12px; background:#f2f4f7; color:#667085; font-size:10px; font-weight:650; }
.signal-search-wrap { margin-bottom:7px; }
.signal-search-wrap input { font-size:12px; padding:7px 9px; }
.signal-quick-actions { display:flex; flex-wrap:wrap; gap:4px 9px; margin-bottom:7px; }
.text-btn { border:0; background:transparent; padding:2px 0; color:#667085; font-size:10px; cursor:pointer; }
.text-btn:hover { color:#344054; text-decoration:underline; }
.primary-text-btn { color:var(--primary-axis); }
.secondary-text-btn { color:var(--secondary-axis); }
.signal-list { border:1px solid var(--line); border-radius:9px; max-height:min(390px,42vh); overflow-y:auto; overscroll-behavior:auto; background:#fff; }
.signal-row {
    display:grid; grid-template-columns:minmax(0,1fr) 86px 80px 54px; gap:4px 6px; align-items:center;
    min-height:54px; padding:6px 7px; border-bottom:1px solid #f0f1f3; border-left:2px solid transparent;
}
.signal-row:last-child { border-bottom:0; }
.signal-row:hover { background:#f9fafb; }
.primary-row { border-left-color:var(--primary-axis); }
.secondary-row { border-left-color:var(--secondary-axis); }
.signal-check { grid-column:1 / -1; display:flex; align-items:flex-start; gap:6px; font-size:12px; min-width:0; cursor:pointer; }
.signal-check input { flex:0 0 auto; margin:2px 0 0; }
.signal-name {
    display:block; min-width:0; flex:1 1 auto; overflow-wrap:anywhere;
    white-space:normal; line-height:1.4; color:#344054; font-weight:600;
}
.signal-row:hover .signal-name { color:#101828; }
.signal-row > .axis-select { grid-column:2; }
.signal-row > .line-style-select { grid-column:3; }
.signal-row > .color-control { grid-column:4; }
.axis-select, .line-style-select { width:100%; padding:3px 4px; border-radius:6px; font-size:10px; min-width:0; }
.color-control { display:flex; align-items:center; gap:3px; justify-content:flex-end; }
.signal-color { width:27px; height:23px; padding:1px; border:1px solid var(--line-strong); border-radius:6px; background:#fff; cursor:pointer; }
.color-control.auto .signal-color { opacity:.72; }
.auto-color-btn { width:22px; height:24px; padding:0; border:1px solid var(--line); border-radius:6px; background:#fff; color:#667085; font-size:9px; font-weight:700; cursor:pointer; }
.auto-color-btn:hover { background:#f2f4f7; color:#344054; }
.color-control.auto .auto-color-btn { color:#175cd3; background:#eff8ff; border-color:#b2ddff; }
.preset-badge {
    display:none; flex:0 0 auto; color:#8a6100; background:#fffaeb; border:1px solid #fedf89;
    border-radius:5px; font-size:8px; line-height:14px; min-width:16px; height:16px; padding:0 3px;
    text-align:center; font-weight:700;
}
.preset-badge.show { display:inline-block; }
.main-column { min-width:0; display:flex; flex-direction:column; gap:18px; }
.plot-card { overflow:hidden; }
.plot-toolbar, .analysis-head { padding:14px 16px; display:flex; align-items:center; justify-content:space-between; gap:12px; border-bottom:1px solid var(--line); }
.plot-card h2, .analysis-card h2 { font-size:16px; }
.toolbar-help { color:var(--muted); font-size:10px; margin-top:3px; }
.toolbar-actions, .cursor-actions { display:flex; align-items:center; gap:6px; flex-wrap:wrap; justify-content:flex-end; }
.interactive-plot-wrap { width:100%; height:clamp(520px, 64vh, 720px); background:#fff; overflow:hidden; }
.interactive-plot-wrap.dual-legend { height:calc(clamp(520px, 64vh, 720px) + 95px); }
.interactive-plot { width:100%; height:100%; min-height:0; }
.plot-loading { width:100%; height:100%; display:flex; align-items:center; justify-content:center; gap:9px; color:var(--muted); font-size:12px; }
.spinner { width:15px; height:15px; border:2px solid #e4e7ec; border-top-color:#667085; border-radius:50%; animation:spin .8s linear infinite; }
@keyframes spin { to { transform:rotate(360deg); } }
.plot-error { min-height:300px; display:grid; place-items:center; color:var(--error); padding:20px; text-align:center; }
.analysis-card { overflow:hidden; }
.cursor-a-btn { border-color:#84adff; color:#175cd3; background:#eff8ff; }
.cursor-b-btn { border-color:#fda29b; color:#b42318; background:#fef3f2; }
.cursor-a-btn.armed, .cursor-b-btn.armed { box-shadow:0 0 0 3px rgba(23,105,170,.12); font-weight:700; }
.cursor-instruction { padding:8px 16px; color:#667085; background:#fafbfc; border-bottom:1px solid var(--line); font-size:10px; }
.analysis-metrics { display:grid; grid-template-columns:repeat(4,minmax(0,1fr)); gap:1px; background:var(--line); border-bottom:1px solid var(--line); }
.analysis-metric { padding:11px 14px; background:#fff; min-width:0; }
.analysis-metric span { display:block; color:var(--muted); font-size:9px; text-transform:uppercase; letter-spacing:.04em; margin-bottom:4px; }
.analysis-metric strong { display:block; font-size:13px; font-weight:650; white-space:nowrap; overflow:hidden; text-overflow:ellipsis; }
.cursor-a-metric { box-shadow:inset 3px 0 #2e90fa; }
.cursor-b-metric { box-shadow:inset 3px 0 #f04438; }
.stats-wrap { width:100%; overflow:auto; max-height:360px; }
.stats-table { width:100%; border-collapse:separate; border-spacing:0; font-size:10px; }
.stats-table th { position:sticky; top:0; z-index:2; background:#f9fafb; color:#667085; text-align:right; font-weight:650; padding:8px 10px; border-bottom:1px solid var(--line); white-space:nowrap; }
.stats-table th:first-child, .stats-table th:nth-child(2) { text-align:left; }
.stats-table td { padding:7px 10px; border-bottom:1px solid #f0f1f3; text-align:right; white-space:nowrap; }
.stats-table td:first-child, .stats-table td:nth-child(2) { text-align:left; }
.stats-table tr:hover td { background:#fafbfc; }
.axis-pill { display:inline-flex; align-items:center; border-radius:10px; padding:2px 6px; font-size:8px; font-weight:650; }
.axis-pill.primary { color:var(--primary-axis); background:#eef4f8; }
.axis-pill.secondary { color:var(--secondary-axis); background:#fff2ef; }
.empty-table { text-align:center !important; color:var(--subtle); padding:24px !important; }

.analysis-note { padding:8px 14px; color:var(--subtle); background:#fafbfc; font-size:9px; border-top:1px solid var(--line); }
.cycle-card { margin-top:0; }
.cycle-head { align-items:flex-end; }
.cycle-controls { display:flex; align-items:flex-end; justify-content:flex-end; gap:8px; flex-wrap:wrap; }
.cycle-controls label { display:flex; flex-direction:column; gap:4px; color:#667085; font-size:9px; text-transform:uppercase; letter-spacing:.035em; }
.cycle-controls select, .cycle-controls input { height:32px; border:1px solid var(--line-strong); border-radius:7px; background:#fff; color:var(--text); font-size:10px; padding:0 8px; }
.cycle-controls select { min-width:180px; max-width:300px; text-transform:none; }
.cycle-controls input { width:96px; }
.cycle-analyze-btn { background:#6941c6; border-color:#6941c6; color:#fff; }
.cycle-analyze-btn:hover { background:#53389e; border-color:#53389e; }
.cycle-status { padding:8px 16px; color:#667085; background:#fafbfc; border-bottom:1px solid var(--line); font-size:10px; }
.cycle-status.success { color:#067647; background:#ecfdf3; }
.cycle-status.warning { color:#b54708; background:#fffaeb; }
.cycle-status.error { color:#b42318; background:#fef3f2; }
.cycle-metrics { display:grid; grid-template-columns:repeat(6,minmax(0,1fr)); gap:1px; background:var(--line); border-bottom:1px solid var(--line); }
.cycle-metric { padding:11px 12px; background:#fff; min-width:0; }
.cycle-metric span { display:block; color:var(--muted); font-size:8px; text-transform:uppercase; letter-spacing:.04em; margin-bottom:4px; }
.cycle-metric strong { display:block; font-size:13px; font-weight:650; white-space:nowrap; overflow:hidden; text-overflow:ellipsis; }
.cycle-table-wrap { max-height:330px; }
.empty-state { text-align:center; padding:90px 20px; background:#fff; border:1px solid var(--line); border-radius:var(--radius); box-shadow:var(--shadow); }
.empty-icon { width:48px; height:48px; display:grid; place-items:center; margin:0 auto 12px; border-radius:13px; background:#eff8ff; color:#175cd3; font-size:22px; }
.empty-state h2 { margin:0 0 5px; font-size:17px; }
.empty-state p { margin:0; color:var(--muted); font-size:12px; }
.footer { text-align:center; color:var(--subtle); font-size:9px; margin:22px 0 8px; }

/* V4.3.1 compact header menus + appearance themes */
.header-menu { position:relative; flex:0 0 auto; }
.header-menu > summary { list-style:none; }
.header-menu > summary::-webkit-details-marker { display:none; }
.header-menu > summary::after { content:"⌄"; font-size:10px; margin-left:4px; opacity:.72; }
.header-menu[open] > summary::after { transform:rotate(180deg); }
.menu-panel {
    position:absolute; top:calc(100% + 7px); right:0; z-index:80;
    min-width:230px; padding:8px; border:1px solid var(--line);
    border-radius:10px; background:var(--surface); box-shadow:0 12px 30px rgba(16,24,40,.16);
}
.menu-panel.wide { width:310px; min-width:310px; max-width:calc(100vw - 24px); }
.menu-title { padding:3px 7px 7px; color:var(--subtle); font-size:9px; font-weight:750; letter-spacing:.07em; text-transform:uppercase; }
.menu-item {
    width:100%; min-height:34px; padding:7px 9px; border:0; border-radius:7px;
    display:flex; align-items:center; justify-content:space-between; gap:10px;
    background:transparent; color:var(--text); font-size:11px; text-decoration:none; cursor:pointer; text-align:left;
}
.menu-item:hover { background:var(--surface-soft); }
.menu-item .menu-note { color:var(--muted); font-size:9px; font-weight:500; }
.menu-separator { height:1px; margin:6px 0; background:var(--line); }
.menu-form {
    margin:0; width:100%; min-width:0;
    display:grid; grid-template-columns:minmax(0, 1fr); gap:7px;
}
.menu-form .preset-file { width:100%; min-width:0; margin:0; display:block; }
.menu-form .btn { width:100%; min-width:0; display:flex; }
.theme-trigger { min-width:84px; }

.theme-switch-wrap {height:36px;display:inline-flex;align-items:center;gap:7px;padding:0 8px;border:1px solid var(--line-strong);border-radius:999px;background:var(--surface-soft);flex:0 0 auto;box-shadow:inset 0 1px 2px rgba(16,24,40,.04)}
.theme-switch-icon {width:16px;height:16px;display:grid;place-items:center;color:#8b98a9;transition:color .18s ease}
.theme-switch-icon svg {width:15px;height:15px;fill:none;stroke:currentColor;stroke-width:1.8;stroke-linecap:round;stroke-linejoin:round}
.theme-switch-icon.moon svg path {fill:currentColor;stroke:none}
.theme-switch {position:relative;width:38px;height:20px;padding:0;border:0;border-radius:999px;background:#5f8292;cursor:pointer;flex:0 0 auto;box-shadow:inset 0 1px 3px rgba(16,24,40,.20);transition:background .18s ease,box-shadow .18s ease}
.theme-switch:focus-visible {outline:none;box-shadow:0 0 0 3px rgba(23,105,170,.16),inset 0 1px 3px rgba(16,24,40,.20)}
.theme-switch-knob {position:absolute;top:3px;left:3px;width:14px;height:14px;border-radius:50%;background:#fff;box-shadow:0 1px 3px rgba(16,24,40,.30);transition:transform .18s cubic-bezier(.2,.8,.2,1)}
.theme-switch[aria-checked="true"] {background:#355c6e}
.theme-switch[aria-checked="true"] .theme-switch-knob {transform:translateX(18px)}
html[data-theme="light"] .theme-switch-icon.sun {color:#4f6471}
html[data-theme="light"] .theme-switch-icon.moon {color:#93a1ad}
html[data-theme="dark"] .theme-switch-wrap {background:#17313d;border-color:#2f4f5e;box-shadow:inset 0 1px 3px rgba(0,0,0,.28)}
html[data-theme="dark"] .theme-switch-icon.sun {color:#78909c}
html[data-theme="dark"] .theme-switch-icon.moon {color:#b9d0da}

.preset-summary { max-width:210px; overflow:hidden; text-overflow:ellipsis; }

html[data-theme="dark"] {
    color-scheme:dark;
    --bg:#171c23;
    --surface:#202730;
    --surface-soft:#29323d;
    --text:#eef2f7;
    --muted:#bac4d0;
    --subtle:#96a2b2;
    --line:#3a4654;
    --line-strong:#4c5b6c;
    --accent:#58afe7;
    --accent-hover:#78c0ed;
    --primary-axis:#79c3ef;
    --secondary-axis:#ff9484;
    --success:#6cdda0;
    --warning:#f8c772;
    --error:#ff9388;
    --shadow:0 1px 2px rgba(0,0,0,.24),0 1px 4px rgba(0,0,0,.18);
}
html[data-theme="dark"] .app-header { background:rgba(29,35,44,.96); }
html[data-theme="dark"] .brand-mark { background:linear-gradient(145deg,#142638,#192f44); border-color:#29445d; color:#75bdec; box-shadow:none; }
html[data-theme="dark"] input[type="text"],
html[data-theme="dark"] input[type="search"],
html[data-theme="dark"] input[type="number"],
html[data-theme="dark"] input[type="file"],
html[data-theme="dark"] select { background:#252d38; color:var(--text); border-color:var(--line-strong); }
html[data-theme="dark"] .header-tools input[type="file"]::file-selector-button { background:#303946; color:#e4e9ef; border-color:var(--line); }
html[data-theme="dark"] .file-picker::file-selector-button { background:#17334a !important; color:#82c7f3 !important; }
html[data-theme="dark"] .preset-file::file-selector-button { background:#2c2445 !important; color:#c4b5fd !important; }
html[data-theme="dark"] .btn { background:#29323d; color:#e4e9ef; border-color:var(--line-strong); }
html[data-theme="dark"] .btn:hover { background:#343e4b; border-color:#627286; }
html[data-theme="dark"] .btn.primary { background:#176ca8; color:#fff; border-color:#2d83bd; }
html[data-theme="dark"] .preset-mode-badge { background:#1b2028; border-color:var(--line); color:#b7c0cc; }
html[data-theme="dark"] .preset-mode-badge.active { background:#14281e; border-color:#285d40; color:#77d49a; }
html[data-theme="dark"] .btn.preset-import { background:#292341; color:#c4b5fd; border-color:#554a82; }
html[data-theme="dark"] .btn.preset-export { background:#173326; color:#72d49a; border-color:#2d6747; }
html[data-theme="dark"] .btn.normal-mode { background:#222832; color:#d4dae4; border-color:#414a57; }
html[data-theme="dark"] .btn.pdf-report { background:#3a211f; color:#ff9388; border-color:#74423d; }
html[data-theme="dark"] .btn.word-report { background:#172f46; color:#82c7f3; border-color:#315d80; }
html[data-theme="dark"] .mini-metric,
html[data-theme="dark"] .axis-block,
html[data-theme="dark"] .count-badge,
html[data-theme="dark"] .signal-list,
html[data-theme="dark"] .interactive-plot-wrap,
html[data-theme="dark"] .analysis-metric,
html[data-theme="dark"] .cycle-metric,
html[data-theme="dark"] .empty-state { background:var(--surface); border-color:var(--line); }
html[data-theme="dark"] .signal-row { border-bottom-color:#394452; }
html[data-theme="dark"] .signal-row:hover,
html[data-theme="dark"] .stats-table tr:hover td,
html[data-theme="dark"] .cursor-instruction,
html[data-theme="dark"] .analysis-note,
html[data-theme="dark"] .cycle-status { background:var(--surface-soft); }
html[data-theme="dark"] .signal-name,
html[data-theme="dark"] .axis-block-title,
html[data-theme="dark"] .toggle-row { color:#d8dee9; }
html[data-theme="dark"] .signal-row:hover .signal-name { color:#fff; }
html[data-theme="dark"] .text-btn { color:#9fa9b8; }
html[data-theme="dark"] .text-btn:hover { color:#e6ebf2; }
html[data-theme="dark"] .signal-color,
html[data-theme="dark"] .auto-color-btn,
html[data-theme="dark"] .cycle-controls select,
html[data-theme="dark"] .cycle-controls input { background:#252d38; color:var(--text); border-color:var(--line-strong); }
html[data-theme="dark"] .stats-table th { background:#303946; color:#c6d0dc; }
html[data-theme="dark"] .stats-table td { border-bottom-color:#394452; }
html[data-theme="dark"] .axis-pill.primary { background:#162a39; }
html[data-theme="dark"] .axis-pill.secondary { background:#39201d; }
html[data-theme="dark"] .empty-icon { background:#17334a; color:#82c7f3; }
html[data-theme="dark"] .spinner { border-color:#39414d; border-top-color:#a9b3c1; }
html[data-theme="dark"] .cycle-status.success { color:#77d49a; background:#14281e; }
html[data-theme="dark"] .cycle-status.warning { color:#f6c76d; background:#342915; }
html[data-theme="dark"] .cycle-status.error { color:#ff9388; background:#3a211f; }
html[data-theme="dark"] .status.info { color:#82c7f3; background:#17334a; border-color:#315d80; }
html[data-theme="dark"] .status.success { color:#77d49a; background:#14281e; border-color:#2d6747; }
html[data-theme="dark"] .status.warning { color:#f6c76d; background:#342915; border-color:#705823; }
html[data-theme="dark"] .status.error { color:#ff9388; background:#3a211f; border-color:#74423d; }
html[data-theme="dark"] .tab-link:hover { background:var(--surface); }
html[data-theme="dark"] .tab-link.active { color:#82c7f3; background:var(--surface); border-color:#315d80; }

.annotation-card { overflow:hidden; }
.annotation-card .details-body { gap:10px; }
.annotation-create-stack { display:flex; flex-direction:column; gap:8px; min-width:0; }
.ann-create-panel { border:1px solid var(--line); border-radius:9px; background:var(--panel); overflow:hidden; min-width:0; }
.ann-create-panel > summary { list-style:none; cursor:pointer; display:flex; align-items:center; justify-content:space-between; gap:8px; padding:9px 10px; font-size:10px; font-weight:750; color:var(--text); }
.ann-create-panel > summary::-webkit-details-marker { display:none; }
.ann-create-panel > summary::after { content:"⌄"; margin-left:4px; color:var(--subtle); font-size:11px; }
.ann-create-panel[open] > summary::after { transform:rotate(180deg); }
.ann-create-panel > summary > span:first-child { display:flex; align-items:center; gap:7px; min-width:0; }
.ann-create-action { margin-left:auto; color:var(--subtle); font-size:9px; font-weight:650; }
.ann-create-panel > .ann-create { padding:10px; border-top:1px solid var(--line); background:var(--soft); }
.ann-add-title { font-size:11px; font-weight:800; color:var(--muted); text-transform:uppercase; letter-spacing:.05em; }
.ann-create, .annotation-edit form { display:flex; flex-direction:column; gap:8px; }
.ann-create label, .annotation-edit label, .comment-modal label { display:flex; flex-direction:column; gap:4px; font-size:10px; font-weight:700; color:var(--muted); }
.ann-create input, .ann-create select, .annotation-edit input, .annotation-edit select, .annotation-edit textarea, .comment-modal input, .comment-modal select, .comment-modal textarea { width:100%; min-width:0; border:1px solid var(--line); border-radius:7px; background:var(--panel); color:var(--text); padding:7px 8px; font:inherit; box-sizing:border-box; }
.ann-create input[type=color], .annotation-edit input[type=color], .comment-modal input[type=color] { height:34px; padding:3px; }
.ann-grid { display:grid; grid-template-columns:minmax(0,1fr) minmax(0,1fr); gap:7px; }
.ann-grid.three { grid-template-columns:minmax(0,1fr) minmax(0,.8fr) minmax(0,.8fr); }
.check-label { justify-content:center; align-items:flex-start; }
.check-label input { width:auto !important; margin-top:6px; }
.ann-divider { height:1px; background:var(--line); margin:3px 0; }
.annotation-drag-tip { margin-top:10px; padding:8px 9px; border-radius:7px; background:var(--surface-soft); color:var(--muted); font-size:9px; line-height:1.4; border:1px solid var(--line); }
.annotation-list { display:flex; flex-direction:column; gap:6px; margin-top:8px; }
.annotation-item { border:1px solid var(--line); border-radius:8px; overflow:hidden; background:var(--panel); }
.annotation-item > summary { cursor:pointer; display:grid; grid-template-columns:auto minmax(0,1fr) auto; gap:7px; align-items:center; padding:7px 8px; list-style:none; }
.annotation-item > summary::-webkit-details-marker { display:none; }
.ann-kind { font-size:8px; font-weight:800; border-radius:999px; padding:3px 5px; }
.line-kind { background:#fef3f2; color:#b42318; } .band-kind { background:#ecfdf3; color:#067647; } .comment-kind { background:#eff8ff; color:#175cd3; }
.ann-name { overflow:hidden; text-overflow:ellipsis; white-space:nowrap; font-size:10px; font-weight:700; }
.ann-value { font-size:9px; color:var(--muted); }
.annotation-edit { padding:8px; border-top:1px solid var(--line); background:var(--soft); }
.ann-actions { display:flex; gap:6px; align-items:center; }
.visibility-toggle { margin-left:auto; display:flex !important; flex-direction:row !important; align-items:center; gap:4px !important; font-size:9px !important; }
.visibility-toggle input { width:auto !important; }
.danger-soft { color:#b42318 !important; border-color:#fecdca !important; background:#fff5f4 !important; }
.comment-add-btn { border-color:#b2ddff !important; color:#175cd3 !important; background:#eff8ff !important; }
.annotation-empty { font-size:10px; color:var(--muted); text-align:center; padding:8px; }
.modal-backdrop {
    position:fixed; inset:0; z-index:10000; background:rgba(8,15,28,.72);
    -webkit-backdrop-filter:blur(3px); backdrop-filter:blur(3px);
    display:flex; align-items:center; justify-content:center; padding:24px;
}
.modal-backdrop[hidden] { display:none; }
body.modal-open { overflow:hidden; }
.comment-modal {
    width:min(500px,calc(100vw - 36px)); max-height:calc(100vh - 48px);
    background:var(--surface); color:var(--text); border:1px solid var(--line-strong);
    border-radius:14px; box-shadow:0 28px 90px rgba(0,0,0,.48); overflow:auto;
    isolation:isolate;
}
.modal-head { position:sticky; top:0; z-index:1; display:flex; justify-content:space-between; gap:12px; padding:16px 18px 12px; background:var(--surface); border-bottom:1px solid var(--line); }
.modal-head h3 { margin:2px 0 0; font-size:16px; } .icon-close { border:0; background:transparent; color:var(--muted); font-size:24px; cursor:pointer; }
.modal-body { padding:16px 18px; display:flex; flex-direction:column; gap:11px; background:var(--surface); }
.modal-point { padding:8px 10px; background:var(--surface-soft); border:1px solid var(--line); border-radius:8px; font-size:10px; color:var(--muted); }
.modal-help { font-size:10px; line-height:1.45; color:var(--muted); padding:8px 10px; border-left:3px solid #84adff; background:var(--surface-soft); border-radius:6px; }
.modal-actions { position:sticky; bottom:0; padding:12px 18px 16px; display:flex; justify-content:flex-end; gap:8px; background:var(--surface); border-top:1px solid var(--line); }
.csv-bind-modal { width:min(760px, calc(100vw - 28px)); }
.csv-file-note { font-size:10px; line-height:1.45; color:var(--muted); }
.csv-file-list { display:flex; flex-direction:column; gap:7px; max-height:330px; overflow:auto; padding:2px; }
.csv-file-row {
    display:grid; grid-template-columns:28px minmax(0,1fr) auto; gap:10px; align-items:center;
    border:1px solid var(--line); border-radius:9px; padding:9px 10px;
    background:var(--surface-soft); cursor:grab; user-select:none;
}
.csv-file-row:active { cursor:grabbing; }
.csv-file-row.dragging { opacity:.45; }
.csv-file-row.drag-over { border-color:#4aa3df; box-shadow:0 0 0 2px rgba(23,105,170,.12); }
.csv-drag-handle { color:var(--muted); font-size:16px; letter-spacing:-2px; text-align:center; }
.csv-file-main { min-width:0; display:flex; flex-direction:column; gap:2px; }
.csv-file-name { font-size:11px; font-weight:700; color:var(--text); overflow:hidden; text-overflow:ellipsis; white-space:nowrap; }
.csv-file-meta { font-size:9.5px; color:var(--muted); }
.csv-file-order { font-size:10px; font-weight:800; color:#175cd3; background:#eff8ff; border:1px solid #b2ddff; border-radius:999px; padding:3px 7px; }
.csv-bind-summary { display:flex; gap:8px; flex-wrap:wrap; font-size:10px; color:var(--muted); }
.csv-bind-summary span { background:var(--surface-soft); border:1px solid var(--line); border-radius:999px; padding:4px 8px; }
.csv-progress-wrap { display:flex; flex-direction:column; gap:6px; }
.csv-progress-track { height:9px; overflow:hidden; border-radius:999px; background:var(--surface-soft); border:1px solid var(--line); }
.csv-progress-bar { width:0%; height:100%; background:linear-gradient(90deg,#1769aa,#48b8d8,#9bdc50); transition:width .2s ease; }
.csv-progress-text { font-size:10px; color:var(--muted); }
.csv-modal-actions-left { margin-right:auto; display:flex; gap:7px; }
.csv-bind-success { color:#067647 !important; }
.csv-bind-error { color:#b42318 !important; }
html[data-theme="dark"] .csv-file-order { background:#172f46; color:#82c7f3; border-color:#315d80; }
.export-modal { width:min(650px, calc(100vw - 28px)); }
.export-modal .modal-body label { display:flex; flex-direction:column; gap:5px; font-size:11px; font-weight:700; color:var(--muted); }
.export-modal input[type=text] { width:100%; box-sizing:border-box; border:1px solid var(--line); background:var(--panel); color:var(--text); border-radius:8px; padding:9px 10px; font:inherit; }
.export-section-head { display:flex; justify-content:space-between; align-items:center; gap:12px; font-size:11px; font-weight:700; color:var(--muted); margin:2px 0 7px; }
.link-button { border:0; background:transparent; color:#1677c8; padding:0; font:inherit; cursor:pointer; }
.export-tab-grid { display:grid; grid-template-columns:repeat(2,minmax(0,1fr)); gap:7px; max-height:300px; overflow:auto; padding:2px; }
.export-tab-option { display:flex !important; flex-direction:row !important; align-items:flex-start; gap:8px !important; padding:9px 10px; border:1px solid var(--line); border-radius:8px; background:var(--surface-soft); cursor:pointer; }
.export-tab-option input { margin-top:2px; width:auto !important; }
.export-tab-option span { min-width:0; display:flex; flex-direction:column; gap:2px; }
.export-tab-option strong { color:var(--text); font-size:11px; }
.export-tab-option small { color:var(--muted); font-size:10px; overflow:hidden; text-overflow:ellipsis; white-space:nowrap; }
.save-path-note { font-size:10px; line-height:1.5; color:var(--muted); background:var(--surface-soft); border:1px solid var(--line); border-radius:8px; padding:9px 10px; }
@media (max-width:600px) { .export-tab-grid { grid-template-columns:1fr; } }
html[data-theme="dark"] .line-kind { background:#3a211f; color:#ff9388; }
html[data-theme="dark"] .band-kind { background:#173326; color:#72d49a; }
html[data-theme="dark"] .comment-kind { background:#172f46; color:#82c7f3; }
html[data-theme="dark"] .danger-soft { background:#3a211f !important; color:#ff9388 !important; border-color:#74423d !important; }
html[data-theme="dark"] .comment-add-btn { background:#172f46 !important; color:#82c7f3 !important; border-color:#315d80 !important; }

@media (max-width: 1500px) {
    .topbar { width:calc(100% - 24px); }
    .brand { min-width:220px; }
    .file-picker { width:190px; }
    .preset-file { width:140px; }
    .header-tools { gap:6px; }
    .header-tools .btn { padding-left:9px; padding-right:9px; }
}
@media (max-width: 1250px) {
    .topbar { padding:8px 0; flex-wrap:wrap; align-items:center; }
    .header-tools { width:100%; justify-content:flex-start; flex-wrap:wrap; overflow:visible; padding-bottom:2px; }
    .workspace { grid-template-columns:1fr; }
    .sidebar { position:static; height:auto; max-height:none; display:grid; grid-template-columns:1fr 1fr; overflow:visible; padding:0 0 28px; }
    .sidebar > form { display:contents; }
    .signals-card { grid-column:1 / -1; }
}
@media (max-width: 850px) {
    .topbar, .page-shell { width:calc(100% - 18px); }
    .topbar { padding:9px 0; align-items:flex-start; flex-direction:column; }
    .header-tools { width:100%; flex-wrap:wrap; overflow:visible; }
    .upload-form, .preset-form { width:auto; }
    .file-picker { width:min(260px, 55vw); }
    .preset-tools { padding-left:0; border-left:0; flex-wrap:wrap; }
    .menu-panel { left:0; right:auto; }
    .sidebar { display:flex; height:auto; max-height:none; overflow:visible; padding-bottom:28px; }
    .sidebar > form { display:flex; flex-direction:column; gap:18px; }
    .analysis-metrics { grid-template-columns:1fr 1fr; }
    .cycle-metrics { grid-template-columns:repeat(3,1fr); }
    .plot-toolbar, .analysis-head { align-items:flex-start; flex-direction:column; }
    .toolbar-actions, .cursor-actions { justify-content:flex-start; }
}
@media (max-width: 560px) {
    .upload-form, .preset-form { align-items:stretch; flex-direction:column; }
    .file-picker { width:100%; }
    .report-actions { width:100%; }
    .report-actions .btn { flex:1; }
    .header-menu { width:auto; }
    .menu-panel, .menu-panel.wide { min-width:min(310px, calc(100vw - 28px)); }
    .signal-row { grid-template-columns:minmax(0,1fr) 86px 80px 54px; }
    .preset-badge { display:none !important; }
    .analysis-metrics { grid-template-columns:1fr; }
    .cycle-metrics { grid-template-columns:1fr 1fr; }
}
"""


V42_JS = r"""
const ACTIVE_TAB = __ACTIVE_TAB__;
const ACTIVE_REVISION = __REVISION__;
const EXPORT_BASE = __EXPORT_BASE_JSON__;
const ACTIVE_PRESET_NAME = __PRESET_NAME_JSON__;
const DATASET_TOKEN = __DATASET_TOKEN_JSON__;
const TAB_VIEW_TOKENS = __TAB_VIEW_TOKENS_JSON__;
const THEME_SESSION_KEY = 'signalworks_studio_theme_session_mode';
let interactiveOriginalLayout = null;
let interactiveViewReady = false;

// V4.4.5 multi-file CSV selection/binding state.
let selectedCsvFiles = [];
let csvDragSourceIndex = null;
let boundCsvReady = false;
let boundCsvFilename = '';
let boundCsvInfo = null;


function csvFilePickerEl() { return document.getElementById('csvFilePicker'); }
function csvSelectionModalEl() { return document.getElementById('csvSelectionModal'); }

function formatFileBytes(bytes) {
    const n = Number(bytes || 0);
    if (n < 1024) return `${n} B`;
    if (n < 1024*1024) return `${(n/1024).toFixed(1)} KB`;
    if (n < 1024*1024*1024) return `${(n/(1024*1024)).toFixed(1)} MB`;
    return `${(n/(1024*1024*1024)).toFixed(2)} GB`;
}

function formatFileTime(ms) {
    const n = Number(ms || 0);
    if (!Number.isFinite(n) || n <= 0) return 'Time unavailable';
    const d = new Date(n);
    const pad = v => String(v).padStart(2,'0');
    return `${pad(d.getDate())}/${pad(d.getMonth()+1)}/${d.getFullYear()} ${pad(d.getHours())}:${pad(d.getMinutes())}:${pad(d.getSeconds())}`;
}

function setCsvBindProgress(percent, message, kind='normal') {
    const bar = document.getElementById('csvBindProgressBar');
    const text = document.getElementById('csvBindProgressText');
    const p = Math.max(0, Math.min(100, Number(percent || 0)));
    if (bar) bar.style.width = `${p}%`;
    if (text) {
        text.textContent = message || '';
        text.classList.toggle('csv-bind-success', kind === 'success');
        text.classList.toggle('csv-bind-error', kind === 'error');
    }
}

function invalidateBoundCsv(message='File order changed. Bind Data again before loading.') {
    boundCsvReady = false;
    boundCsvFilename = '';
    boundCsvInfo = null;
    const loadBtn = document.getElementById('csvLoadBtn');
    const dlBtn = document.getElementById('csvDownloadBoundBtn');
    if (selectedCsvFiles.length > 1) {
        if (loadBtn) loadBtn.disabled = true;
        if (dlBtn) dlBtn.disabled = true;
        setCsvBindProgress(0, message);
    }
}

function updateCsvModalButtons() {
    const count = selectedCsvFiles.length;
    const bindBtn = document.getElementById('csvBindBtn');
    const loadBtn = document.getElementById('csvLoadBtn');
    const dlBtn = document.getElementById('csvDownloadBoundBtn');
    const heading = document.getElementById('csvSelectionCount');

    if (heading) heading.textContent = `${count} file${count === 1 ? '' : 's'} selected`;
    if (bindBtn) {
        bindBtn.hidden = count <= 1;
        bindBtn.disabled = count <= 1;
    }
    if (loadBtn) {
        loadBtn.disabled = count === 0 || (count > 1 && !boundCsvReady);
        loadBtn.textContent = count > 1 ? 'Load Bound CSV' : 'Load CSV';
    }
    if (dlBtn) {
        dlBtn.hidden = count <= 1;
        dlBtn.disabled = !boundCsvReady;
    }
}

function renderCsvSelectionList() {
    const list = document.getElementById('csvFileList');
    if (!list) return;
    list.replaceChildren();

    selectedCsvFiles.forEach((file, index) => {
        const row = document.createElement('div');
        row.className = 'csv-file-row';
        row.draggable = true;
        row.dataset.index = String(index);

        const handle = document.createElement('div');
        handle.className = 'csv-drag-handle';
        handle.textContent = '⋮⋮';

        const main = document.createElement('div');
        main.className = 'csv-file-main';
        const name = document.createElement('div');
        name.className = 'csv-file-name';
        name.textContent = file.name;
        name.title = file.name;
        const meta = document.createElement('div');
        meta.className = 'csv-file-meta';
        meta.textContent = `${formatFileBytes(file.size)} · Last modified ${formatFileTime(file.lastModified)}`;
        main.append(name, meta);

        const order = document.createElement('div');
        order.className = 'csv-file-order';
        order.textContent = `#${index + 1}`;

        row.append(handle, main, order);

        row.addEventListener('dragstart', event => {
            csvDragSourceIndex = index;
            row.classList.add('dragging');
            try { event.dataTransfer.effectAllowed = 'move'; event.dataTransfer.setData('text/plain', String(index)); } catch (_) {}
        });
        row.addEventListener('dragend', () => {
            csvDragSourceIndex = null;
            document.querySelectorAll('.csv-file-row').forEach(el => el.classList.remove('dragging','drag-over'));
        });
        row.addEventListener('dragover', event => {
            event.preventDefault();
            row.classList.add('drag-over');
            try { event.dataTransfer.dropEffect = 'move'; } catch (_) {}
        });
        row.addEventListener('dragleave', () => row.classList.remove('drag-over'));
        row.addEventListener('drop', event => {
            event.preventDefault();
            row.classList.remove('drag-over');
            const from = csvDragSourceIndex ?? Number(event.dataTransfer?.getData('text/plain'));
            const to = index;
            if (!Number.isInteger(from) || from < 0 || from >= selectedCsvFiles.length || from === to) return;
            const [moved] = selectedCsvFiles.splice(from, 1);
            selectedCsvFiles.splice(to, 0, moved);
            invalidateBoundCsv();
            renderCsvSelectionList();
            updateCsvModalButtons();
        });
        list.appendChild(row);
    });
}

function openCsvSelectionModal() {
    const modal = csvSelectionModalEl();
    if (!modal || !selectedCsvFiles.length) return;
    renderCsvSelectionList();
    updateCsvModalButtons();
    modal.hidden = false;
    document.body.classList.add('modal-open');
}

function closeCsvSelectionModal(resetPicker=false) {
    const modal = csvSelectionModalEl();
    if (modal) modal.hidden = true;
    document.body.classList.remove('modal-open');
    if (resetPicker) {
        selectedCsvFiles = [];
        boundCsvReady = false;
        boundCsvFilename = '';
        boundCsvInfo = null;
        const picker = csvFilePickerEl();
        if (picker) picker.value = '';
    }
}

function handleCsvFileSelection(fileList) {
    selectedCsvFiles = Array.from(fileList || []).filter(file => /\.csv$/i.test(file.name || ''));
    selectedCsvFiles.sort((a,b) => {
        const ta = Number(a.lastModified || 0), tb = Number(b.lastModified || 0);
        if (ta !== tb) return ta - tb;
        return String(a.name || '').localeCompare(String(b.name || ''), undefined, {numeric:true, sensitivity:'base'});
    });
    boundCsvReady = false;
    boundCsvFilename = '';
    boundCsvInfo = null;
    setCsvBindProgress(0, selectedCsvFiles.length > 1 ? 'Review the order, then click Bind Data.' : 'Ready to load this CSV.');
    if (selectedCsvFiles.length) openCsvSelectionModal();
}

function bindSelectedCsvFiles() {
    if (selectedCsvFiles.length < 2) return;
    const bindBtn = document.getElementById('csvBindBtn');
    const loadBtn = document.getElementById('csvLoadBtn');
    const dlBtn = document.getElementById('csvDownloadBoundBtn');
    if (bindBtn) bindBtn.disabled = true;
    if (loadBtn) loadBtn.disabled = true;
    if (dlBtn) dlBtn.disabled = true;
    boundCsvReady = false;

    const form = new FormData();
    selectedCsvFiles.forEach(file => form.append('files', file, file.name));

    const xhr = new XMLHttpRequest();
    xhr.open('POST', '/bind/csv', true);
    xhr.responseType = 'json';
    xhr.upload.onprogress = event => {
        if (!event.lengthComputable) return;
        const percent = 5 + Math.round((event.loaded / event.total) * 65);
        setCsvBindProgress(percent, `Uploading ${formatFileBytes(event.loaded)} / ${formatFileBytes(event.total)}…`);
    };
    xhr.upload.onload = () => setCsvBindProgress(74, 'Upload complete. Validating headers and binding data…');
    xhr.onload = () => {
        const data = xhr.response;
        if (xhr.status >= 200 && xhr.status < 300 && data?.ok) {
            boundCsvReady = true;
            boundCsvFilename = data.filename || 'Bound_Data.csv';
            boundCsvInfo = data;
            setCsvBindProgress(
                100,
                `Bind complete · ${Number(data.total_rows || 0).toLocaleString()} rows · ${formatFileBytes(data.output_bytes)} · ${data.source_files?.length || selectedCsvFiles.length} files`,
                'success'
            );
        } else {
            const message = data?.error || `Bind failed (${xhr.status})`;
            setCsvBindProgress(0, message, 'error');
            boundCsvReady = false;
        }
        if (bindBtn) bindBtn.disabled = false;
        updateCsvModalButtons();
    };
    xhr.onerror = () => {
        setCsvBindProgress(0, 'Network/server error while binding CSV files.', 'error');
        if (bindBtn) bindBtn.disabled = false;
        boundCsvReady = false;
        updateCsvModalButtons();
    };
    setCsvBindProgress(5, 'Preparing files for upload…');
    xhr.send(form);
}

async function loadCsvSelection() {
    if (!selectedCsvFiles.length) return;
    const btn = document.getElementById('csvLoadBtn');
    if (btn) { btn.disabled = true; btn.textContent = 'Loading…'; }

    try {
        if (selectedCsvFiles.length === 1) {
            const form = new FormData();
            form.append('file', selectedCsvFiles[0], selectedCsvFiles[0].name);
            const response = await fetch('/upload', {method:'POST', body:form, cache:'no-store'});
            if (!response.ok) throw new Error(`CSV load failed (${response.status})`);
        } else {
            if (!boundCsvReady) throw new Error('Bind Data before loading multiple files.');
            const response = await fetch('/bind/load', {method:'POST', cache:'no-store'});
            if (!response.ok) {
                let message = '';
                try { message = (await response.json())?.error || ''; } catch (_) {}
                throw new Error(message || `Bound CSV load failed (${response.status})`);
            }
        }
        window.location.href = '/?tab=0';
    } catch (err) {
        alert(`Load CSV error: ${String(err)}`);
        if (btn) { btn.disabled = false; btn.textContent = selectedCsvFiles.length > 1 ? 'Load Bound CSV' : 'Load CSV'; }
    }
}

async function downloadBoundCsv() {
    if (!boundCsvReady) return;
    const filename = boundCsvFilename || 'Bound_Data.csv';
    const target = await chooseNativeSaveHandle(filename, 'text/csv', '.csv');
    if (target.cancelled) return;
    try {
        const blob = await fetchExportBlob('/bind/download');
        await writeBlobToSaveTarget(blob, filename, target);
    } catch (err) {
        alert(`Download error: ${String(err)}`);
    }
}

function initCsvMultiFileUI() {
    const picker = csvFilePickerEl();
    if (!picker) return;
    // If a browser restores a file-input selection after back navigation, clear it
    // so the confirmation workflow always starts from a deliberate new selection.
    try { picker.value = ''; } catch (_) {}
}

function systemPreferredTheme() {
    return (window.matchMedia && window.matchMedia('(prefers-color-scheme: dark)').matches) ? 'dark' : 'light';
}
function currentThemeMode() {
    try {
        const saved = sessionStorage.getItem(THEME_SESSION_KEY);
        if (saved === 'light' || saved === 'dark') return saved;
    } catch (_) {}
    return systemPreferredTheme();
}
function updateThemeToggleUI(mode) {
    const toggle = document.getElementById('themeToggle');
    if (!toggle) return;
    const dark = mode === 'dark';
    toggle.setAttribute('aria-checked', dark ? 'true' : 'false');
    toggle.title = dark ? 'Switch to Light theme' : 'Switch to Dark theme';
}
function plotThemeUpdate() {
    const dark = document.documentElement.dataset.theme === 'dark';
    return {
        'paper_bgcolor': dark ? '#202730' : '#ffffff', 'plot_bgcolor': dark ? '#202730' : '#ffffff',
        'font.color': dark ? '#e6ebf2' : '#344054', 'title.font.color': dark ? '#f4f7fa' : '#172033',
        'xaxis.gridcolor': dark ? '#46515f' : '#e5e9ee', 'xaxis.linecolor': dark ? '#8390a0' : '#555f6b',
        'xaxis.zerolinecolor': dark ? '#586575' : '#c8ced6', 'xaxis.spikecolor': dark ? '#aeb8c5' : '#8b95a1',
        'yaxis.gridcolor': dark ? '#46515f' : '#dfe4ea', 'yaxis.zerolinecolor': dark ? '#586575' : '#c8ced6',
        'legend.bgcolor': dark ? 'rgba(41,50,61,0.94)' : 'rgba(255,255,255,0.88)',
        'legend.bordercolor': dark ? '#79c3ef' : '#1f5a99',
        'legend2.bgcolor': dark ? 'rgba(41,50,61,0.94)' : 'rgba(255,255,255,0.88)',
        'legend2.bordercolor': dark ? '#ff9484' : '#c43c20',
        'legend.title.font.color': dark ? '#79c3ef' : '#1f5a99',
        'legend2.title.font.color': dark ? '#ff9484' : '#c43c20'
    };
}
function positionAxisLegends(plotDiv) {
    if (!plotDiv?.data || !plotDiv.clientHeight) return;
    const hasPrimary = plotDiv.data.some(trace => trace.legend !== 'legend2');
    const hasSecondary = plotDiv.data.some(trace => trace.legend === 'legend2');
    const dual = hasPrimary && hasSecondary;
    plotDiv.parentElement?.classList.toggle('dual-legend', dual);
    const height = plotDiv.clientHeight;
    return Plotly.relayout(plotDiv, {
        'legend.y': (dual ? 190 : 95) / height,
        'legend2.y': 95 / height,
        'margin.b': dual ? 240 : 145
    });
}
function applyThemeToPlot() {
    const plotDiv = plotDivEl();
    if (!plotDiv || typeof Plotly === 'undefined' || !plotDiv.data) return;
    try { Plotly.relayout(plotDiv, plotThemeUpdate()); } catch (_) {}
}
function setThemeMode(mode, persist=true) {
    mode = mode === 'dark' ? 'dark' : 'light';
    if (persist) { try { sessionStorage.setItem(THEME_SESSION_KEY, mode); } catch (_) {} }
    document.documentElement.dataset.themeMode = mode;
    document.documentElement.dataset.theme = mode;
    updateThemeToggleUI(mode);
    applyThemeToPlot();
}
function toggleThemeMode() {
    setThemeMode(document.documentElement.dataset.theme === 'dark' ? 'light' : 'dark', true);
}

function initExclusiveHeaderMenus() {
    const menus = Array.from(document.querySelectorAll('details.header-menu'));

    menus.forEach(menu => {
        menu.addEventListener('toggle', () => {
            if (!menu.open) return;
            menus.forEach(other => {
                if (other !== menu && other.open) other.removeAttribute('open');
            });
        });
    });

    document.addEventListener('pointerdown', event => {
        if (event.target.closest('.header-menu')) return;
        menus.forEach(menu => { if (menu.open) menu.removeAttribute('open'); });
    });

    document.addEventListener('keydown', event => {
        if (event.key !== 'Escape') return;
        menus.forEach(menu => { if (menu.open) menu.removeAttribute('open'); });
    });
}

function sidebarStateKey(key) { return SIDEBAR_STATE_PREFIX + key; }
function saveSidebarUiState() {
    try {
        document.querySelectorAll('.sidebar details[data-ui-key]').forEach(el => {
            sessionStorage.setItem(sidebarStateKey(el.dataset.uiKey), el.open ? '1' : '0');
        });
        const sidebar = document.querySelector('.sidebar');
        if (sidebar) sessionStorage.setItem(sidebarStateKey('scrollTop'), String(sidebar.scrollTop || 0));
    } catch (_) {}
}
function restoreSidebarUiState() {
    try {
        document.querySelectorAll('.sidebar details[data-ui-key]').forEach(el => {
            const saved = sessionStorage.getItem(sidebarStateKey(el.dataset.uiKey));
            if (saved !== null) el.open = saved === '1';
            el.addEventListener('toggle', saveSidebarUiState);
        });
        const sidebar = document.querySelector('.sidebar');
        if (sidebar) {
            const top = Number(sessionStorage.getItem(sidebarStateKey('scrollTop')) || 0);
            requestAnimationFrame(() => { sidebar.scrollTop = Number.isFinite(top) ? top : 0; });
            sidebar.addEventListener('scroll', () => {
                clearTimeout(sidebar._gpScrollSaveTimer);
                sidebar._gpScrollSaveTimer = setTimeout(saveSidebarUiState, 80);
            }, {passive:true});
        }
        document.querySelectorAll('.sidebar form').forEach(form => form.addEventListener('submit', saveSidebarUiState));
    } catch (_) {}
}

function initNestedScrollHandoff() {
    const sidebar = document.querySelector('.sidebar');
    if (!sidebar) return;
    document.querySelectorAll('.signal-list').forEach(inner => {
        inner.addEventListener('wheel', event => {
            const atTop = inner.scrollTop <= 0;
            const atBottom = inner.scrollTop + inner.clientHeight >= inner.scrollHeight - 1;
            if ((event.deltaY < 0 && atTop) || (event.deltaY > 0 && atBottom)) {
                event.preventDefault();
                sidebar.scrollTop += event.deltaY;
            }
        }, {passive:false});
    });
}

if (window.matchMedia) {
    const systemThemeQuery = window.matchMedia('(prefers-color-scheme: dark)');
    const onSystemThemeChange = () => {
        let manual = false;
        try { const saved = sessionStorage.getItem(THEME_SESSION_KEY); manual = saved === 'light' || saved === 'dark'; } catch (_) {}
        if (!manual) setThemeMode(systemPreferredTheme(), false);
    };
    if (systemThemeQuery.addEventListener) systemThemeQuery.addEventListener('change', onSystemThemeChange);
    else if (systemThemeQuery.addListener) systemThemeQuery.addListener(onSystemThemeChange);
}
let cursorMode = null;
let cursorA = null;
let cursorB = null;

// V4.4 cursor/annotation interaction state.
// These variables MUST be declared before the Plotly click handler reads them.
// Without the declarations, the first graph click throws
// "ReferenceError: commentMode is not defined" and Cursor A/B never receives it.
let commentMode = false;
let pendingCommentPoint = null;
let engineeringBaseShapes = [];
let engineeringBaseAnnotations = [];

let lastCycleAnalysis = null;
let relayoutTimer = null;
let engineeringShapeMap = [];
let engineeringAnnotationMap = [];
let annotationDragSyncTimer = null;
let annotationRelayoutGuard = false;
const SIDEBAR_STATE_PREFIX = `signalworks_studio_sidebar_tab_${ACTIVE_TAB}_`;

function plotDivEl() { return document.getElementById('interactivePlot'); }

function isDateLike(value) {
    if (typeof value !== 'string') return false;
    return /^\d{4}-\d{2}-\d{2}/.test(value) || (!Number.isFinite(Number(value)) && Number.isFinite(Date.parse(value)));
}

function xComparable(value) {
    if (value === null || value === undefined) return NaN;
    if (typeof value === 'number') return value;
    if (isDateLike(value)) return Date.parse(value);
    const n = Number(value);
    if (Number.isFinite(n)) return n;
    const d = Date.parse(value);
    return Number.isFinite(d) ? d : NaN;
}

function formatX(value) {
    if (value === null || value === undefined) return '—';
    if (isDateLike(value)) {
        const d = new Date(value);
        if (!Number.isNaN(d.getTime())) {
            const pad = n => String(n).padStart(2, '0');
            return `${pad(d.getDate())}/${pad(d.getMonth()+1)}/${d.getFullYear()} ${pad(d.getHours())}:${pad(d.getMinutes())}:${pad(d.getSeconds())}`;
        }
    }
    if (typeof value === 'number') return Number.isInteger(value) ? String(value) : value.toPrecision(7).replace(/\.?0+$/, '');
    return String(value);
}

function escapeHtml(value) {
    return String(value ?? '').replace(/[&<>\"']/g, ch => ({'&':'&amp;','<':'&lt;','>':'&gt;','\"':'&quot;',"'":'&#39;'}[ch]));
}

function formatNumber(value) {
    const n = Number(value);
    if (!Number.isFinite(n)) return '—';
    if (n === 0) return '0';
    const a = Math.abs(n);
    if (a >= 100000 || a < 0.001) return n.toExponential(4);
    return Number(n.toPrecision(7)).toString();
}

function formatDuration(delta) {
    if (!Number.isFinite(delta)) return '—';
    const ms = Math.abs(delta);
    const sign = delta < 0 ? '-' : '';
    if (ms < 1000) return `${sign}${ms.toFixed(0)} ms`;
    const totalSec = ms / 1000;
    if (totalSec < 60) return `${sign}${totalSec.toFixed(3).replace(/\.?0+$/, '')} s`;
    const totalMin = totalSec / 60;
    if (totalMin < 60) return `${sign}${Math.floor(totalMin)} min ${(totalSec % 60).toFixed(1)} s`;
    const totalHr = totalMin / 60;
    if (totalHr < 24) return `${sign}${Math.floor(totalHr)} h ${Math.floor(totalMin % 60)} min ${(totalSec % 60).toFixed(0)} s`;
    return `${sign}${Math.floor(totalHr / 24)} d ${Math.floor(totalHr % 24)} h ${Math.floor(totalMin % 60)} min`;
}

function traceIsVisible(trace) {
    return trace && trace.visible !== false && trace.visible !== 'legendonly';
}

function nearestIndex(xs, target) {
    if (!xs || xs.length === 0 || target === null || target === undefined) return -1;
    const t = xComparable(target);
    if (!Number.isFinite(t)) return -1;
    let lo = 0, hi = xs.length - 1;
    let first = xComparable(xs[0]);
    let last = xComparable(xs[hi]);
    if (!Number.isFinite(first) || !Number.isFinite(last) || first > last) {
        let best = -1, bestDiff = Infinity;
        for (let i = 0; i < xs.length; i++) {
            const xv = xComparable(xs[i]);
            if (!Number.isFinite(xv)) continue;
            const diff = Math.abs(xv - t);
            if (diff < bestDiff) { best = i; bestDiff = diff; }
        }
        return best;
    }
    while (lo <= hi) {
        const mid = (lo + hi) >> 1;
        const mv = xComparable(xs[mid]);
        if (!Number.isFinite(mv)) break;
        if (mv < t) lo = mid + 1;
        else if (mv > t) hi = mid - 1;
        else return mid;
    }
    const candidates = [Math.max(0, Math.min(xs.length - 1, lo)), Math.max(0, Math.min(xs.length - 1, hi))];
    let best = candidates[0], bestDiff = Infinity;
    for (const idx of candidates) {
        const diff = Math.abs(xComparable(xs[idx]) - t);
        if (diff < bestDiff) { best = idx; bestDiff = diff; }
    }
    return best;
}

function valueAtCursor(trace, cursorX) {
    if (!trace || cursorX === null) return null;
    const idx = nearestIndex(trace.x || [], cursorX);
    if (idx < 0 || !trace.y || idx >= trace.y.length) return null;
    const v = Number(trace.y[idx]);
    return Number.isFinite(v) ? v : null;
}

function currentXRange(plotDiv) {
    if (!plotDiv || !plotDiv._fullLayout || !plotDiv._fullLayout.xaxis) return [null, null];
    const r = plotDiv._fullLayout.xaxis.range;
    if (!r || r.length < 2) return [null, null];
    return [r[0], r[1]];
}

function analysisRange(plotDiv) {
    if (cursorA !== null && cursorB !== null) {
        const a = xComparable(cursorA), b = xComparable(cursorB);
        return [a <= b ? cursorA : cursorB, a <= b ? cursorB : cursorA, 'Cursor A ↔ B'];
    }
    const r = currentXRange(plotDiv);
    return [r[0], r[1], 'Current view'];
}

function statisticsForTrace(trace, x0, x1) {
    if (!trace || !trace.x || !trace.y) return {min:null,max:null,avg:null,count:0};
    let a = xComparable(x0), b = xComparable(x1);
    if (Number.isFinite(a) && Number.isFinite(b) && a > b) [a,b] = [b,a];
    let min = Infinity, max = -Infinity, sum = 0, count = 0;
    for (let i = 0; i < trace.y.length; i++) {
        const y = Number(trace.y[i]);
        if (!Number.isFinite(y)) continue;
        const x = xComparable(trace.x[i]);
        if (Number.isFinite(a) && Number.isFinite(x) && x < a) continue;
        if (Number.isFinite(b) && Number.isFinite(x) && x > b) continue;
        if (y < min) min = y;
        if (y > max) max = y;
        sum += y; count++;
    }
    return count ? {min, max, avg:sum/count, count} : {min:null,max:null,avg:null,count:0};
}

function armCursor(which) {
    // Cursor and comment placement are mutually exclusive.
    commentMode = false;
    pendingCommentPoint = null;
    cursorMode = which;
    document.getElementById('cursorABtn')?.classList.toggle('armed', which === 'A');
    document.getElementById('cursorBBtn')?.classList.toggle('armed', which === 'B');
    const plotDiv = plotDivEl();
    if (plotDiv) plotDiv.style.cursor = 'crosshair';
    const msg = document.getElementById('cursorInstruction');
    if (msg) msg.textContent = `Cursor ${which} armed — click a signal point on the graph.`;
}

function clearCursorArmedState() {
    cursorMode = null;
    document.getElementById('cursorABtn')?.classList.remove('armed');
    document.getElementById('cursorBBtn')?.classList.remove('armed');
    const plotDiv = plotDivEl();
    if (plotDiv && !commentMode) plotDiv.style.cursor = '';
}

function setCursor(which, xValue) {
    if (which === 'A') cursorA = xValue;
    else cursorB = xValue;
    clearCursorArmedState();
    updateCursorShapes();
    updateAnalysis();
    const msg = document.getElementById('cursorInstruction');
    if (msg) msg.textContent = `Cursor ${which} set at ${formatX(xValue)}.`;
}

function clearCursors() {
    cursorA = null; cursorB = null; clearCursorArmedState();
    updateCursorShapes(); updateAnalysis();
    const msg = document.getElementById('cursorInstruction');
    if (msg) msg.textContent = 'Choose Cursor A or B, then click a point on the graph.';
}

function updateCursorShapes() {
    const plotDiv = plotDivEl();
    if (!plotDiv) return;
    const shapes = JSON.parse(JSON.stringify(engineeringBaseShapes || []));
    const annotations = JSON.parse(JSON.stringify(engineeringBaseAnnotations || []));
    if (cursorA !== null && cursorB !== null) {
        const a = xComparable(cursorA), b = xComparable(cursorB);
        const x0 = a <= b ? cursorA : cursorB;
        const x1 = a <= b ? cursorB : cursorA;
        shapes.push({type:'rect', xref:'x', yref:'paper', x0, x1, y0:0, y1:1, fillcolor:'rgba(23,105,170,.055)', line:{width:0}, layer:'below', editable:false});
    }
    if (cursorA !== null) {
        shapes.push({type:'line', xref:'x', yref:'paper', x0:cursorA, x1:cursorA, y0:0, y1:1, line:{color:'#2e90fa', width:1.6, dash:'dash'}, editable:false});
        annotations.push({xref:'x', yref:'paper', x:cursorA, y:1, text:'A', showarrow:false, yshift:10, bgcolor:'#eff8ff', bordercolor:'#84adff', borderwidth:1, font:{size:10,color:'#175cd3'}});
    }
    if (cursorB !== null) {
        shapes.push({type:'line', xref:'x', yref:'paper', x0:cursorB, x1:cursorB, y0:0, y1:1, line:{color:'#f04438', width:1.6, dash:'dash'}, editable:false});
        annotations.push({xref:'x', yref:'paper', x:cursorB, y:1, text:'B', showarrow:false, yshift:10, bgcolor:'#fef3f2', bordercolor:'#fda29b', borderwidth:1, font:{size:10,color:'#b42318'}});
    }
    Plotly.relayout(plotDiv, {shapes, annotations});
}

function updateAnalysis() {
    const plotDiv = plotDivEl();
    if (!plotDiv || !plotDiv.data) return;

    document.getElementById('cursorAValue').textContent = formatX(cursorA);
    document.getElementById('cursorBValue').textContent = formatX(cursorB);

    const datetimeAxis = plotDiv.layout?.xaxis?.type === 'date' || (plotDiv.data[0]?.x?.length && isDateLike(plotDiv.data[0].x[0]));
    const deltaLabel = document.getElementById('deltaXLabel');
    const deltaValue = document.getElementById('deltaXValue');
    if (datetimeAxis) {
        if (deltaLabel) deltaLabel.textContent = 'Δ Time';
        if (deltaValue) deltaValue.textContent = (cursorA !== null && cursorB !== null) ? formatDuration(xComparable(cursorB) - xComparable(cursorA)) : '—';
    } else {
        if (deltaLabel) deltaLabel.textContent = 'Δ X';
        if (deltaValue) deltaValue.textContent = (cursorA !== null && cursorB !== null) ? formatNumber(xComparable(cursorB) - xComparable(cursorA)) : '—';
    }

    const [x0, x1, mode] = analysisRange(plotDiv);
    const modeEl = document.getElementById('analysisModeValue');
    if (modeEl) modeEl.textContent = mode;
    const rangeText = document.getElementById('analysisRangeText');
    if (rangeText) rangeText.textContent = `${mode}: ${formatX(x0)} → ${formatX(x1)}`;

    const tbody = document.getElementById('analysisTableBody');
    if (!tbody) return;
    const rows = [];
    plotDiv.data.forEach(trace => {
        if (!traceIsVisible(trace)) return;
        const aVal = valueAtCursor(trace, cursorA);
        const bVal = valueAtCursor(trace, cursorB);
        const delta = (aVal !== null && bVal !== null) ? bVal - aVal : null;
        const s = statisticsForTrace(trace, x0, x1);
        const isSecondary = trace.yaxis === 'y2';
        rows.push(`
            <tr>
                <td title="${escapeHtml(trace.name)}">${escapeHtml(trace.name)}</td>
                <td><span class="axis-pill ${isSecondary ? 'secondary' : 'primary'}">${isSecondary ? 'Secondary' : 'Primary'}</span></td>
                <td>${formatNumber(aVal)}</td>
                <td>${formatNumber(bVal)}</td>
                <td>${formatNumber(delta)}</td>
                <td>${formatNumber(s.min)}</td>
                <td>${formatNumber(s.max)}</td>
                <td>${formatNumber(s.avg)}</td>
                <td>${s.count}</td>
            </tr>`);
    });
    tbody.innerHTML = rows.length ? rows.join('') : '<tr><td colspan="9" class="empty-table">No visible signal traces.</td></tr>';
}

function setCycleStatus(message, kind='') {
    const el = document.getElementById('cycleStatus');
    if (!el) return;
    el.className = 'cycle-status' + (kind ? ' ' + kind : '');
    el.textContent = message;
}

function setCycleMetric(id, value) {
    const el = document.getElementById(id);
    if (el) el.textContent = value ?? '—';
}

async function runCycleAnalysis() {
    const signalEl = document.getElementById('cycleSignal');
    const signal = signalEl?.value || '';
    if (!signal) { setCycleStatus('Select a plotted signal first.', 'warning'); return; }

    const profileMode = document.getElementById('cycleProfileMode')?.value || 'auto';
    const tolerance = Number(document.getElementById('cycleTolerance')?.value || 2);
    const sampleTime = Number(document.getElementById('cycleSampleTime')?.value || 10);
    if (!Number.isFinite(tolerance) || tolerance <= 0) { setCycleStatus('Tolerance must be greater than 0.', 'error'); return; }
    if (!Number.isFinite(sampleTime) || sampleTime <= 0) { setCycleStatus('Fallback sample time must be greater than 0.', 'error'); return; }

    const btn = document.getElementById('cycleAnalyzeBtn');
    if (btn) { btn.disabled = true; btn.textContent = 'Analyzing…'; }
    setCycleStatus(`Analyzing ${signal}…`);
    try {
        const response = await fetch(`/api/cycle-analysis/${ACTIVE_TAB}`, {
            method:'POST',
            headers:{'Content-Type':'application/json'},
            body:JSON.stringify({signal, profile_mode: profileMode, tolerance, sample_time: sampleTime}),
            cache:'no-store'
        });
        const data = await response.json();
        if (!response.ok) throw new Error(data?.detail || data?.error || JSON.stringify(data));
        lastCycleAnalysis = data;

        setCycleMetric('cycleCount', String(data.cycles ?? 0));
        setCycleMetric('cycleTmin', Number.isFinite(Number(data.tmin)) ? `${formatNumber(data.tmin)} °C` : '—');
        setCycleMetric('cycleTmax', Number.isFinite(Number(data.tmax)) ? `${formatNumber(data.tmax)} °C` : '—');
        setCycleMetric('cycleAvgTime', data.avg_cycle_time || '—');
        setCycleMetric('cycleAvgSoakMin', data.avg_soak_tmin || '—');
        setCycleMetric('cycleAvgSoakMax', data.avg_soak_tmax || '—');

        const tbody = document.getElementById('cycleTableBody');
        if (tbody) {
            if (data.details?.length) {
                tbody.innerHTML = data.details.map(row => `<tr>
                    <td>${row.cycle}</td>
                    <td>${escapeHtml(row.start_state)}</td>
                    <td>${escapeHtml(row.cycle_start)}</td>
                    <td>${escapeHtml(row.cycle_end)}</td>
                    <td>${escapeHtml(row.soak_tmin)}</td>
                    <td>${escapeHtml(row.soak_tmax)}</td>
                    <td>${escapeHtml(row.cycle_time)}</td>
                </tr>`).join('');
            } else {
                tbody.innerHTML = '<tr><td colspan="7" class="empty-table">No complete cycles detected for this signal.</td></tr>';
            }
        }
        const source = document.getElementById('cycleTimeSource');
        if (source) {
            const observed = `Observed min/max: ${formatNumber(data.observed_tmin)} / ${formatNumber(data.observed_tmax)} °C`;
            const analysisLevels = `Analysis Tmin/Tmax: ${formatNumber(data.tmin)} / ${formatNumber(data.tmax)} °C`;
            const transition = data.median_transition_seconds == null ? 'Transition: —' : `Median transition: ${data.median_transition}`;
            const mode = data.profile_mode_requested === 'auto'
                ? `Profile: Auto → ${data.detected_profile_display || data.profile_mode_display}`
                : `Profile: ${data.profile_mode_display}`;
            source.textContent = `${mode}. ${observed}. ${analysisLevels}. ${transition}. Tolerance: ±${data.tolerance} °C. Tmin soaks: ${data.tmin_soak_count}; Tmax soaks: ${data.tmax_soak_count}.`;
        }
        const exportBtn = document.getElementById('cycleExportBtn');
        if (exportBtn) exportBtn.disabled = false;

        if ((data.cycles || 0) > 0) {
            setCycleStatus(`Cycle pattern detected: ${data.cycles} complete cycle(s) in ${data.signal}. ${data.note || ''}`.trim(), 'success');
        } else {
            setCycleStatus(data.note || 'No complete cycle pattern detected.', 'warning');
        }
    } catch (err) {
        lastCycleAnalysis = null;
        setCycleStatus(`Cycle analysis error: ${String(err)}`, 'error');
        const exportBtn = document.getElementById('cycleExportBtn');
        if (exportBtn) exportBtn.disabled = true;
    } finally {
        if (btn) { btn.disabled = false; btn.textContent = 'Analyze cycles'; }
    }
}

function csvEscape(value) {
    const text = String(value ?? '');
    return '"' + text.replaceAll('"', '""') + '"';
}

function buildCycleAnalysisCsvBlob() {
    const d = lastCycleAnalysis;
    if (!d) return null;
    const columns = ['Signal','Profile Requested','Profile Used','Detected Profile','Observed Tmin','Observed Tmax','Analysis Tmin','Analysis Tmax','Plateau Tmin','Plateau Tmax','Median Transition (seconds)','Transition Rate (°C/min)','Tolerance','Time Source','Complete Cycles','Cycle','Cycle Start State','Cycle Start','Cycle End','Soak Tmin (seconds)','Soak Tmax (seconds)','Cycle Time (seconds)','Average Soak Tmin (seconds)','Average Soak Tmax (seconds)','Average Cycle Time (seconds)'];
    const rows = [];
    if (d.details?.length) {
        d.details.forEach(r => rows.push([
            d.signal,d.profile_mode_requested,d.profile_mode_display,d.detected_profile_display,
            d.observed_tmin,d.observed_tmax,d.tmin,d.tmax,d.plateau_tmin,d.plateau_tmax,
            d.median_transition_seconds,d.transition_rate_c_per_min,d.tolerance,d.time_source,d.cycles,
            r.cycle,r.start_state,r.cycle_start,r.cycle_end,
            r.soak_tmin_seconds,r.soak_tmax_seconds,r.cycle_time_seconds,
            d.avg_soak_tmin_seconds,d.avg_soak_tmax_seconds,d.avg_cycle_time_seconds
        ]));
    } else {
        rows.push([
            d.signal,d.profile_mode_requested,d.profile_mode_display,d.detected_profile_display,
            d.observed_tmin,d.observed_tmax,d.tmin,d.tmax,d.plateau_tmin,d.plateau_tmax,
            d.median_transition_seconds,d.transition_rate_c_per_min,d.tolerance,d.time_source,d.cycles,
            '','','','','','','',d.avg_soak_tmin_seconds,d.avg_soak_tmax_seconds,d.avg_cycle_time_seconds
        ]);
    }
    const text = [columns, ...rows].map(row => row.map(csvEscape).join(',')).join('\r\n');
    return new Blob(['\ufeff' + text], {type:'text/csv;charset=utf-8'});
}

function exportCycleAnalysisCSV() {
    const d = lastCycleAnalysis;
    if (!d) { setCycleStatus('Run Cycle Analysis before exporting.', 'warning'); return; }
    const safeSignal = String(d.signal || 'signal').replace(/[^a-zA-Z0-9_-]+/g, '_');
    openExportDialog({
        mode:'cycle',
        title:'Export Cycle Analysis CSV',
        filename:`${EXPORT_BASE}_Cycle_Analysis_${safeSignal}.csv`,
        extension:'.csv',
        mime:'text/csv',
        showTabs:false,
        showPresetName:false,
        help:'Save the current cycle-analysis result as CSV.'
    });
}


let exportDialogContext = null;

function ensureExtension(filename, extension) {
    let name = String(filename || '').trim();
    if (!name) name = `SignalWorksStudio${extension}`;
    if (!name.toLowerCase().endsWith(extension.toLowerCase())) name += extension;
    return name.replace(/[\\/:*?"<>|]+/g, '_');
}

function openExportDialog(ctx) {
    exportDialogContext = {...ctx};
    const modal = document.getElementById('exportModal');
    if (!modal) return;
    const title = document.getElementById('exportModalTitle');
    const help = document.getElementById('exportModalHelp');
    const filename = document.getElementById('exportFileName');
    const tabSection = document.getElementById('exportTabSection');
    const presetRow = document.getElementById('exportPresetNameRow');
    const presetName = document.getElementById('exportPresetName');
    if (title) title.textContent = ctx.title || 'Export';
    if (help) help.textContent = ctx.help || 'Choose the output settings and file name.';
    if (filename) filename.value = ctx.filename || `SignalWorksStudio${ctx.extension || ''}`;
    if (tabSection) tabSection.hidden = !ctx.showTabs;
    if (presetRow) presetRow.hidden = !ctx.showPresetName;
    if (presetName && ctx.showPresetName) presetName.value = ctx.presetName || ACTIVE_PRESET_NAME || 'Signal Preset';
    modal.hidden = false;
    document.body.classList.add('modal-open');
    try { filename?.focus(); filename?.select(); } catch (_) {}
}

function closeExportModal() {
    const modal = document.getElementById('exportModal');
    if (modal) modal.hidden = true;
    document.body.classList.remove('modal-open');
    exportDialogContext = null;
}

function setAllExportTabs(checked) {
    document.querySelectorAll('#exportTabGrid input[type="checkbox"]').forEach(cb => cb.checked = !!checked);
}

function selectedExportTabs() {
    return Array.from(document.querySelectorAll('#exportTabGrid input[type="checkbox"]:checked')).map(cb => Number(cb.value));
}

function openReportExportDialog(kind) {
    const map = {
        wide: {title:'Export Wide PDF Report', extension:'.pdf', mime:'application/pdf', endpoint:'/report/pdf', suffix:'Test_Report_Wide', help:'Select only the plot pages you want in the custom-wide PDF report.'},
        a4: {title:'Export A4 PDF Report', extension:'.pdf', mime:'application/pdf', endpoint:'/report/pdf-a4', suffix:'Test_Report_A4', help:'Select only the plot pages you want in the A4 landscape PDF report.'},
        word: {title:'Export Word Report', extension:'.docx', mime:'application/vnd.openxmlformats-officedocument.wordprocessingml.document', endpoint:'/report/word', suffix:'Test_Report', help:'Select only the plot pages you want in the Word report.'}
    };
    const cfg = map[kind]; if (!cfg) return;
    openExportDialog({mode:'report', kind, ...cfg, filename:`${EXPORT_BASE}_${cfg.suffix}${cfg.extension}`, showTabs:true, showPresetName:false});
}

function openPresetExportDialog() {
    const preset = ACTIVE_PRESET_NAME || 'Signal Preset';
    const safe = preset.replace(/[^a-zA-Z0-9_-]+/g, '_').replace(/^_+|_+$/g,'') || 'Signal_Preset';
    openExportDialog({mode:'preset', title:'Export Signal Preset', extension:'.preset', mime:'application/json', filename:`${safe}.preset`, showTabs:false, showPresetName:true, presetName:preset, help:'Set the preset name and file name, then choose where to save the portable .preset file.'});
}

function openPlotExportDialog(tabId) {
    openExportDialog({mode:'plot', title:`Save Plot ${Number(tabId)+1} as PNG`, extension:'.png', mime:'image/png', filename:`${EXPORT_BASE}_Tab${Number(tabId)+1}.png`, showTabs:false, showPresetName:false, tabId:Number(tabId), help:'Choose a file name and save location for this plot image.'});
}

function filePickerType(mime, extension) {
    const description = extension === '.pdf' ? 'PDF document' : extension === '.docx' ? 'Word document' : extension === '.png' ? 'PNG image' : extension === '.csv' ? 'CSV file' : extension === '.preset' ? 'SignalWorks Studio preset' : 'File';
    return [{description, accept:{[mime || 'application/octet-stream']:[extension]}}];
}

async function chooseNativeSaveHandle(filename, mime, extension) {
    if (!window.showSaveFilePicker || !window.isSecureContext) return {fallback:true};
    try {
        const handle = await window.showSaveFilePicker({suggestedName:filename, types:filePickerType(mime, extension)});
        return {handle};
    } catch (err) {
        if (err && err.name === 'AbortError') return {cancelled:true};
        console.warn('Native Save As unavailable; falling back to browser download.', err);
        return {fallback:true};
    }
}

async function writeBlobToSaveTarget(blob, filename, target) {
    if (target?.cancelled) return false;
    if (target?.handle) {
        const writable = await target.handle.createWritable();
        await writable.write(blob);
        await writable.close();
        return true;
    }
    const url = URL.createObjectURL(blob);
    const a = document.createElement('a');
    a.href = url; a.download = filename;
    document.body.appendChild(a); a.click(); a.remove();
    setTimeout(() => URL.revokeObjectURL(url), 1500);
    return true;
}

async function fetchExportBlob(url) {
    const response = await fetch(url, {cache:'no-store'});
    if (!response.ok) throw new Error(await response.text() || `Export failed (${response.status})`);
    return await response.blob();
}

async function confirmExportDialog() {
    const ctx = exportDialogContext;
    if (!ctx) return;
    const filenameEl = document.getElementById('exportFileName');
    const presetNameEl = document.getElementById('exportPresetName');
    const confirmBtn = document.getElementById('exportConfirmBtn');
    const filename = ensureExtension(filenameEl?.value, ctx.extension || '');
    if (filenameEl) filenameEl.value = filename;
    let tabs = [];
    if (ctx.showTabs) {
        tabs = selectedExportTabs();
        if (!tabs.length) { alert('Select at least one report page.'); return; }
    }

    // Native Save As must be requested directly from this click before long-running report generation.
    const target = await chooseNativeSaveHandle(filename, ctx.mime, ctx.extension);
    if (target.cancelled) return;
    if (confirmBtn) { confirmBtn.disabled = true; confirmBtn.textContent = 'Preparing…'; }
    try {
        let blob;
        if (ctx.mode === 'report') {
            const views = exportViewStates(tabs);
            const url = `${ctx.endpoint}?tabs=${encodeURIComponent(tabs.join(','))}&views=${encodeURIComponent(JSON.stringify(views))}`;
            blob = await fetchExportBlob(url);
        } else if (ctx.mode === 'preset') {
            const presetName = String(presetNameEl?.value || 'Signal Preset').trim() || 'Signal Preset';
            blob = await fetchExportBlob(`/preset/export?name=${encodeURIComponent(presetName)}`);
        } else if (ctx.mode === 'plot') {
            const view = exportViewStates([ctx.tabId])[String(ctx.tabId)] || {};
            blob = await fetchExportBlob(`/download/plot/${ctx.tabId}?view=${encodeURIComponent(JSON.stringify(view))}`);
        } else if (ctx.mode === 'cycle') {
            blob = buildCycleAnalysisCsvBlob();
            if (!blob) throw new Error('Run Cycle Analysis before exporting.');
        } else {
            throw new Error('Unknown export type.');
        }
        const saved = await writeBlobToSaveTarget(blob, filename, target);
        if (saved) closeExportModal();
    } catch (err) {
        alert(`Export error: ${String(err)}`);
    } finally {
        if (confirmBtn) { confirmBtn.disabled = false; confirmBtn.textContent = 'Save As…'; }
    }
}

function plotViewStorageKey(tabId) {
    return `signalworks_studio_plot_view_${DATASET_TOKEN || 'no-data'}_tab_${Number(tabId)}_${TAB_VIEW_TOKENS[tabId] || ''}`;
}

function readStoredPlotView(tabId) {
    try {
        const raw = sessionStorage.getItem(plotViewStorageKey(tabId));
        if (!raw) return null;
        const parsed = JSON.parse(raw);
        return parsed && typeof parsed === 'object' ? parsed : null;
    } catch (_) { return null; }
}

function currentPlotViewState(plotDiv) {
    if (!plotDiv?.layout) return null;
    const axisRange = axis => {
        if (!axis || !Array.isArray(axis.range) || axis.range.length !== 2) return null;
        return [axis.range[0], axis.range[1]];
    };
    const visibleSignals = {};
    (plotDiv.data || []).forEach(trace => {
        if (!trace?.name) return;
        visibleSignals[String(trace.name)] = trace.visible !== false && trace.visible !== 'legendonly';
    });
    return {
        xaxis_range: axisRange(plotDiv.layout.xaxis),
        yaxis_range: axisRange(plotDiv.layout.yaxis),
        yaxis2_range: axisRange(plotDiv.layout.yaxis2),
        dragmode: String(plotDiv.layout.dragmode || 'zoom'),
        visible_signals: visibleSignals
    };
}

function savePlotViewState(tabId=ACTIVE_TAB) {
    const plotDiv = plotDivEl();
    if (!interactiveViewReady || !plotDiv || Number(tabId) !== Number(ACTIVE_TAB)) return;
    const state = currentPlotViewState(plotDiv);
    if (!state) return;
    try { sessionStorage.setItem(plotViewStorageKey(tabId), JSON.stringify(state)); } catch (_) {}
}

async function restorePlotViewState(plotDiv, tabId=ACTIVE_TAB) {
    const state = readStoredPlotView(tabId);
    if (!state || !plotDiv) return;
    const relayout = {};
    if (Array.isArray(state.xaxis_range) && state.xaxis_range.length === 2) {
        relayout['xaxis.autorange'] = false; relayout['xaxis.range'] = state.xaxis_range;
    }
    if (Array.isArray(state.yaxis_range) && state.yaxis_range.length === 2) {
        relayout['yaxis.autorange'] = false; relayout['yaxis.range'] = state.yaxis_range;
    }
    if (Array.isArray(state.yaxis2_range) && state.yaxis2_range.length === 2) {
        relayout['yaxis2.autorange'] = false; relayout['yaxis2.range'] = state.yaxis2_range;
    }
    if (['zoom','pan'].includes(state.dragmode)) relayout.dragmode = state.dragmode;
    if (Object.keys(relayout).length) await Plotly.relayout(plotDiv, relayout);
    if (state.visible_signals && typeof state.visible_signals === 'object') {
        const vis = (plotDiv.data || []).map(trace => state.visible_signals[String(trace.name)] === false ? 'legendonly' : true);
        if (vis.length) await Plotly.restyle(plotDiv, {visible: vis});
    }
}

function exportViewStates(tabIds) {
    savePlotViewState(ACTIVE_TAB);
    const out = {};
    (tabIds || []).forEach(id => {
        const state = readStoredPlotView(id);
        if (state) out[String(id)] = state;
    });
    return out;
}

function setStaticPlotHref(link) {
    const view = exportViewStates([ACTIVE_TAB])[String(ACTIVE_TAB)] || {};
    link.href = `/plot/${ACTIVE_TAB}.png?view=${encodeURIComponent(JSON.stringify(view))}`;
}

function scheduleAnalysisUpdate() {
    clearTimeout(relayoutTimer);
    relayoutTimer = setTimeout(updateAnalysis, 120);
}

async function loadInteractivePlot() {
    interactiveViewReady = false;
    const plotDiv = plotDivEl();
    if (!plotDiv) return;
    if (typeof Plotly === 'undefined') {
        plotDiv.innerHTML = '<div class="plot-error">Plotly runtime could not be loaded. Refresh the page or restart the WebServer.</div>';
        return;
    }
    try {
        const response = await fetch(`/api/plot/${ACTIVE_TAB}?v=${ACTIVE_REVISION}`, {cache:'no-store'});
        if (!response.ok) throw new Error(await response.text());
        const fig = await response.json();
        interactiveOriginalLayout = JSON.parse(JSON.stringify(fig.layout || {}));
        const config = {
            responsive:true,
            scrollZoom:true,
            displaylogo:false,
            editable:true,
            edits:{
                shapePosition:true, annotationPosition:true, annotationTail:true,
                axisTitleText:false, colorbarPosition:false, colorbarTitleText:false,
                legendPosition:false, legendText:false, titleText:false
            },
            doubleClick:'reset+autosize',
            modeBarButtonsToRemove:['lasso2d','select2d'],
            toImageButtonOptions:{format:'png', filename:'CSV_Plot_Interactive_Tab'+(ACTIVE_TAB+1), scale:2}
        };
        // Remove the loading placeholder before Plotly creates its own DOM.
        // Without this, Plotly appends beside the placeholder and the spinner keeps
        // occupying a full graph-height area underneath the rendered graph.
        try { Plotly.purge(plotDiv); } catch (_) {}
        plotDiv.replaceChildren();

        await Plotly.newPlot(plotDiv, fig.data || [], fig.layout || {}, config);
        engineeringBaseShapes = JSON.parse(JSON.stringify(fig.layout?.shapes || []));
        engineeringBaseAnnotations = JSON.parse(JSON.stringify(fig.layout?.annotations || []));
        engineeringShapeMap = JSON.parse(JSON.stringify(fig.layout?.meta?.signalworks_shape_map || fig.layout?.meta?.graphplot_shape_map || []));
        engineeringAnnotationMap = JSON.parse(JSON.stringify(fig.layout?.meta?.signalworks_annotation_map || fig.layout?.meta?.graphplot_annotation_map || []));
        await Plotly.relayout(plotDiv, plotThemeUpdate());
        await positionAxisLegends(plotDiv);
        await restorePlotViewState(plotDiv, ACTIVE_TAB);
        interactiveViewReady = true;
        requestAnimationFrame(() => {
            try { Plotly.Plots.resize(plotDiv); } catch (_) {}
        });

        plotDiv.on('plotly_click', evt => {
            if (!evt?.points?.length) return;
            const pt = evt.points[0];
            if (commentMode) {
                pendingCommentPoint = {x: pt.x, y: pt.y, axis: pt.data?.yaxis === 'y2' ? 'right' : 'left'};
                commentMode = false;
                openCommentModal();
                return;
            }
            if (cursorMode) setCursor(cursorMode, pt.x);
        });
        plotDiv.on('plotly_relayout', evt => { handleAnnotationRelayout(evt); savePlotViewState(); scheduleAnalysisUpdate(); });
        plotDiv.on('plotly_restyle', () => { savePlotViewState(); scheduleAnalysisUpdate(); });
        plotDiv.on('plotly_legendclick', () => setTimeout(() => { savePlotViewState(); updateAnalysis(); }, 100));
        plotDiv.on('plotly_legenddoubleclick', () => setTimeout(() => { savePlotViewState(); updateAnalysis(); }, 100));
        updateAnalysis();
    } catch (err) {
        plotDiv.innerHTML = `<div class="plot-error">Interactive plot error: ${String(err)}</div>`;
    }
}

function changedIndexedKeys(evt, prefix) {
    const out = new Map();
    Object.entries(evt || {}).forEach(([key, value]) => {
        const m = key.match(new RegExp('^' + prefix + '\\[(\\d+)\\]\\.(.+)$'));
        if (!m) return;
        const idx = Number(m[1]);
        if (!out.has(idx)) out.set(idx, {});
        out.get(idx)[m[2]] = value;
    });
    return out;
}
function updateManagerField(kind, id, field, value) {
    const form = document.querySelector(`form[data-ann-kind="${kind}"][data-ann-id="${CSS.escape(String(id))}"]`);
    if (!form) return;
    const input = form.querySelector(`[name="${field}"]`);
    if (input) input.value = value;
    const details = form.closest('.annotation-item');
    const summaryValue = details?.querySelector('.ann-value');
    if (summaryValue && kind === 'line' && field === 'value') summaryValue.textContent = formatNumber(value);
    if (summaryValue && kind === 'band' && (field === 'min' || field === 'max')) {
        const lo = form.querySelector('[name="min"]')?.value ?? '';
        const hi = form.querySelector('[name="max"]')?.value ?? '';
        summaryValue.textContent = `${formatNumber(lo)}…${formatNumber(hi)}`;
    }
}
async function syncDraggedAnnotation(payload) {
    try {
        const r = await fetch(`/api/annotation/position/${ACTIVE_TAB}`, {
            method:'POST', headers:{'Content-Type':'application/json'}, body:JSON.stringify(payload)
        });
        if (!r.ok) throw new Error(await r.text());
    } catch (err) {
        console.warn('Annotation drag sync failed', err);
    }
}
function queueDraggedAnnotation(payload) {
    clearTimeout(annotationDragSyncTimer);
    annotationDragSyncTimer = setTimeout(() => syncDraggedAnnotation(payload), 160);
}
function handleAnnotationRelayout(evt) {
    if (annotationRelayoutGuard || !evt) return;
    const plotDiv = plotDivEl();
    if (!plotDiv) return;

    const shapeChanges = changedIndexedKeys(evt, 'shapes');
    shapeChanges.forEach((changes, idx) => {
        if (idx >= engineeringShapeMap.length) return; // Cursor overlay or non-engineering shape.
        const map = engineeringShapeMap[idx];
        const shape = plotDiv.layout?.shapes?.[idx];
        if (!map || !shape) return;
        engineeringBaseShapes[idx] = JSON.parse(JSON.stringify(shape));
        if (map.kind === 'line') {
            const y0 = Number(shape.y0), y1 = Number(shape.y1);
            const value = Number.isFinite(y0) && Number.isFinite(y1) ? (y0 + y1) / 2 : Number(shape.y0 ?? shape.y1);
            if (!Number.isFinite(value)) return;
            shape.x0 = 0; shape.x1 = 1; shape.y0 = value; shape.y1 = value;
            engineeringBaseShapes[idx] = JSON.parse(JSON.stringify(shape));
            annotationRelayoutGuard = true;
            Plotly.relayout(plotDiv, {
                [`shapes[${idx}].x0`]:0, [`shapes[${idx}].x1`]:1,
                [`shapes[${idx}].y0`]:value, [`shapes[${idx}].y1`]:value
            }).finally(() => { annotationRelayoutGuard=false; });
            updateManagerField('line', map.id, 'value', value);
            const labelIdx = engineeringAnnotationMap.findIndex(x => x.id === map.id && x.kind === 'line-label');
            if (labelIdx >= 0 && engineeringBaseAnnotations[labelIdx]) {
                engineeringBaseAnnotations[labelIdx].y = value;
                annotationRelayoutGuard = true;
                Plotly.relayout(plotDiv, {[`annotations[${labelIdx}].y`]:value}).finally(() => { annotationRelayoutGuard=false; });
            }
            queueDraggedAnnotation({kind:'line', id:map.id, value});
        } else if (map.kind === 'band') {
            let lo = Number(shape.y0), hi = Number(shape.y1);
            if (!Number.isFinite(lo) || !Number.isFinite(hi)) return;
            if (lo > hi) [lo, hi] = [hi, lo];
            shape.x0 = 0; shape.x1 = 1; shape.y0 = lo; shape.y1 = hi;
            engineeringBaseShapes[idx] = JSON.parse(JSON.stringify(shape));
            annotationRelayoutGuard = true;
            Plotly.relayout(plotDiv, {
                [`shapes[${idx}].x0`]:0, [`shapes[${idx}].x1`]:1,
                [`shapes[${idx}].y0`]:lo, [`shapes[${idx}].y1`]:hi
            }).finally(() => { annotationRelayoutGuard=false; });
            updateManagerField('band', map.id, 'min', lo);
            updateManagerField('band', map.id, 'max', hi);
            const labelIdx = engineeringAnnotationMap.findIndex(x => x.id === map.id && x.kind === 'band-label');
            if (labelIdx >= 0 && engineeringBaseAnnotations[labelIdx]) {
                engineeringBaseAnnotations[labelIdx].y = hi;
                annotationRelayoutGuard = true;
                Plotly.relayout(plotDiv, {[`annotations[${labelIdx}].y`]:hi}).finally(() => { annotationRelayoutGuard=false; });
            }
            queueDraggedAnnotation({kind:'band', id:map.id, min:lo, max:hi});
        }
    });

    const annotationChanges = changedIndexedKeys(evt, 'annotations');
    annotationChanges.forEach((changes, idx) => {
        if (idx >= engineeringAnnotationMap.length) return;
        const map = engineeringAnnotationMap[idx];
        if (!map || map.kind !== 'comment') return; // Limit labels remain attached to their limit.
        const ann = plotDiv.layout?.annotations?.[idx];
        if (!ann) return;
        engineeringBaseAnnotations[idx] = JSON.parse(JSON.stringify(ann));
        const payload = {kind:'comment', id:map.id};
        if (ann.x !== undefined) payload.x = ann.x;
        if (ann.y !== undefined && Number.isFinite(Number(ann.y))) payload.y = Number(ann.y);
        if (ann.ax !== undefined && Number.isFinite(Number(ann.ax))) payload.ax = Number(ann.ax);
        if (ann.ay !== undefined && Number.isFinite(Number(ann.ay))) payload.ay = Number(ann.ay);
        if ('x' in payload) updateManagerField('comment', map.id, 'x', payload.x);
        if ('y' in payload) updateManagerField('comment', map.id, 'y', payload.y);
        if ('ax' in payload) updateManagerField('comment', map.id, 'ax', payload.ax);
        if ('ay' in payload) updateManagerField('comment', map.id, 'ay', payload.ay);
        queueDraggedAnnotation(payload);
    });
}

function resetInteractiveView() {
    const plotDiv = plotDivEl();
    if (!plotDiv || !interactiveOriginalLayout) return;
    const update = {'xaxis.autorange':true};
    if (interactiveOriginalLayout.yaxis?.range) {
        update['yaxis.autorange'] = false; update['yaxis.range'] = interactiveOriginalLayout.yaxis.range;
    } else update['yaxis.autorange'] = true;
    if (interactiveOriginalLayout.yaxis2?.range) {
        update['yaxis2.autorange'] = false; update['yaxis2.range'] = interactiveOriginalLayout.yaxis2.range;
    } else update['yaxis2.autorange'] = true;
    Plotly.relayout(plotDiv, update).then(scheduleAnalysisUpdate);
}

function autoscaleInteractiveView() {
    const plotDiv = plotDivEl();
    if (!plotDiv) return;
    Plotly.relayout(plotDiv, {'xaxis.autorange':true,'yaxis.autorange':true,'yaxis2.autorange':true}).then(scheduleAnalysisUpdate);
}

function showAllInteractiveSignals() {
    const plotDiv = plotDivEl();
    if (!plotDiv || !plotDiv.data) return;
    Plotly.restyle(plotDiv, {visible:true}).then(scheduleAnalysisUpdate);
}

function armGraphComment() {
    commentMode = true; pendingCommentPoint = null; clearCursorArmedState();
    const plotDiv = plotDivEl();
    if (plotDiv) plotDiv.style.cursor = 'crosshair';
    const msg = document.getElementById('cursorInstruction');
    if (msg) msg.textContent = 'Comment mode armed — click the graph point to annotate.';
}
function openCommentModal() {
    const modal = document.getElementById('commentModal');
    if (!modal || !pendingCommentPoint) return;
    document.getElementById('commentAxis').value = pendingCommentPoint.axis || 'left';
    document.getElementById('commentPointHint').textContent = `Point: ${formatX(pendingCommentPoint.x)} · Y ${formatNumber(pendingCommentPoint.y)} · ${pendingCommentPoint.axis === 'right' ? 'Secondary' : 'Primary'}`;
    document.getElementById('commentText').value = '';
    modal.hidden = false;
    document.body.classList.add('modal-open');
    document.getElementById('commentText').focus();
    const plotDiv = plotDivEl(); if (plotDiv) plotDiv.style.cursor = '';
}
function closeCommentModal() {
    const modal = document.getElementById('commentModal'); if (modal) modal.hidden = true;
    document.body.classList.remove('modal-open');
    commentMode = false; pendingCommentPoint = null;
    const plotDiv = plotDivEl(); if (plotDiv) plotDiv.style.cursor = '';
}
async function saveGraphComment() {
    if (!pendingCommentPoint) return;
    const text = document.getElementById('commentText')?.value?.trim();
    if (!text) { alert('Please enter a comment.'); return; }
    const payload = {
        text, x: pendingCommentPoint.x, y: pendingCommentPoint.y,
        axis: document.getElementById('commentAxis')?.value || pendingCommentPoint.axis || 'left',
        text_color: document.getElementById('commentTextColor')?.value || '#101828',
        bg_color: document.getElementById('commentBgColor')?.value || '#FFFFFF',
        border_color: document.getElementById('commentBorderColor')?.value || '#101828',
        border_width: Number(document.getElementById('commentBorderWidth')?.value || 1),
        box_width: Number(document.getElementById('commentBoxWidth')?.value || 0),
        box_height: Number(document.getElementById('commentBoxHeight')?.value || 0),
        font_size: Number(document.getElementById('commentFontSize')?.value || 12),
        show_arrow: !!document.getElementById('commentArrow')?.checked,
        arrow_color: document.getElementById('commentArrowColor')?.value || '#101828',
        arrow_width: Number(document.getElementById('commentArrowWidth')?.value || 1.2)
    };
    const r = await fetch(`/api/annotation/comment/${ACTIVE_TAB}`, {method:'POST', headers:{'Content-Type':'application/json'}, body:JSON.stringify(payload)});
    if (!r.ok) { alert('Failed to save comment: ' + await r.text()); return; }
    saveSidebarUiState();
    window.location.reload();
}
async function deleteAnnotation(kind, id) {
    if (!confirm('Delete this annotation?')) return;
    const r = await fetch(`/api/annotation/${ACTIVE_TAB}/${encodeURIComponent(kind)}/${encodeURIComponent(id)}`, {method:'DELETE'});
    if (!r.ok) { alert('Delete failed: ' + await r.text()); return; }
    saveSidebarUiState();
    window.location.reload();
}
async function toggleAnnotation(kind, id, visible) {
    const r = await fetch(`/api/annotation/toggle/${ACTIVE_TAB}`, {method:'POST', headers:{'Content-Type':'application/json'}, body:JSON.stringify({kind,id,visible})});
    if (!r.ok) { alert('Update failed: ' + await r.text()); return; }
    await loadInteractivePlot();
}

function filterSignals() {
    const q = (document.getElementById('signalSearch')?.value || '').toLowerCase();
    document.querySelectorAll('[data-signal-row]').forEach(row => { row.style.display = row.dataset.name.includes(q) ? '' : 'none'; });
}
function selectVisible(checked) {
    document.querySelectorAll('[data-signal-row]').forEach(row => {
        if (row.style.display !== 'none') { const cb=row.querySelector('input[type=checkbox]'); if(cb) cb.checked=checked; }
    });
}
function assignChecked(axis) {
    document.querySelectorAll('[data-signal-row]').forEach(row => {
        const cb=row.querySelector('input[type=checkbox]'); const sel=row.querySelector('.axis-select');
        if(cb && cb.checked && sel) { sel.value=axis; updateSignalRowClass(sel); refreshAutoColorForRow(sel); }
    });
}
function updateSignalRowClass(select) {
    const row=select.closest('[data-signal-row]'); if(!row) return;
    row.classList.remove('primary-row','secondary-row'); row.classList.add(select.value==='right'?'secondary-row':'primary-row');
}
function setManualColor(input) {
    const control=input.closest('.color-control'); if(!control) return;
    const mode=control.querySelector('.color-mode'); if(mode) mode.value='manual';
    control.classList.remove('auto');
}
function setAutoColor(button) {
    const control=button.closest('.color-control'); if(!control) return;
    const mode=control.querySelector('.color-mode'); if(mode) mode.value='auto';
    control.classList.add('auto');
    const row=control.closest('[data-signal-row]');
    const select=row?.querySelector('.axis-select');
    const picker=control.querySelector('.signal-color');
    if(picker && select) picker.value = select.value==='right' ? picker.dataset.autoSecondary : picker.dataset.autoPrimary;
}
function refreshAutoColorForRow(select) {
    const row=select.closest('[data-signal-row]'); if(!row) return;
    const control=row.querySelector('.color-control');
    const mode=control?.querySelector('.color-mode');
    const picker=control?.querySelector('.signal-color');
    if(control && mode && picker && mode.value==='auto') {
        picker.value = select.value==='right' ? picker.dataset.autoSecondary : picker.dataset.autoPrimary;
    }
}
const CLIENT_ID = (() => {
    const key = 'signalworks_studio_client_id';
    try {
        let id = sessionStorage.getItem(key);
        if (!id) {
            id = (globalThis.crypto && typeof crypto.randomUUID === 'function')
                ? crypto.randomUUID()
                : `gp-${Date.now()}-${Math.random().toString(16).slice(2)}`;
            sessionStorage.setItem(key, id);
        }
        return id;
    } catch (_) {
        return `gp-${Date.now()}-${Math.random().toString(16).slice(2)}`;
    }
})();
let heartbeatTimer = null;

async function sendServerHeartbeat() {
    try {
        await fetch('/api/heartbeat', {
            method: 'POST',
            headers: {'X-SignalWorks-Studio-Client': CLIENT_ID},
            cache: 'no-store',
            keepalive: true
        });
    } catch (_) {}
}

function startServerHeartbeat() {
    sendServerHeartbeat();
    if (heartbeatTimer !== null) clearInterval(heartbeatTimer);
    heartbeatTimer = setInterval(sendServerHeartbeat, 2500);
}

function releaseServerClient() {
    try {
        if (heartbeatTimer !== null) clearInterval(heartbeatTimer);
        heartbeatTimer = null;
        const url = `/api/client-close?client=${encodeURIComponent(CLIENT_ID)}`;
        if (navigator.sendBeacon) navigator.sendBeacon(url, new Blob([], {type:'text/plain'}));
        else fetch(url, {method:'POST', keepalive:true}).catch(() => {});
    } catch (_) {}
}

window.addEventListener('DOMContentLoaded', () => {
    initExclusiveHeaderMenus();
    initCsvMultiFileUI();
    restoreSidebarUiState();
    initNestedScrollHandoff();
    setThemeMode(currentThemeMode(), false);
    startServerHeartbeat();
    loadInteractivePlot();
});
window.addEventListener('resize', () => {
    if (interactiveViewReady) positionAxisLegends(plotDivEl());
});
window.addEventListener('pagehide', () => { savePlotViewState(); saveSidebarUiState(); releaseServerClient(); });
"""


def render_page(state: BrowserSession, active_tab: int) -> str:
    active_tab = max(0, min(active_tab, len(TAB_PRESETS) - 1))
    has_data = state.df is not None
    export_base = re.sub(r"[^A-Za-z0-9_-]+", "_", os.path.splitext(state.filename or "SignalWorksStudio")[0]).strip("_") or "SignalWorksStudio"
    export_tab_options = "".join(
        f'<label class="export-tab-option"><input type="checkbox" value="{idx}" checked><span><strong>Page {idx + 1}</strong><small>{esc(preset.title)}</small></span></label>'
        for idx, preset in enumerate(TAB_PRESETS)
    )

    if has_data:
        dataset_strip = f"""
        <div class="dataset-strip">
            <div class="dataset-main">
                <span class="dataset-dot"></span>
                <span class="dataset-name">{esc(state.filename)}</span>
                <span>· {state.row_count:,} rows</span>
                {f'<span>· bound from {len(state.loaded_source_files)} files</span>' if len(state.loaded_source_files) > 1 else ''}
                <span>· {len(state.numeric_cols)} numeric signals</span>
                <span>· header row {state.header_row_index + 1} ({esc(state.header_detection)})</span>
                <span>· {esc(state.parsed_info)}</span>
            </div>
        </div>
        """
    else:
        dataset_strip = '<div class="dataset-strip"><div class="dataset-main"><span class="dataset-name">No CSV loaded</span></div></div>'

    if has_data:
        main_content = f'<nav class="tabs">{render_tabs(state, active_tab)}</nav>{render_active_tab(state, active_tab)}'
        report_menu = """
            <details class="header-menu">
                <summary class="btn">Export</summary>
                <div class="menu-panel">
                    <div class="menu-title">Reports</div>
                    <button class="menu-item" type="button" onclick="openReportExportDialog('wide')"><span>PDF wide report</span><span class="menu-note">Select pages · 1200 × 650 pt</span></button>
                    <button class="menu-item" type="button" onclick="openReportExportDialog('a4')"><span>PDF A4 report</span><span class="menu-note">Select pages · 842 × 595 pt</span></button>
                    <button class="menu-item" type="button" onclick="openReportExportDialog('word')"><span>Word report</span><span class="menu-note">Select pages · DOCX</span></button>
                </div>
            </details>
        """
    else:
        main_content = """
            <section class="empty-state">
                <div class="empty-icon">⌁</div>
                <h2>Load a CSV file to begin</h2>
                <p>Upload any CSV log file. The header row is detected automatically from Timestamp / Date&Time / Date / Time in the first column. Without a .preset file the application starts in Normal Plot Mode.</p>
            </section>
        """
        report_menu = ""

    js = (
        V42_JS.replace('__ACTIVE_TAB__', str(active_tab))
        .replace('__REVISION__', str(state.tabs[active_tab].revision if has_data else 0))
        .replace('__EXPORT_BASE_JSON__', json.dumps(export_base, ensure_ascii=False))
        .replace('__PRESET_NAME_JSON__', json.dumps(state.preset_name or "Signal Preset", ensure_ascii=False))
        .replace('__DATASET_TOKEN_JSON__', json.dumps(state.dataset_token or "no-data", ensure_ascii=False))
        .replace('__TAB_VIEW_TOKENS_JSON__', json.dumps([tab.view_token for tab in state.tabs]))
    )

    return f"""<!doctype html>
<html lang="en">
<head>
    <meta charset="utf-8">
    <meta name="viewport" content="width=device-width, initial-scale=1">
    <title>{esc(APP_TITLE)} — {esc(APP_TAGLINE)}</title>
    <script>
    (() => {{
        const key='signalworks_studio_theme_session_mode';
        let mode='system';
        try {{ mode=sessionStorage.getItem(key)||'system'; }} catch (_) {{}}
        const resolved=(mode==='dark'||mode==='light') ? mode : ((window.matchMedia&&window.matchMedia('(prefers-color-scheme: dark)').matches)?'dark':'light');
        document.documentElement.dataset.themeMode=mode;
        document.documentElement.dataset.theme=resolved;
    }})();
    </script>
    <style>{V42_CSS}</style>
</head>
<body>
    <header class="app-header">
        <div class="topbar">
            <div class="brand">
                <div class="brand-mark" aria-label="SignalWorks Studio brand mark" title="SignalWorks Studio">
                    <svg viewBox="0 0 36 36" aria-hidden="true" focusable="false">
                        <path class="signal-wave" d="M3 19h6l3-8 4 16 4-20 4 12h9" />
                    </svg>
                </div>
                <div>
                    <h1>SignalWorks <span class="brand-studio">Studio</span></h1>
                    <div class="version">{esc(APP_TAGLINE)} · WebServer {esc(APP_VERSION)}</div>
                </div>
            </div>
            <div class="header-tools">
                <div class="upload-form">
                    <span class="header-inline-label">CSV</span>
                    <input id="csvFilePicker" class="file-picker" type="file" accept=".csv,text/csv" multiple onchange="handleCsvFileSelection(this.files)" title="Choose one or more CSV files">
                </div>

                <span class="preset-mode-badge {'active' if state.loaded_preset else ''}" title="{'Active preset: ' + esc(state.preset_name) if state.loaded_preset else 'No signal preset loaded'}">
                    {'Preset · ' + esc(state.preset_name) if state.loaded_preset else 'Normal mode'}
                </span>

                <details class="header-menu">
                    <summary class="btn">Preset</summary>
                    <div class="menu-panel wide">
                        <div class="menu-title">Signal preset</div>
                        <form class="menu-form" method="post" action="/preset/import" enctype="multipart/form-data">
                            <input class="preset-file" type="file" name="preset_file" accept=".preset,application/json" required title="Import portable signal preset">
                            <button class="btn preset-import" type="submit">Import preset</button>
                        </form>
                        <div class="menu-separator"></div>
                        <button class="menu-item" type="button" onclick="openPresetExportDialog()"><span>Export current preset</span><span class="menu-note">Name + Save As</span></button>
                        {('<form class="menu-form" method="post" action="/preset/clear"><button class="menu-item" type="submit"><span>Switch to Normal mode</span><span class="menu-note">Unload preset</span></button></form>' if state.loaded_preset else '')}
                    </div>
                </details>

                {report_menu}

                <div class="theme-switch-wrap" title="Switch Light / Dark theme">
                    <span class="theme-switch-icon sun" aria-hidden="true">
                        <svg viewBox="0 0 24 24"><circle cx="12" cy="12" r="3.6"></circle><path d="M12 2.4v2M12 19.6v2M2.4 12h2M19.6 12h2M5.2 5.2l1.4 1.4M17.4 17.4l1.4 1.4M18.8 5.2l-1.4 1.4M6.6 17.4l-1.4 1.4"></path></svg>
                    </span>
                    <button id="themeToggle" class="theme-switch" type="button" role="switch" aria-checked="false" aria-label="Switch Light / Dark theme" onclick="toggleThemeMode()">
                        <span class="theme-switch-knob"></span>
                    </button>
                    <span class="theme-switch-icon moon" aria-hidden="true">
                        <svg viewBox="0 0 24 24"><path d="M19.2 15.5A8.2 8.2 0 0 1 8.5 4.8 8.2 8.2 0 1 0 19.2 15.5Z"></path></svg>
                    </span>
                </div>
            </div>
        </div>
    </header>

    <div class="page-shell">
        {dataset_strip}
        <div class="{status_class(state.status_type)}">{esc(state.status)}</div>
        {main_content}
        <div id="csvSelectionModal" class="modal-backdrop" hidden onclick="if(event.target===this) closeCsvSelectionModal(true)">
          <div class="comment-modal csv-bind-modal" role="dialog" aria-modal="true" aria-labelledby="csvSelectionModalTitle">
            <div class="modal-head">
              <div><div class="eyebrow">CSV DATA</div><h3 id="csvSelectionModalTitle">Confirm CSV files</h3></div>
              <button class="icon-close" type="button" onclick="closeCsvSelectionModal(true)">×</button>
            </div>
            <div class="modal-body">
              <div class="modal-help">
                Files are initially sorted by the time available to the browser (<strong>Last Modified</strong>).
                Browsers do not expose Windows Creation Time. Drag and drop rows to correct the binding order if needed.
              </div>
              <div class="export-section-head"><span id="csvSelectionCount">0 files selected</span><span>Top = first data file</span></div>
              <div id="csvFileList" class="csv-file-list"></div>
              <div class="csv-progress-wrap">
                <div class="csv-progress-track"><div id="csvBindProgressBar" class="csv-progress-bar"></div></div>
                <div id="csvBindProgressText" class="csv-progress-text">Select CSV files to begin.</div>
              </div>
              <div class="csv-bind-summary">
                <span>1 file → Load directly</span>
                <span>2+ files → Bind Data first</span>
                <span>Subsequent headers are removed automatically</span>
              </div>
            </div>
            <div class="modal-actions">
              <div class="csv-modal-actions-left">
                <button id="csvDownloadBoundBtn" class="btn" type="button" onclick="downloadBoundCsv()" disabled hidden>Download Bound CSV</button>
              </div>
              <button class="btn" type="button" onclick="closeCsvSelectionModal(true)">Cancel</button>
              <button id="csvBindBtn" class="btn" type="button" onclick="bindSelectedCsvFiles()" hidden>Bind Data</button>
              <button id="csvLoadBtn" class="btn primary" type="button" onclick="loadCsvSelection()" disabled>Load CSV</button>
            </div>
          </div>
        </div>
        <div id="exportModal" class="modal-backdrop" hidden onclick="if(event.target===this) closeExportModal()">
          <div class="comment-modal export-modal" role="dialog" aria-modal="true" aria-labelledby="exportModalTitle">
            <div class="modal-head"><div><div class="eyebrow">SAVE / EXPORT</div><h3 id="exportModalTitle">Export</h3></div><button class="icon-close" type="button" onclick="closeExportModal()">×</button></div>
            <div class="modal-body">
              <div id="exportModalHelp" class="modal-help">Choose the output settings and file name.</div>
              <div id="exportPresetNameRow" hidden><label>Preset name<input id="exportPresetName" type="text" value="{esc(state.preset_name or 'Signal Preset')}" placeholder="Preset display name"></label></div>
              <label>File name<input id="exportFileName" type="text" value="{esc(export_base)}" spellcheck="false"></label>
              <div id="exportTabSection" hidden>
                <div class="export-section-head"><span>Report pages</span><span><button class="link-button" type="button" onclick="setAllExportTabs(true)">Select all</button> · <button class="link-button" type="button" onclick="setAllExportTabs(false)">Clear</button></span></div>
                <div id="exportTabGrid" class="export-tab-grid">{export_tab_options}</div>
              </div>
              <div class="save-path-note"><strong>Save location:</strong> Edge/Chrome will show a native Save As dialog when the File System Access API is available. On LAN HTTP sessions, browser download settings may control the destination folder.</div>
            </div>
            <div class="modal-actions"><button class="btn" type="button" onclick="closeExportModal()">Cancel</button><button id="exportConfirmBtn" class="btn primary" type="button" onclick="confirmExportDialog()">Save As…</button></div>
          </div>
        </div>
        <div class="footer">SignalWorks Studio · Visualize · Evaluate · Analyze · Report · {esc(APP_VERSION)}</div>
    </div>

    <script src="/static/plotly.min.js"></script>
    <script>{js}</script>
</body>
</html>
"""


# ------------------------------------------------------------------------------
# WebServer lifecycle helpers (V4.2.4)
# ------------------------------------------------------------------------------
def _normalize_client_id(value: Optional[str]) -> str:
    value = (value or "").strip()
    if not value:
        return ""
    # Browser-generated UUIDs are short. Cap the value to avoid unbounded headers.
    return value[:128]


def register_browser_heartbeat(client_id: str) -> None:
    global _ever_had_browser_client, _no_client_since
    client_id = _normalize_client_id(client_id)
    if not client_id:
        return
    now = time.monotonic()
    with _lifecycle_lock:
        _client_last_seen[client_id] = now
        _ever_had_browser_client = True
        _no_client_since = None


def release_browser_client(client_id: str) -> None:
    global _no_client_since
    client_id = _normalize_client_id(client_id)
    now = time.monotonic()
    with _lifecycle_lock:
        if client_id:
            _client_last_seen.pop(client_id, None)
        if _ever_had_browser_client and not _client_last_seen and _no_client_since is None:
            _no_client_since = now


def request_server_shutdown() -> bool:
    server = _managed_uvicorn_server
    if server is None:
        return False
    try:
        server.should_exit = True
        return True
    except Exception:
        return False


def lifecycle_watchdog() -> None:
    """Stop Uvicorn after all browser tabs have disappeared.

    A clean page close removes the client immediately and starts the short idle
    countdown. If a browser crashes or sleeps without pagehide, stale heartbeats
    are pruned after CLIENT_HEARTBEAT_STALE_SECONDS, providing a fallback cleanup.
    """
    global _no_client_since
    while True:
        time.sleep(1.0)
        server = _managed_uvicorn_server
        if server is None:
            continue
        if getattr(server, "should_exit", False):
            return
        if not _auto_shutdown_enabled:
            continue

        now = time.monotonic()
        should_stop = False
        with _lifecycle_lock:
            stale = [cid for cid, seen in _client_last_seen.items() if now - seen > CLIENT_HEARTBEAT_STALE_SECONDS]
            for cid in stale:
                _client_last_seen.pop(cid, None)

            if _client_last_seen:
                _no_client_since = None
            elif _ever_had_browser_client:
                if _no_client_since is None:
                    _no_client_since = now
                elif now - _no_client_since >= _idle_shutdown_seconds:
                    should_stop = True

        if should_stop:
            try:
                server.should_exit = True
            except Exception:
                pass
            return


# ------------------------------------------------------------------------------
# FastAPI application
# ------------------------------------------------------------------------------
app = FastAPI(title=APP_TITLE, version=APP_VERSION)


@app.get("/api/health")
async def server_health():
    with _lifecycle_lock:
        active_clients = len(_client_last_seen)
    return {"ok": True, "app": APP_TITLE, "version": APP_VERSION, "active_clients": active_clients}


@app.post("/api/heartbeat")
async def browser_heartbeat(request: Request):
    client_id = request.headers.get("X-SignalWorks-Studio-Client") or request.headers.get("X-GraphPlot-Client") or request.query_params.get("client") or ""
    register_browser_heartbeat(client_id)
    return {"ok": True}


@app.post("/api/client-close")
async def browser_client_close(request: Request):
    client_id = request.headers.get("X-SignalWorks-Studio-Client") or request.headers.get("X-GraphPlot-Client") or request.query_params.get("client") or ""
    release_browser_client(client_id)
    return {"ok": True}


@app.post("/api/shutdown")
async def manual_server_shutdown(request: Request):
    # Keep the explicit shutdown endpoint local-only. LAN browser clients can use
    # the normal heartbeat/close lifecycle but cannot remotely stop the service.
    host = request.client.host if request.client else ""
    if host not in {"127.0.0.1", "::1", "localhost"}:
        return Response(content="Shutdown is allowed from localhost only.", status_code=403, media_type="text/plain")
    ok = request_server_shutdown()
    return {"ok": ok}


@app.get("/static/plotly.min.js")
async def plotly_javascript():
    # Plotly runtime is served from the installed Python package for offline use.
    response = Response(content=PLOTLY_JS, media_type="application/javascript")
    response.headers["Cache-Control"] = "public, max-age=86400"
    return response


@app.post("/api/cycle-analysis/{tab_id}")
async def cycle_analysis_api(request: Request, tab_id: int):
    state, _ = get_or_create_session(request)
    if state.df is None:
        return Response(content=json.dumps({"error": "No CSV loaded."}), status_code=400, media_type="application/json")
    if not (0 <= tab_id < len(state.tabs)):
        return Response(content=json.dumps({"error": "Invalid tab."}), status_code=400, media_type="application/json")
    try:
        payload = await request.json()
        signal = str(payload.get("signal") or "").strip()
        profile_mode = str(payload.get("profile_mode", "auto"))
        tolerance = float(payload.get("tolerance", 2.0))
        sample_time = float(payload.get("sample_time", 10.0))
        if signal not in state.tabs[tab_id].selected_signals:
            raise ValueError("Cycle Analysis can only use a signal currently selected for plotting on this tab.")
        result = analyze_cycle_signal(
            state,
            signal,
            tolerance=tolerance,
            fallback_sample_time=sample_time,
            profile_mode=profile_mode,
        )
        return Response(content=json.dumps(result, ensure_ascii=False), media_type="application/json")
    except Exception as exc:
        return Response(content=json.dumps({"error": str(exc)}, ensure_ascii=False), status_code=400, media_type="application/json")


@app.get("/api/plot/{tab_id}")
async def interactive_plot_json(request: Request, tab_id: int):
    state, is_new = get_or_create_session(request)
    if state.df is None or not (0 <= tab_id < len(TAB_PRESETS)):
        response = Response(
            content=json.dumps({"error": "No plot available"}),
            media_type="application/json",
            status_code=404,
        )
        return response_with_session_cookie(response, state, is_new)

    try:
        fig = build_interactive_figure(state, tab_id)
        payload = pio.to_json(fig, validate=False, pretty=False, remove_uids=True)
        response = Response(content=payload, media_type="application/json")
        response.headers["Cache-Control"] = "no-store, max-age=0"
    except Exception as exc:
        response = Response(
            content=json.dumps({"error": str(exc)}),
            media_type="application/json",
            status_code=500,
        )
    return response_with_session_cookie(response, state, is_new)


@app.get("/", response_class=HTMLResponse)
async def index(request: Request, tab: int = 0):
    state, is_new = get_or_create_session(request)
    active_tab = max(0, min(int(tab), len(TAB_PRESETS) - 1))
    response = HTMLResponse(render_page(state, active_tab))
    return response_with_session_cookie(response, state, is_new)


@app.post("/bind/csv")
async def bind_csv_files(request: Request, files: List[UploadFile] = File(...)):
    state, is_new = get_or_create_session(request)
    try:
        if len(files) < 2:
            raise ValueError("Select at least two CSV files before using Bind Data.")

        payloads: List[Tuple[str, bytes]] = []
        for upload in files:
            name = os.path.basename(upload.filename or "data.csv")
            if not name.lower().endswith(".csv"):
                raise ValueError(f"'{name}' is not a .csv file.")
            payloads.append((name, await upload.read()))

        bound_bytes, info = bind_csv_payloads(payloads)
        state.bound_csv_bytes = bound_bytes
        state.bound_filename = info["filename"]
        state.bound_source_files = list(info["source_files"])
        state.bound_source_rows = list(info["source_rows"])
        state.bound_total_rows = int(info["total_rows"])

        response = JSONResponse({
            "ok": True,
            **info,
        })
        return response_with_session_cookie(response, state, is_new)
    except Exception as exc:
        response = JSONResponse({"ok": False, "error": str(exc)}, status_code=400)
        return response_with_session_cookie(response, state, is_new)


@app.post("/bind/load")
async def load_bound_csv(request: Request):
    state, is_new = get_or_create_session(request)
    try:
        if not state.bound_csv_bytes:
            raise ValueError("No bound CSV is available. Select files and run Bind Data first.")
        load_csv_payload_into_state(
            state,
            state.bound_csv_bytes,
            state.bound_filename or "Bound_Data.csv",
            source_files=state.bound_source_files,
        )
        response = JSONResponse({
            "ok": True,
            "filename": state.filename,
            "rows": state.row_count,
            "source_files": list(state.loaded_source_files),
        })
        return response_with_session_cookie(response, state, is_new)
    except Exception as exc:
        state.status = f"Load error: {exc}"
        state.status_type = "error"
        response = JSONResponse({"ok": False, "error": str(exc)}, status_code=400)
        return response_with_session_cookie(response, state, is_new)


@app.get("/bind/download")
async def download_bound_csv(request: Request):
    state, is_new = get_or_create_session(request)
    if not state.bound_csv_bytes:
        response = Response("No bound CSV is available.", status_code=404, media_type="text/plain")
        return response_with_session_cookie(response, state, is_new)

    filename = state.bound_filename or "Bound_Data.csv"
    response = Response(
        content=state.bound_csv_bytes,
        media_type="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}",
            "Cache-Control": "no-store",
        },
    )
    return response_with_session_cookie(response, state, is_new)


@app.post("/upload")
async def upload_csv(request: Request, file: UploadFile = File(...)):
    state, is_new = get_or_create_session(request)

    try:
        file_bytes = await file.read()
        filename = os.path.basename(file.filename or "uploaded.csv")
        load_csv_payload_into_state(state, file_bytes, filename, source_files=[filename])

        # A direct single-file upload is a separate workflow from any previous bind.
        state.bound_csv_bytes = b""
        state.bound_filename = ""
        state.bound_source_files = []
        state.bound_source_rows = []
        state.bound_total_rows = 0

    except Exception as exc:
        state.status = f"Load error: {exc}"
        state.status_type = "error"

    response = RedirectResponse(url="/?tab=0", status_code=303)
    return response_with_session_cookie(response, state, is_new)


@app.post("/preset/import")
async def import_signal_preset(request: Request, preset_file: UploadFile = File(...)):
    state, is_new = get_or_create_session(request)
    try:
        raw = await preset_file.read()
        if not raw:
            raise ValueError("Preset file is empty.")
        if len(raw) > 2 * 1024 * 1024:
            raise ValueError("Preset file is too large (maximum 2 MB).")
        try:
            payload = json.loads(raw.decode("utf-8-sig"))
        except Exception as exc:
            raise ValueError(f"Cannot read preset JSON: {exc}") from exc

        payload = normalize_preset_payload(payload)
        state.loaded_preset = payload
        state.preset_name = str(payload.get("name") or os.path.splitext(os.path.basename(preset_file.filename or "preset"))[0])
        state.preset_filename = os.path.basename(preset_file.filename or "imported.preset")
        matched, requested = apply_loaded_preset(state)

        if state.df is not None:
            state.status = f"Preset '{state.preset_name}' imported · matched {matched}/{requested} configured signals in current CSV"
        else:
            state.status = f"Preset '{state.preset_name}' imported · it will be matched automatically when a CSV is loaded"
        state.status_type = "success"
    except Exception as exc:
        state.status = f"Preset import error: {exc}"
        state.status_type = "error"

    response = RedirectResponse(url="/?tab=0", status_code=303)
    return response_with_session_cookie(response, state, is_new)


@app.get("/preset/export")
async def export_signal_preset(request: Request):
    state, is_new = get_or_create_session(request)
    requested_name = (request.query_params.get("name") or "").strip()
    export_name = requested_name or state.preset_name or "Signal Preset"
    payload = build_preset_payload_from_state(state, preset_name=export_name)
    content = json.dumps(payload, indent=2, ensure_ascii=False).encode("utf-8")
    base_name = export_name
    safe_name = "".join(c if (c.isalnum() or c in "-_ ") else "_" for c in base_name).strip().replace(" ", "_") or "Signal_Preset"
    filename = f"{safe_name}.preset"
    response = Response(content=content, media_type="application/json; charset=utf-8")
    response.headers["Content-Disposition"] = f'attachment; filename="{filename}"'
    response.headers["Cache-Control"] = "no-store, max-age=0"
    return response_with_session_cookie(response, state, is_new)


@app.post("/preset/clear")
async def clear_signal_preset(request: Request):
    state, is_new = get_or_create_session(request)
    old_name = state.preset_name
    state.loaded_preset = None
    state.preset_name = ""
    state.preset_filename = ""
    initialize_tabs_normal(state)
    state.status = f"Preset '{old_name}' unloaded. Normal Plot Mode is active." if old_name else "Normal Plot Mode is active."
    state.status_type = "success"
    response = RedirectResponse(url="/?tab=0", status_code=303)
    return response_with_session_cookie(response, state, is_new)

def _annotation_collection(ts: TabState, kind: str) -> List[dict]:
    if kind == "line":
        return ts.reference_lines
    if kind in ("band", "range"):
        return ts.range_bands
    if kind == "comment":
        return ts.comments
    raise ValueError("Unknown annotation type")


def _bump_annotation_revision(state: BrowserSession, tab_id: int) -> None:
    state.tabs[tab_id].revision += 1
    state.plot_cache.pop(tab_id, None)


@app.post("/annotation/line/{tab_id}")
async def save_reference_line(request: Request, tab_id: int):
    state, is_new = get_or_create_session(request)
    if not (0 <= tab_id < len(state.tabs)):
        return response_with_session_cookie(RedirectResponse(url="/", status_code=303), state, is_new)
    form = await request.form(); ts = state.tabs[tab_id]
    try:
        item_id = str(form.get("item_id") or "").strip()
        raw = {
            "id": item_id or _new_annotation_id("line"), "label": str(form.get("label") or "Reference"),
            "value": float(form.get("value")), "axis": str(form.get("axis") or "left"),
            "color": str(form.get("color") or "#D92D20"), "width": form.get("width") or 1.5,
            "dash": str(form.get("dash") or "dash"), "show_label": form.get("show_label") is not None, "visible": True,
        }
        item = _normalize_reference_lines([raw])[0]
        existing = next((x for x in ts.reference_lines if x.get("id") == item_id), None) if item_id else None
        if existing:
            item["visible"] = existing.get("visible", True); existing.clear(); existing.update(item)
        else:
            ts.reference_lines.append(item)
        _bump_annotation_revision(state, tab_id); state.status = "Reference line saved."; state.status_type = "success"
    except Exception as exc:
        state.status = f"Reference line error: {exc}"; state.status_type = "error"
    return response_with_session_cookie(RedirectResponse(url=f"/?tab={tab_id}", status_code=303), state, is_new)


@app.post("/annotation/range/{tab_id}")
async def save_range_band(request: Request, tab_id: int):
    state, is_new = get_or_create_session(request)
    if not (0 <= tab_id < len(state.tabs)):
        return response_with_session_cookie(RedirectResponse(url="/", status_code=303), state, is_new)
    form = await request.form(); ts = state.tabs[tab_id]
    try:
        item_id = str(form.get("item_id") or "").strip()
        raw = {
            "id": item_id or _new_annotation_id("band"), "label": str(form.get("label") or "Accept range"),
            "min": float(form.get("min")), "max": float(form.get("max")), "axis": str(form.get("axis") or "left"),
            "line_color": str(form.get("line_color") or "#12B76A"), "fill_color": str(form.get("fill_color") or "#12B76A"),
            "opacity": form.get("opacity") or 10, "width": form.get("width") or 1.2, "dash": str(form.get("dash") or "dash"),
            "show_label": form.get("show_label") is not None, "visible": True,
        }
        item = _normalize_range_bands([raw])[0]
        existing = next((x for x in ts.range_bands if x.get("id") == item_id), None) if item_id else None
        if existing:
            item["visible"] = existing.get("visible", True); existing.clear(); existing.update(item)
        else:
            ts.range_bands.append(item)
        _bump_annotation_revision(state, tab_id); state.status = "Range band saved."; state.status_type = "success"
    except Exception as exc:
        state.status = f"Range band error: {exc}"; state.status_type = "error"
    return response_with_session_cookie(RedirectResponse(url=f"/?tab={tab_id}", status_code=303), state, is_new)


@app.post("/api/annotation/comment/{tab_id}")
async def api_add_graph_comment(request: Request, tab_id: int):
    state, is_new = get_or_create_session(request)
    if not (0 <= tab_id < len(state.tabs)):
        return response_with_session_cookie(Response(content="Invalid tab", status_code=400), state, is_new)
    try:
        payload = await request.json(); payload["id"] = _new_annotation_id("comment"); payload["visible"] = True
        items = _normalize_comments([payload])
        if not items:
            raise ValueError("Invalid comment data")
        item = items[0]; state.tabs[tab_id].comments.append(item); _bump_annotation_revision(state, tab_id)
        response = Response(content=json.dumps({"ok": True, "id": item["id"]}), media_type="application/json")
    except Exception as exc:
        response = Response(content=str(exc), status_code=400, media_type="text/plain")
    return response_with_session_cookie(response, state, is_new)


@app.post("/annotation/comment-edit/{tab_id}")
async def edit_graph_comment(request: Request, tab_id: int):
    state, is_new = get_or_create_session(request)
    if not (0 <= tab_id < len(state.tabs)):
        return response_with_session_cookie(RedirectResponse(url="/", status_code=303), state, is_new)
    form = await request.form(); ts = state.tabs[tab_id]
    try:
        item_id = str(form.get("item_id") or ""); existing = next(x for x in ts.comments if x.get("id") == item_id)
        raw = dict(existing)
        raw.update({
            "text": str(form.get("text") or existing.get("text", "")),
            "x": form.get("x") if form.get("x") is not None else existing.get("x"),
            "y": form.get("y") if form.get("y") not in (None, "") else existing.get("y"),
            "ax": form.get("ax") if form.get("ax") not in (None, "") else existing.get("ax", 32),
            "ay": form.get("ay") if form.get("ay") not in (None, "") else existing.get("ay", -38),
            "text_color": str(form.get("text_color") or existing.get("text_color")),
            "bg_color": str(form.get("bg_color") or existing.get("bg_color")),
            "border_color": str(form.get("border_color") or existing.get("border_color", existing.get("text_color", "#101828"))),
            "border_width": form.get("border_width") if form.get("border_width") not in (None, "") else existing.get("border_width", 1),
            "box_width": form.get("box_width") if form.get("box_width") not in (None, "") else existing.get("box_width", 0),
            "box_height": form.get("box_height") if form.get("box_height") not in (None, "") else existing.get("box_height", 0),
            "font_size": form.get("font_size") or existing.get("font_size", 12),
            "show_arrow": form.get("show_arrow") is not None,
            "arrow_color": str(form.get("arrow_color") or existing.get("arrow_color", existing.get("border_color", "#101828"))),
            "arrow_width": form.get("arrow_width") if form.get("arrow_width") not in (None, "") else existing.get("arrow_width", 1.2),
        })
        item = _normalize_comments([raw])[0]; item["visible"] = existing.get("visible", True)
        existing.clear(); existing.update(item); _bump_annotation_revision(state, tab_id)
        state.status = "Comment updated."; state.status_type = "success"
    except Exception as exc:
        state.status = f"Comment update error: {exc}"; state.status_type = "error"
    return response_with_session_cookie(RedirectResponse(url=f"/?tab={tab_id}", status_code=303), state, is_new)


@app.post("/api/annotation/position/{tab_id}")
async def api_update_annotation_position(request: Request, tab_id: int):
    """Persist direct Plotly drag edits without reloading the Web UI."""
    state, is_new = get_or_create_session(request)
    try:
        if not (0 <= tab_id < len(state.tabs)):
            raise ValueError("Invalid tab")
        payload = await request.json()
        kind = str(payload.get("kind") or "")
        item_id = str(payload.get("id") or "")
        collection = _annotation_collection(state.tabs[tab_id], kind)
        item = next(x for x in collection if x.get("id") == item_id)

        if kind == "line":
            item["value"] = float(payload.get("value"))
        elif kind in ("band", "range"):
            lo = float(payload.get("min")); hi = float(payload.get("max"))
            if lo > hi: lo, hi = hi, lo
            item["min"], item["max"] = lo, hi
        elif kind == "comment":
            if "x" in payload:
                item["x"] = payload.get("x")
            if "y" in payload:
                item["y"] = float(payload.get("y"))
            if "ax" in payload:
                item["ax"] = max(-800.0, min(800.0, float(payload.get("ax"))))
            if "ay" in payload:
                item["ay"] = max(-800.0, min(800.0, float(payload.get("ay"))))
        else:
            raise ValueError("Unsupported annotation type")

        _bump_annotation_revision(state, tab_id)
        response = Response(content=json.dumps({"ok": True, "revision": state.tabs[tab_id].revision}), media_type="application/json")
    except Exception as exc:
        response = Response(content=str(exc), status_code=400, media_type="text/plain")
    return response_with_session_cookie(response, state, is_new)


@app.post("/api/annotation/toggle/{tab_id}")
async def api_toggle_annotation(request: Request, tab_id: int):
    state, is_new = get_or_create_session(request)
    try:
        payload = await request.json(); kind = str(payload.get("kind")); item_id = str(payload.get("id")); visible = bool(payload.get("visible"))
        collection = _annotation_collection(state.tabs[tab_id], kind); item = next(x for x in collection if x.get("id") == item_id)
        item["visible"] = visible; _bump_annotation_revision(state, tab_id)
        response = Response(content='{"ok":true}', media_type="application/json")
    except Exception as exc:
        response = Response(content=str(exc), status_code=400, media_type="text/plain")
    return response_with_session_cookie(response, state, is_new)


@app.delete("/api/annotation/{tab_id}/{kind}/{item_id}")
async def api_delete_annotation(request: Request, tab_id: int, kind: str, item_id: str):
    state, is_new = get_or_create_session(request)
    try:
        collection = _annotation_collection(state.tabs[tab_id], kind); before = len(collection)
        collection[:] = [x for x in collection if x.get("id") != item_id]
        if len(collection) == before:
            raise ValueError("Annotation not found")
        _bump_annotation_revision(state, tab_id)
        response = Response(content='{"ok":true}', media_type="application/json")
    except Exception as exc:
        response = Response(content=str(exc), status_code=400, media_type="text/plain")
    return response_with_session_cookie(response, state, is_new)


def _form_value(form: FormData, name: str, fallback: str = "") -> str:
    value = form.get(name, fallback)
    return str(value) if value is not None else fallback


@app.post("/update/{tab_id}")
async def update_tab(request: Request, tab_id: int):
    state, is_new = get_or_create_session(request)

    if state.df is None:
        state.status = "Please load a CSV file first."
        state.status_type = "warning"
        response = RedirectResponse(url="/", status_code=303)
        return response_with_session_cookie(response, state, is_new)

    if not (0 <= tab_id < len(TAB_PRESETS)):
        response = RedirectResponse(url="/", status_code=303)
        return response_with_session_cookie(response, state, is_new)

    form = await request.form()
    ts = state.tabs[tab_id]

    try:
        old_axes = (ts.left_min, ts.left_max, ts.left_step, ts.right_min, ts.right_max, ts.right_step,
                    list(ts.selected_signals), dict(ts.axis_assignments))
        ts.title = _form_value(form, "title", ts.title)
        ts.auto_title = form.get("auto_title") is not None
        ts.x_label = _form_value(form, "x_label", ts.x_label)
        ts.left_label = _form_value(form, "left_label", ts.left_label)
        ts.right_label = _form_value(form, "right_label", ts.right_label)
        ts.left_min = _form_value(form, "left_min", ts.left_min)
        ts.left_max = _form_value(form, "left_max", ts.left_max)
        ts.left_step = _form_value(form, "left_step", ts.left_step)
        ts.right_min = _form_value(form, "right_min", ts.right_min)
        ts.right_max = _form_value(form, "right_max", ts.right_max)
        ts.right_step = _form_value(form, "right_step", ts.right_step)

        selected_signals: List[str] = []
        assignments: Dict[str, str] = dict(ts.axis_assignments)
        colors: Dict[str, str] = dict(ts.color_assignments)
        line_styles: Dict[str, str] = dict(ts.line_style_assignments)

        for idx, signal in enumerate(state.numeric_cols):
            if form.get(f"selected_{idx}") is not None:
                selected_signals.append(signal)
            axis = _form_value(form, f"axis_{idx}", assignments.get(signal, default_axis_for_signal(tab_id, signal)))
            assignments[signal] = axis if axis in ("left", "right") else "left"

            color_mode = _form_value(form, f"color_mode_{idx}", "auto").lower()
            color_value = _form_value(form, f"color_{idx}", "").upper()
            if color_mode == "manual" and _valid_hex_color(color_value):
                colors[signal] = color_value
            else:
                colors.pop(signal, None)

            line_style = _form_value(form, f"line_style_{idx}", "auto").lower()
            if line_style in SIGNAL_LINE_STYLES:
                line_styles[signal] = line_style
            else:
                line_styles.pop(signal, None)

        ts.selected_signals = selected_signals
        ts.axis_assignments = assignments
        ts.color_assignments = colors
        ts.line_style_assignments = line_styles

        # Validate by building a figure now. This catches bad range input before redirect.
        fig = build_figure(state, tab_id)
        plt.close(fig)

        new_axes = (ts.left_min, ts.left_max, ts.left_step, ts.right_min, ts.right_max, ts.right_step,
                    list(ts.selected_signals), dict(ts.axis_assignments))
        if old_axes != new_axes:
            ts.view_token = uuid.uuid4().hex[:16]
        ts.revision += 1
        state.plot_cache.pop(tab_id, None)

        left_count = sum(1 for sig in selected_signals if assignments.get(sig) == "left")
        right_count = len(selected_signals) - left_count
        state.status = f"Tab {tab_id + 1} plotted: {left_count} Primary / {right_count} Secondary signals"
        state.status_type = "success"

    except Exception as exc:
        state.status = f"Plot error: {exc}"
        state.status_type = "error"

    response = RedirectResponse(url=f"/?tab={tab_id}", status_code=303)
    return response_with_session_cookie(response, state, is_new)


@app.post("/reset/{tab_id}")
async def reset_tab(request: Request, tab_id: int):
    state, is_new = get_or_create_session(request)
    if 0 <= tab_id < len(TAB_PRESETS):
        reset_tab_to_loaded_mode(state, tab_id)
        state.status = (f"Tab {tab_id + 1} reloaded from preset." if state.loaded_preset else f"Tab {tab_id + 1} cleared for Normal Plot Mode.")
        state.status_type = "success"
    response = RedirectResponse(url=f"/?tab={max(0, min(tab_id, len(TAB_PRESETS)-1))}", status_code=303)
    return response_with_session_cookie(response, state, is_new)


@app.get("/plot/{tab_id}.png")
async def plot_png(request: Request, tab_id: int):
    state, is_new = get_or_create_session(request)
    if state.df is None or not (0 <= tab_id < len(TAB_PRESETS)):
        response = Response(status_code=404)
        return response_with_session_cookie(response, state, is_new)

    try:
        raw_view = request.query_params.get("view")
        view_state = _parse_export_view(raw_view) if raw_view is not None else None
        image = render_plot_png(state, tab_id, dpi=150, use_cache=True, view_state=view_state)
        response = Response(content=image, media_type="image/png")
        response.headers["Cache-Control"] = "no-store, max-age=0"
    except Exception as exc:
        response = Response(content=f"Plot error: {exc}", media_type="text/plain", status_code=500)
    return response_with_session_cookie(response, state, is_new)


@app.get("/download/plot/{tab_id}")
async def download_plot(request: Request, tab_id: int):
    state, is_new = get_or_create_session(request)
    if state.df is None or not (0 <= tab_id < len(TAB_PRESETS)):
        response = Response(content="No plot available.", media_type="text/plain", status_code=404)
        return response_with_session_cookie(response, state, is_new)

    try:
        view_state = _parse_export_view(request.query_params.get("view"))
        image = render_plot_png(state, tab_id, dpi=300, use_cache=False, view_state=view_state)
        base = os.path.splitext(state.filename)[0] or "plot"
        tab_title = state.tabs[tab_id].title.strip() or f"Plot_{tab_id + 1}"
        safe_title = "".join(c if (c.isalnum() or c in "-_ ") else "_" for c in tab_title).strip().replace(" ", "_")
        filename = f"{base}_Tab{tab_id + 1}_{safe_title}.png"
        response = Response(content=image, media_type="image/png")
        response.headers["Content-Disposition"] = f'attachment; filename="{filename}"'
    except Exception as exc:
        response = Response(content=f"Save plot error: {exc}", media_type="text/plain", status_code=500)
    return response_with_session_cookie(response, state, is_new)


def _sanitize_export_view_state(value: object) -> dict:
    # Keep only safe view fields sent by this browser for static export.
    if not isinstance(value, dict):
        return {}
    out: dict = {}
    for key in ("xaxis_range", "yaxis_range", "yaxis2_range"):
        rng = value.get(key)
        if isinstance(rng, list) and len(rng) == 2:
            out[key] = [rng[0], rng[1]]
    if str(value.get("dragmode", "")) in ("zoom", "pan"):
        out["dragmode"] = str(value.get("dragmode"))
    vis = value.get("visible_signals")
    if isinstance(vis, dict):
        out["visible_signals"] = {str(k)[:240]: bool(v) for k, v in list(vis.items())[:500]}
    return out


def _parse_export_view(raw: Optional[str]) -> dict:
    if not raw or len(str(raw)) > 20000:
        return {}
    try:
        return _sanitize_export_view_state(json.loads(str(raw)))
    except Exception:
        return {}


def _parse_export_views(raw: Optional[str]) -> Dict[int, dict]:
    if not raw or len(str(raw)) > 60000:
        return {}
    try:
        payload = json.loads(str(raw))
    except Exception:
        return {}
    if not isinstance(payload, dict):
        return {}
    out: Dict[int, dict] = {}
    for key, value in list(payload.items())[:len(TAB_PRESETS)]:
        try:
            tab_id = int(key)
        except Exception:
            continue
        if 0 <= tab_id < len(TAB_PRESETS):
            out[tab_id] = _sanitize_export_view_state(value)
    return out


def _parse_report_tab_ids(raw_tabs: str | None) -> List[int]:
    """Parse a comma-separated zero-based tab list. Empty means all tabs."""
    if raw_tabs is None or not str(raw_tabs).strip():
        return list(range(len(TAB_PRESETS)))
    result: List[int] = []
    seen = set()
    for part in str(raw_tabs).split(","):
        part = part.strip()
        if not part:
            continue
        try:
            tab_id = int(part)
        except ValueError:
            continue
        if 0 <= tab_id < len(TAB_PRESETS) and tab_id not in seen:
            seen.add(tab_id)
            result.append(tab_id)
    return result


def _report_tab_ids_from_request(request: Request) -> List[int]:
    return _parse_report_tab_ids(request.query_params.get("tabs"))


def build_wide_pdf_bytes(state: BrowserSession, tab_ids: Optional[List[int]] = None, view_states: Optional[Dict[int, dict]] = None) -> bytes:
    """Create a multi-page PDF using a custom page size larger than A4 landscape.

    Each plot tab remains on its own PDF page.  The default page size is
    WIDE_REPORT_WIDTH_PT x WIDE_REPORT_HEIGHT_PT points.
    """
    selected_tab_ids = list(tab_ids) if tab_ids is not None else list(range(len(TAB_PRESETS)))
    if not selected_tab_ids:
        raise ValueError("Select at least one report page/tab.")
    stream = io.BytesIO()
    with PdfPages(stream) as pdf:
        for tab_id in selected_tab_ids:
            fig = build_figure(
                state,
                tab_id,
                report_header=True,
                figsize=WIDE_REPORT_FIGSIZE_IN,
                view_state=(view_states or {}).get(tab_id),
            )
            try:
                # bbox_inches=None is intentional: preserve the exact custom PDF page size.
                pdf.savefig(fig, bbox_inches=None, pad_inches=0)
            finally:
                plt.close(fig)
    stream.seek(0)
    return stream.getvalue()


@app.get("/report/pdf")
async def generate_pdf_report(request: Request):
    state, is_new = get_or_create_session(request)
    if state.df is None:
        response = Response(content="Please load a CSV first.", media_type="text/plain", status_code=400)
        return response_with_session_cookie(response, state, is_new)

    try:
        tab_ids = _report_tab_ids_from_request(request)
        if not tab_ids:
            raise ValueError("Select at least one report page/tab.")
        view_states = _parse_export_views(request.query_params.get("views"))
        pdf_bytes = build_wide_pdf_bytes(state, tab_ids, view_states=view_states)
        base = os.path.splitext(state.filename)[0] or "Nissan_OBC"
        filename = f"{base}_Test_Report_Wide.pdf"
        response = Response(content=pdf_bytes, media_type="application/pdf")
        response.headers["Content-Disposition"] = f'attachment; filename="{filename}"'
        state.status = f"Wide PDF Report generated successfully ({len(tab_ids)} selected page(s), {WIDE_REPORT_WIDTH_PT} x {WIDE_REPORT_HEIGHT_PT} pt)."
        state.status_type = "success"
    except Exception as exc:
        response = Response(content=f"PDF generation error: {exc}", media_type="text/plain", status_code=500)
    return response_with_session_cookie(response, state, is_new)


@app.get("/report/pdf-a4")
async def generate_pdf_report_a4(request: Request):
    """A4 landscape multi-page PDF export retained for compatibility."""
    state, is_new = get_or_create_session(request)
    if state.df is None:
        response = Response(content="Please load a CSV first.", media_type="text/plain", status_code=400)
        return response_with_session_cookie(response, state, is_new)
    stream = io.BytesIO()
    try:
        tab_ids = _report_tab_ids_from_request(request)
        if not tab_ids:
            raise ValueError("Select at least one report page/tab.")
        view_states = _parse_export_views(request.query_params.get("views"))
        with PdfPages(stream) as pdf:
            for tab_id in tab_ids:
                fig = build_figure(state, tab_id, report_header=True, view_state=view_states.get(tab_id))
                try:
                    pdf.savefig(fig, bbox_inches="tight", pad_inches=0.08)
                finally:
                    plt.close(fig)
        stream.seek(0)
        base = os.path.splitext(state.filename)[0] or "Nissan_OBC"
        filename = f"{base}_Test_Report_A4.pdf"
        response = Response(content=stream.getvalue(), media_type="application/pdf")
        response.headers["Content-Disposition"] = f'attachment; filename="{filename}"'
        state.status = f"A4 PDF Report generated successfully ({len(tab_ids)} selected page(s))."
        state.status_type = "success"
    except Exception as exc:
        response = Response(content=f"PDF generation error: {exc}", media_type="text/plain", status_code=500)
    return response_with_session_cookie(response, state, is_new)


def _png_pixel_size(png_bytes: bytes) -> Tuple[int, int]:
    # Return PNG pixel dimensions using only the standard library.
    if len(png_bytes) < 24 or png_bytes[:8] != b"\x89PNG\r\n\x1a\n":
        return 1600, 900
    try:
        width, height = struct.unpack(">II", png_bytes[16:24])
        return max(1, int(width)), max(1, int(height))
    except Exception:
        return 1600, 900


def _build_builtin_docx(image_pages: List[bytes]) -> bytes:
    # Minimal landscape DOCX: one rendered plot image per selected report page.
    if not image_pages:
        raise ValueError("No report pages were generated.")

    relationships = []
    paragraphs = []
    media = []
    max_width_emu = int(10.5 * 914400)
    max_height_emu = int(6.65 * 914400)

    for idx, image in enumerate(image_pages, start=1):
        width_px, height_px = _png_pixel_size(image)
        ratio = width_px / max(1.0, float(height_px))
        cx = max_width_emu
        cy = int(cx / max(0.01, ratio))
        if cy > max_height_emu:
            cy = max_height_emu
            cx = int(cy * ratio)
        rid = f"rId{idx}"
        relationships.append(
            f'<Relationship Id="{rid}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/image{idx}.png"/>'
        )
        paragraphs.append(f"""<w:p><w:pPr><w:jc w:val="center"/></w:pPr><w:r><w:drawing>
<wp:inline distT="0" distB="0" distL="0" distR="0">
<wp:extent cx="{cx}" cy="{cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>
<wp:docPr id="{idx}" name="SignalWorks Studio Plot {idx}"/><wp:cNvGraphicFramePr/>
<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">
<pic:pic><pic:nvPicPr><pic:cNvPr id="0" name="image{idx}.png"/><pic:cNvPicPr/></pic:nvPicPr>
<pic:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>
<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr>
</pic:pic></a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>""")
        if idx < len(image_pages):
            paragraphs.append('<w:p><w:r><w:br w:type="page"/></w:r></w:p>')
        media.append((f"word/media/image{idx}.png", image))

    document_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture"><w:body>{''.join(paragraphs)}<w:sectPr><w:pgSz w:w="16834" w:h="11909" w:orient="landscape"/><w:pgMar w:top="720" w:right="720" w:bottom="720" w:left="720" w:header="360" w:footer="360" w:gutter="0"/></w:sectPr></w:body></w:document>""".encode("utf-8")
    rels_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{''.join(relationships)}</Relationships>""".encode("utf-8")
    content_types = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Default Extension="png" ContentType="image/png"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>"""
    root_rels = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>"""

    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", root_rels)
        archive.writestr("word/document.xml", document_xml)
        archive.writestr("word/_rels/document.xml.rels", rels_xml)
        for name, image in media:
            archive.writestr(name, image)
    return output.getvalue()


@app.get("/report/word")
async def generate_word_report(request: Request):
    state, is_new = get_or_create_session(request)
    if state.df is None:
        response = Response(content="Please load a CSV first.", media_type="text/plain", status_code=400)
        return response_with_session_cookie(response, state, is_new)

    try:
        tab_ids = _report_tab_ids_from_request(request)
        if not tab_ids:
            raise ValueError("Select at least one report page/tab.")
        view_states = _parse_export_views(request.query_params.get("views"))

        # Render once; both engines consume the same PNG pages.
        page_images: List[bytes] = []
        for tab_id in tab_ids:
            fig = build_figure(state, tab_id, report_header=True, view_state=view_states.get(tab_id))
            image_stream = io.BytesIO()
            try:
                fig.savefig(image_stream, format="png", dpi=200, bbox_inches="tight", pad_inches=0.08)
            finally:
                plt.close(fig)
            page_images.append(image_stream.getvalue())

        used_builtin_fallback = False
        try:
            import docx
            from docx.enum.section import WD_ORIENT
            from docx.shared import Inches
        except ImportError:
            used_builtin_fallback = True
            docx_bytes = _build_builtin_docx(page_images)
        else:
            doc = docx.Document()
            section = doc.sections[0]
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width = Inches(11.69)
            section.page_height = Inches(8.27)
            section.left_margin = Inches(0.5)
            section.right_margin = Inches(0.5)
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)

            for page_index, image in enumerate(page_images):
                if page_index > 0:
                    doc.add_page_break()
                doc.add_picture(io.BytesIO(image), width=Inches(10.5))

            output = io.BytesIO()
            doc.save(output)
            docx_bytes = output.getvalue()

        base = os.path.splitext(state.filename)[0] or "SignalWorksStudio"
        filename = f"{base}_Test_Report.docx"
        response = Response(
            content=docx_bytes,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        response.headers["Content-Disposition"] = f'attachment; filename="{filename}"'
        engine_note = " · built-in DOCX fallback" if used_builtin_fallback else ""
        state.status = f"Word Report generated successfully ({len(tab_ids)} selected page(s)){engine_note}."
        state.status_type = "success"
    except Exception as exc:
        response = Response(content=f"Word generation error: {exc}", media_type="text/plain", status_code=500)

    return response_with_session_cookie(response, state, is_new)


# ------------------------------------------------------------------------------
# Start WebServer
# ------------------------------------------------------------------------------
def get_lan_ip() -> str:
    """Return the preferred non-loopback IPv4 address for LAN access."""
    candidates: List[str] = []

    sock = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        # This selects the address from the active route; no actual traffic is required.
        sock.connect(("8.8.8.8", 80))
        candidates.append(sock.getsockname()[0])
    except Exception:
        pass
    finally:
        sock.close()

    try:
        candidates.extend(socket.gethostbyname_ex(socket.gethostname())[2])
    except Exception:
        pass
    try:
        for info in socket.getaddrinfo(socket.gethostname(), None, socket.AF_INET, socket.SOCK_STREAM):
            candidates.append(info[4][0])
    except Exception:
        pass

    unique: List[str] = []
    for value in candidates:
        value = str(value or "").strip()
        if value and value not in unique and not value.startswith("127.") and not value.startswith("169.254."):
            unique.append(value)

    for value in unique:
        try:
            if ipaddress.ip_address(value).is_private:
                return value
        except ValueError:
            continue
    return unique[0] if unique else "127.0.0.1"


def open_browser_later(url: str, delay: float = 1.2) -> None:
    def _worker():
        time.sleep(delay)
        try:
            webbrowser.open(url)
        except Exception:
            pass

    threading.Thread(target=_worker, daemon=True).start()


def main() -> None:
    global _managed_uvicorn_server, _auto_shutdown_enabled, _idle_shutdown_seconds

    parser = argparse.ArgumentParser(description="SignalWorks Studio Engineering Data WebServer")
    parser.add_argument("--lan", action="store_true", help="Allow other devices on the same LAN to access the WebServer")
    parser.add_argument("--port", type=int, default=DEFAULT_PORT, help=f"WebServer port (default: {DEFAULT_PORT})")
    parser.add_argument("--no-browser", action="store_true", help="Do not automatically open the local web browser")
    parser.add_argument("--quiet", action="store_true", help="Suppress console output/logging (recommended for hidden launcher)")
    parser.add_argument("--no-auto-shutdown", action="store_true", help="Keep the server running after all browser tabs close")
    parser.add_argument(
        "--idle-shutdown-seconds",
        type=float,
        default=DEFAULT_IDLE_SHUTDOWN_SECONDS,
        help=f"Seconds to wait after the last browser tab closes before stopping (default: {DEFAULT_IDLE_SHUTDOWN_SECONDS:g})",
    )
    args = parser.parse_args()

    try:
        import uvicorn
    except ImportError as exc:
        raise SystemExit(
            "Missing Uvicorn. Install WebServer dependencies with:\n\n"
            "  pip install fastapi uvicorn python-multipart\n\n"
            f"Original error: {exc}"
        )

    host = "0.0.0.0" if args.lan else "127.0.0.1"
    local_url = f"http://127.0.0.1:{args.port}"
    lan_ip = get_lan_ip() if args.lan else ""
    lan_url = f"http://{lan_ip}:{args.port}" if lan_ip and lan_ip != "127.0.0.1" else local_url
    browser_open_url = lan_url if args.lan else local_url
    _auto_shutdown_enabled = not args.no_auto_shutdown
    _idle_shutdown_seconds = max(2.0, float(args.idle_shutdown_seconds))

    if not args.quiet:
        print("=" * 72)
        print(f"{APP_TITLE} - {APP_VERSION}")
        print("=" * 72)
        print(f"Local browser : {local_url}")
        if args.lan:
            print(f"LAN browser   : {lan_url}")
            print("WARNING       : LAN mode has no login/authentication.")
        else:
            print("Access mode   : Localhost only")
            print("LAN access    : Start with --lan if required")
        if _auto_shutdown_enabled:
            print(f"Auto shutdown : ON ({_idle_shutdown_seconds:g}s after all browser tabs close)")
        else:
            print("Auto shutdown : OFF")
        print("Press CTRL+C to stop the WebServer.")
        print("=" * 72)

    if not args.no_browser:
        # In LAN mode open the same address that other PCs should use. This makes
        # the selected access mode visible immediately instead of always showing
        # 127.0.0.1 in the primary PC browser.
        open_browser_later(browser_open_url)

    config = uvicorn.Config(
        app,
        host=host,
        port=args.port,
        log_level="critical" if args.quiet else "info",
        access_log=not args.quiet,
    )
    server = uvicorn.Server(config)
    _managed_uvicorn_server = server
    threading.Thread(target=lifecycle_watchdog, daemon=True, name="SignalWorksStudioLifecycleWatchdog").start()

    try:
        server.run()
    finally:
        _managed_uvicorn_server = None


if __name__ == "__main__":
    main()
